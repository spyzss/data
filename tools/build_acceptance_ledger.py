#!/usr/bin/env python3
"""Build a reusable supplier acceptance ledger from file-contract outputs."""

from __future__ import annotations

import argparse
import json
import logging
import os
import shutil
import sys
import tempfile
from dataclasses import dataclass
from pathlib import Path
from typing import Any, Iterable

import pandas as pd
import yaml
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter


LOGGER = logging.getLogger("build_acceptance_ledger")
GENERATED_SHEETS = (
    "Overview",
    "Asset_Ledger",
    "JDT_Manual_Review",
    "JDT_Review_Queue",
    "DR_Blocker",
    "Data_Dictionary",
)
OVERVIEW_COLUMNS = (
    "supplier_id",
    "total_clip_count",
    "total_frame_count",
    "pass_frame_count",
    "fail_frame_count",
    "review_frame_count",
    "pass_clip_count",
    "fail_clip_count",
    "review_clip_count",
    "sam3_status",
    "manual_review_count",
    "manual_review_window_count",
    "manual_review_asset_count",
    "manual_review_label_rows",
    "note",
)
ASSET_COLUMNS = (
    "supplier_id",
    "asset_id",
    "source_start_frame",
    "source_end_frame",
    "total_frame_count",
    "pass_frame_count",
    "fail_frame_count",
    "review_frame_count",
    "final_status",
    "sam3_status",
    "manual_label_rows",
    "unresolved_review_window_count",
    "fail_intervals_json",
    "review_intervals_json",
    "note",
)


@dataclass
class SupplierLedger:
    overview: dict[str, Any]
    asset_rows: list[dict[str, Any]]
    manual_rows: list[dict[str, Any]]
    review_rows: list[dict[str, Any]]
    blocker_rows: list[dict[str, Any]]


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        description="Build an XLSX acceptance ledger from supplier module outputs."
    )
    parser.add_argument("--quality-archive", required=True, type=Path)
    parser.add_argument("--config", required=True, type=Path)
    parser.add_argument("--output", required=True, type=Path)
    parser.add_argument("--existing-workbook", type=Path)
    parser.add_argument("--overwrite", action="store_true")
    parser.add_argument("--log-level", default="INFO")
    return parser.parse_args()


def main() -> int:
    args = parse_args()
    logging.basicConfig(
        level=getattr(logging, args.log_level.upper()),
        format="%(asctime)s [%(levelname)s] %(name)s: %(message)s",
    )
    output = build_acceptance_ledger(
        config_path=args.config,
        output_path=args.output,
        existing_workbook=args.existing_workbook,
        overwrite=args.overwrite,
        quality_archive=args.quality_archive,
    )
    LOGGER.info("Wrote %s", output)
    return 0


def read_records(path: Path) -> list[dict[str, Any]]:
    suffix = path.suffix.lower()
    if suffix == ".csv":
        frame = pd.read_csv(path, dtype=str, keep_default_na=False)
        rows = frame.to_dict(orient="records")
    elif suffix == ".parquet":
        rows = pd.read_parquet(path).to_dict(orient="records")
    elif suffix in {".jsonl", ".ndjson"}:
        rows = [
            json.loads(line)
            for line in path.read_text(encoding="utf-8").splitlines()
            if line.strip()
        ]
    elif suffix == ".json":
        payload = json.loads(path.read_text(encoding="utf-8"))
        if isinstance(payload, list):
            rows = payload
        elif isinstance(payload, dict):
            sequence = next(
                (value for value in payload.values() if isinstance(value, list)),
                None,
            )
            rows = sequence if sequence is not None else [payload]
        else:
            raise ValueError(f"JSON input must contain an object or list: {path}")
    else:
        raise ValueError(f"unsupported input format {suffix!r}: {path}")
    return [
        {str(key): _plain_value(value) for key, value in row.items()}
        for row in rows
        if isinstance(row, dict)
    ]


def merge_intervals(intervals: Iterable[tuple[int, int]]) -> list[tuple[int, int]]:
    normalized = sorted((min(start, end), max(start, end)) for start, end in intervals)
    merged: list[list[int]] = []
    for start, end in normalized:
        if not merged or start > merged[-1][1] + 1:
            merged.append([start, end])
        else:
            merged[-1][1] = max(merged[-1][1], end)
    return [(start, end) for start, end in merged]


def subtract_intervals(
    intervals: Iterable[tuple[int, int]],
    excluded: Iterable[tuple[int, int]],
) -> list[tuple[int, int]]:
    remaining: list[tuple[int, int]] = []
    exclusions = merge_intervals(excluded)
    for start, end in merge_intervals(intervals):
        pieces = [(start, end)]
        for cut_start, cut_end in exclusions:
            next_pieces: list[tuple[int, int]] = []
            for piece_start, piece_end in pieces:
                if cut_end < piece_start or cut_start > piece_end:
                    next_pieces.append((piece_start, piece_end))
                    continue
                if piece_start < cut_start:
                    next_pieces.append((piece_start, cut_start - 1))
                if cut_end < piece_end:
                    next_pieces.append((cut_end + 1, piece_end))
            pieces = next_pieces
        remaining.extend(pieces)
    return merge_intervals(remaining)


def interval_frame_count(intervals: Iterable[tuple[int, int]]) -> int:
    return sum(end - start + 1 for start, end in merge_intervals(intervals))


def build_supplier_ledger(
    supplier_config: dict[str, Any],
    *,
    config_dir: Path,
) -> SupplierLedger:
    supplier_id = str(supplier_config.get("supplier_id") or "").strip()
    if not supplier_id:
        raise ValueError("supplier configuration missing supplier_id")
    required_inputs = {
        str(value) for value in supplier_config.get("required_inputs", [])
    }
    notes: list[str] = []

    manifest_rows = _load_configured_records(
        supplier_config,
        "manifest",
        config_dir=config_dir,
        required="manifest" in required_inputs,
        notes=notes,
    )
    manual_rows = _load_configured_records(
        supplier_config,
        "manual_labels",
        config_dir=config_dir,
        required="manual_labels" in required_inputs,
        notes=notes,
    )
    review_rows = _load_configured_records(
        supplier_config,
        "review_queue",
        config_dir=config_dir,
        required="review_queue" in required_inputs,
        notes=notes,
    )
    sam3_rows = _load_configured_records(
        supplier_config,
        "sam3_summary",
        config_dir=config_dir,
        required="sam3_summary" in required_inputs,
        notes=notes,
    )
    blocker_rows = _load_configured_records(
        supplier_config,
        "blocker",
        config_dir=config_dir,
        required="blocker" in required_inputs,
        notes=notes,
        note_if_unconfigured=False,
    )
    sam3_output_available = _configured_input_exists(
        supplier_config, "sam3_summary", config_dir
    )

    blocked = _as_bool(supplier_config.get("blocked", False))
    blocked_policy = str(supplier_config.get("blocked_policy") or "").strip()
    if blocked and blocked_policy not in {"", "all_frames_review"}:
        raise ValueError(
            f"unsupported blocked_policy for {supplier_id}: {blocked_policy}"
        )
    if blocked:
        sam3_status = "blocked"
        blocker_reason = _first_nonempty(
            row.get("reason") for row in blocker_rows
        )
        notes.append(blocker_reason or "supplier module blocked")
    elif sam3_output_available:
        sam3_status = "evaluated"
    else:
        sam3_status = "not_evaluated"

    assets = _manifest_assets(manifest_rows, supplier_id)
    manual_by_asset = _group_by_asset(manual_rows)
    review_by_asset = _group_by_asset(review_rows)
    _validate_known_assets(manual_by_asset, assets, supplier_id, "manual_labels")
    _validate_known_assets(review_by_asset, assets, supplier_id, "review_queue")
    _validate_review_windows(review_rows, supplier_id)
    resolved_review_ids, resolved_windows = _resolved_manual_keys(manual_rows)
    asset_rows: list[dict[str, Any]] = []
    for asset_id, asset in assets.items():
        source_bounds = (asset["source_start_frame"], asset["source_end_frame"])
        asset_manual_rows = manual_by_asset.get(asset_id, [])
        asset_review_rows = review_by_asset.get(asset_id, [])
        fail_intervals = _manual_fail_intervals(asset_manual_rows, source_bounds)
        unresolved = [
            row
            for row in asset_review_rows
            if not _review_row_is_resolved(
                row,
                resolved_review_ids=resolved_review_ids,
                resolved_windows=resolved_windows,
            )
        ]
        review_intervals = _window_intervals(unresolved, source_bounds)
        if blocked and blocked_policy == "all_frames_review":
            fail_intervals = []
            review_intervals = [source_bounds]
        else:
            review_intervals = subtract_intervals(review_intervals, fail_intervals)

        fail_count = interval_frame_count(fail_intervals)
        review_count = interval_frame_count(review_intervals)
        pass_count = asset["total_frame_count"] - fail_count - review_count
        if pass_count < 0:
            raise ValueError(f"negative pass frame count for {supplier_id}/{asset_id}")
        final_status = "fail" if fail_count else "review" if review_count else "pass"
        row = {
            "supplier_id": supplier_id,
            "asset_id": asset_id,
            **asset,
            "pass_frame_count": pass_count,
            "fail_frame_count": fail_count,
            "review_frame_count": review_count,
            "final_status": final_status,
            "sam3_status": sam3_status,
            "manual_label_rows": len(asset_manual_rows),
            "unresolved_review_window_count": len(unresolved),
            "fail_intervals_json": _intervals_json(fail_intervals),
            "review_intervals_json": _intervals_json(review_intervals),
            "note": "all frames review: blocked module"
            if blocked and blocked_policy == "all_frames_review"
            else "",
        }
        _validate_asset_balance(row)
        asset_rows.append(row)

    manual_review_ids = {
        str(row.get("review_id")).strip()
        for row in manual_rows
        if str(row.get("review_id") or "").strip()
    }
    manual_review_assets = {
        str(row.get("asset_id")).strip()
        for row in manual_rows
        if str(row.get("asset_id") or "").strip()
    }
    manual_review_window_count = len(manual_review_ids)
    overview = {
        "supplier_id": supplier_id,
        "total_clip_count": len(asset_rows),
        "total_frame_count": sum(row["total_frame_count"] for row in asset_rows),
        "pass_frame_count": sum(row["pass_frame_count"] for row in asset_rows),
        "fail_frame_count": sum(row["fail_frame_count"] for row in asset_rows),
        "review_frame_count": sum(row["review_frame_count"] for row in asset_rows),
        "pass_clip_count": sum(row["final_status"] == "pass" for row in asset_rows),
        "fail_clip_count": sum(row["final_status"] == "fail" for row in asset_rows),
        "review_clip_count": sum(row["final_status"] == "review" for row in asset_rows),
        "sam3_status": sam3_status,
        "manual_review_count": manual_review_window_count,
        "manual_review_window_count": manual_review_window_count,
        "manual_review_asset_count": len(manual_review_assets),
        "manual_review_label_rows": len(manual_rows),
        "note": "; ".join(_dedupe(notes)),
    }
    _validate_overview_balance(overview)
    return SupplierLedger(
        overview=overview,
        asset_rows=asset_rows,
        manual_rows=manual_rows,
        review_rows=review_rows,
        blocker_rows=blocker_rows,
    )


def build_acceptance_ledger(
    *,
    config_path: Path,
    output_path: Path,
    overwrite: bool = False,
    existing_workbook: Path | None = None,
    quality_archive: Path | None = None,
) -> Path:
    config_path = Path(config_path)
    output_path = Path(output_path)
    if output_path.exists() and not overwrite:
        raise FileExistsError(f"output exists; pass --overwrite: {output_path}")
    if quality_archive is not None:
        # The unified QC JSON projection is the formal ledger.  Keep the
        # historical function/API usable for callers that omit this argument,
        # but never let legacy supplier inputs override a canonical verdict.
        from tools.build_qc_json_projection import run_projection_cli

        paths = run_projection_cli(
            Path(quality_archive),
            output_path.parent,
            formats=("csv", "parquet", "xlsx", "markdown"),
        )
        projected_xlsx = paths.get("xlsx")
        if projected_xlsx is None:
            raise RuntimeError("QC projection did not produce an XLSX output")
        output_path.parent.mkdir(parents=True, exist_ok=True)
        shutil.copyfile(projected_xlsx, output_path)
        return output_path
    config = yaml.safe_load(config_path.read_text(encoding="utf-8")) or {}
    if config.get("workbook_mode") == "weekly_template":
        try:
            from tools.acceptance_ledger_weekly import build_weekly_template_workbook
        except ModuleNotFoundError as exc:  # Direct: python tools/<script>.py
            if exc.name != "tools":
                raise
            from acceptance_ledger_weekly import build_weekly_template_workbook

        return build_weekly_template_workbook(
            config=config,
            config_path=config_path,
            output_path=output_path,
            existing_workbook=existing_workbook,
        )
    suppliers = config.get("suppliers")
    if not isinstance(suppliers, list) or not suppliers:
        raise ValueError("config suppliers must be a non-empty list")
    results = [
        build_supplier_ledger(dict(supplier), config_dir=config_path.parent)
        for supplier in suppliers
    ]
    workbook = _load_or_create_workbook(existing_workbook)
    workbook.properties.title = str(config.get("output_label") or "Acceptance Ledger")
    _replace_generated_sheets(workbook)
    _write_sheet(workbook, "Overview", OVERVIEW_COLUMNS, [row.overview for row in results])
    _write_sheet(
        workbook,
        "Asset_Ledger",
        ASSET_COLUMNS,
        [asset for result in results for asset in result.asset_rows],
    )
    jdt = next((result for result in results if result.overview["supplier_id"] == "jdt"), None)
    deepreach = next(
        (result for result in results if result.overview["supplier_id"] == "deepreach"),
        None,
    )
    _write_dynamic_sheet(
        workbook,
        "JDT_Manual_Review",
        [] if jdt is None else jdt.manual_rows,
        fallback_columns=("review_id", "asset_id", "manual_outcome"),
    )
    _write_dynamic_sheet(
        workbook,
        "JDT_Review_Queue",
        [] if jdt is None else jdt.review_rows,
        fallback_columns=("review_id", "asset_id", "window_start_frame", "window_end_frame"),
    )
    _write_dynamic_sheet(
        workbook,
        "DR_Blocker",
        [] if deepreach is None else deepreach.blocker_rows,
        fallback_columns=("status", "reason"),
    )
    _write_sheet(
        workbook,
        "Data_Dictionary",
        ("sheet", "column", "definition"),
        _data_dictionary_rows(),
    )
    _save_workbook_atomic(workbook, output_path)
    return output_path


def _load_configured_records(
    config: dict[str, Any],
    key: str,
    *,
    config_dir: Path,
    required: bool,
    notes: list[str],
    note_if_unconfigured: bool = True,
) -> list[dict[str, Any]]:
    value = config.get(key)
    if not value:
        if required:
            raise FileNotFoundError(f"required input {key} is not configured")
        if note_if_unconfigured:
            notes.append(f"{key}=not_evaluated")
        return []
    path = _resolve_input_path(value, config_dir)
    if not path.exists():
        if required:
            raise FileNotFoundError(f"required input {key} not found: {path}")
        notes.append(f"{key}=not_evaluated")
        return []
    return read_records(path)


def _resolve_input_path(value: Any, config_dir: Path) -> Path:
    path = Path(str(value)).expanduser()
    if path.is_absolute():
        return path
    cwd_path = (Path.cwd() / path).resolve()
    if cwd_path.exists():
        return cwd_path
    config_path = (config_dir / path).resolve()
    return config_path if config_path.exists() else cwd_path


def _configured_input_exists(
    config: dict[str, Any], key: str, config_dir: Path
) -> bool:
    value = config.get(key)
    return bool(value) and _resolve_input_path(value, config_dir).exists()


def _manifest_assets(
    rows: list[dict[str, Any]], supplier_id: str
) -> dict[str, dict[str, int]]:
    assets: dict[str, dict[str, int]] = {}
    for index, row in enumerate(rows):
        asset_id = str(row.get("asset_id") or "").strip()
        if not asset_id:
            raise ValueError(f"manifest row {index} missing asset_id for {supplier_id}")
        if asset_id in assets:
            raise ValueError(f"duplicate manifest asset_id for {supplier_id}: {asset_id}")
        start = _first_int(
            row,
            ("start_frame", "source_start_frame", "clip_start_frame"),
        )
        end = _first_int(
            row,
            ("end_frame", "source_end_frame", "clip_end_frame"),
        )
        has_complete_bounds = start is not None and end is not None
        if has_complete_bounds:
            frame_count = end - start + 1
        else:
            frame_count = _first_int(
                row,
                (
                    "clip_frame_count",
                    "total_frame_count",
                    "total_frames",
                    "num_frames",
                    "frame_count",
                ),
            )
        if frame_count is None:
            raise ValueError(
                f"manifest asset {supplier_id}/{asset_id} lacks frame count or bounds"
            )
        if frame_count <= 0:
            raise ValueError(f"manifest asset {supplier_id}/{asset_id} has no frames")
        source_start = 0 if start is None else start
        source_end = end if has_complete_bounds else source_start + frame_count - 1
        assets[asset_id] = {
            "source_start_frame": source_start,
            "source_end_frame": source_end,
            "total_frame_count": frame_count,
        }
    return assets


def _group_by_asset(rows: list[dict[str, Any]]) -> dict[str, list[dict[str, Any]]]:
    grouped: dict[str, list[dict[str, Any]]] = {}
    for index, row in enumerate(rows):
        asset_id = str(row.get("asset_id") or "").strip()
        if not asset_id:
            raise ValueError(f"input row {index} missing asset_id")
        grouped.setdefault(asset_id, []).append(row)
    return grouped


def _validate_known_assets(
    grouped: dict[str, list[dict[str, Any]]],
    assets: dict[str, dict[str, int]],
    supplier_id: str,
    input_name: str,
) -> None:
    unknown = sorted(set(grouped) - set(assets))
    if unknown:
        raise ValueError(
            f"{input_name} contains unknown assets for {supplier_id}: {unknown}"
        )


def _validate_review_windows(rows: list[dict[str, Any]], supplier_id: str) -> None:
    for index, row in enumerate(rows):
        start, end = _row_window(row)
        if start is None or end is None:
            raise ValueError(
                f"review_queue row {index} for {supplier_id} lacks window bounds"
            )


def _resolved_manual_keys(
    rows: list[dict[str, Any]],
) -> tuple[set[str], set[tuple[str, int, int]]]:
    review_ids: set[str] = set()
    windows: set[tuple[str, int, int]] = set()
    for row in rows:
        if normalize_manual_outcome(row) == "review":
            continue
        review_id = str(row.get("review_id") or "").strip()
        if review_id:
            review_ids.add(review_id)
        start, end = _row_window(row)
        asset_id = str(row.get("asset_id") or "").strip()
        if asset_id and start is not None and end is not None:
            windows.add((asset_id, start, end))
    return review_ids, windows


def _review_row_is_resolved(
    row: dict[str, Any],
    *,
    resolved_review_ids: set[str],
    resolved_windows: set[tuple[str, int, int]],
) -> bool:
    review_id = str(row.get("review_id") or "").strip()
    if review_id and review_id in resolved_review_ids:
        return True
    start, end = _row_window(row)
    if start is None or end is None:
        return False
    return (str(row.get("asset_id") or "").strip(), start, end) in resolved_windows


def _manual_fail_intervals(
    rows: list[dict[str, Any]], bounds: tuple[int, int]
) -> list[tuple[int, int]]:
    intervals: list[tuple[int, int]] = []
    for row in rows:
        if normalize_manual_outcome(row) != "fail":
            continue
        start = _optional_int(row.get("affected_start_frame"))
        end = _optional_int(row.get("affected_end_frame"))
        if start is None or end is None:
            start, end = _row_window(row)
        if start is None or end is None:
            continue
        clipped = _clip_interval((start, end), bounds)
        if clipped is not None:
            intervals.append(clipped)
    return merge_intervals(intervals)


def normalize_manual_outcome(row: dict[str, Any]) -> str:
    """Normalize completed manual decisions for every ledger presentation."""
    outcome = str(row.get("manual_outcome") or "").strip().lower()
    status = str(row.get("acceptance_status") or "").strip().lower()
    if outcome in {"true_positive", "positive", "fail"} or status in {
        "rejected",
        "fail",
    }:
        return "fail"
    if outcome in {"false_positive", "acceptable_flagged", "pass"} or status in {
        "accepted",
        "pass",
    }:
        return "pass"
    return "review"


def _window_intervals(
    rows: list[dict[str, Any]], bounds: tuple[int, int]
) -> list[tuple[int, int]]:
    intervals: list[tuple[int, int]] = []
    for row in rows:
        start, end = _row_window(row)
        if start is None or end is None:
            continue
        clipped = _clip_interval((start, end), bounds)
        if clipped is not None:
            intervals.append(clipped)
    return merge_intervals(intervals)


def _row_window(row: dict[str, Any]) -> tuple[int | None, int | None]:
    start = _first_int(
        row,
        ("window_start_frame", "candidate_start_frame", "start_frame"),
    )
    end = _first_int(
        row,
        ("window_end_frame", "candidate_end_frame", "end_frame"),
    )
    return start, end


def _clip_interval(
    interval: tuple[int, int], bounds: tuple[int, int]
) -> tuple[int, int] | None:
    start, end = sorted(interval)
    clipped = (max(start, bounds[0]), min(end, bounds[1]))
    return clipped if clipped[0] <= clipped[1] else None


def _validate_asset_balance(row: dict[str, Any]) -> None:
    counted = (
        row["pass_frame_count"]
        + row["fail_frame_count"]
        + row["review_frame_count"]
    )
    if counted != row["total_frame_count"]:
        raise ValueError(f"frame totals do not balance for {row['asset_id']}")


def _validate_overview_balance(row: dict[str, Any]) -> None:
    if (
        row["pass_frame_count"]
        + row["fail_frame_count"]
        + row["review_frame_count"]
        != row["total_frame_count"]
    ):
        raise ValueError(f"supplier frame totals do not balance: {row['supplier_id']}")
    if (
        row["pass_clip_count"]
        + row["fail_clip_count"]
        + row["review_clip_count"]
        != row["total_clip_count"]
    ):
        raise ValueError(f"supplier clip totals do not balance: {row['supplier_id']}")


def _load_or_create_workbook(existing_workbook: Path | None):
    if existing_workbook is not None:
        path = Path(existing_workbook)
        if not path.exists():
            raise FileNotFoundError(f"existing workbook not found: {path}")
        return load_workbook(path)
    workbook = Workbook()
    workbook.remove(workbook.active)
    return workbook


def _replace_generated_sheets(workbook) -> None:
    temporary = None
    if workbook.sheetnames and all(name in GENERATED_SHEETS for name in workbook.sheetnames):
        temporary = workbook.create_sheet("__ledger_build_tmp__")
    for name in GENERATED_SHEETS:
        if name in workbook.sheetnames:
            workbook.remove(workbook[name])
    if temporary is not None:
        workbook.remove(temporary)


def _write_sheet(
    workbook,
    name: str,
    columns: Iterable[str],
    rows: Iterable[dict[str, Any]],
) -> None:
    columns = tuple(columns)
    sheet = workbook.create_sheet(name)
    sheet.append(list(columns))
    for row in rows:
        sheet.append([_excel_value(row.get(column, "")) for column in columns])
    _format_sheet(sheet)


def _write_dynamic_sheet(
    workbook,
    name: str,
    rows: list[dict[str, Any]],
    *,
    fallback_columns: Iterable[str],
) -> None:
    columns = list(fallback_columns)
    for row in rows:
        for key in row:
            if key not in columns:
                columns.append(key)
    _write_sheet(workbook, name, columns, rows)


def _format_sheet(sheet) -> None:
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    fill = PatternFill("solid", fgColor="D9EAF7")
    for cell in sheet[1]:
        cell.font = Font(bold=True)
        cell.fill = fill
    for column_index, cells in enumerate(sheet.columns, start=1):
        width = min(
            48,
            max(10, max(len(str(cell.value or "")) for cell in cells) + 2),
        )
        sheet.column_dimensions[get_column_letter(column_index)].width = width


def _save_workbook_atomic(workbook, output_path: Path) -> None:
    output_path.parent.mkdir(parents=True, exist_ok=True)
    with tempfile.NamedTemporaryFile(
        suffix=".xlsx",
        dir=output_path.parent,
        delete=False,
    ) as handle:
        temporary = Path(handle.name)
    try:
        workbook.save(temporary)
        os.replace(temporary, output_path)
    finally:
        temporary.unlink(missing_ok=True)


def _data_dictionary_rows() -> list[dict[str, str]]:
    definitions = {
        "total_clip_count": "Manifest asset row count.",
        "total_frame_count": "Inclusive source-frame count across manifest assets.",
        "fail_frame_count": "Union of manual true-positive/rejected affected intervals.",
        "review_frame_count": "Union of unresolved review windows after removing fail frames.",
        "pass_frame_count": "Total frames minus fail and review frames.",
        "final_status": "Per-clip priority: fail, then review, then pass.",
        "sam3_status": "evaluated, blocked, or not_evaluated from configured inputs.",
        "manual_review_count": (
            "Legacy alias for manual_review_window_count."
        ),
        "manual_review_window_count": (
            "Distinct non-empty review_id values submitted in manual labels."
        ),
        "manual_review_asset_count": (
            "Distinct asset_id values represented in submitted manual labels."
        ),
        "manual_review_label_rows": "Number of submitted manual-label rows.",
        "manual_label_rows": "Manual-label row count for this asset.",
        "note": "Input availability, blocker, or policy context.",
    }
    rows = []
    for sheet, columns in (("Overview", OVERVIEW_COLUMNS), ("Asset_Ledger", ASSET_COLUMNS)):
        rows.extend(
            {
                "sheet": sheet,
                "column": column,
                "definition": definitions.get(column, column.replace("_", " ").capitalize()),
            }
            for column in columns
        )
    return rows


def _first_int(row: dict[str, Any], keys: Iterable[str]) -> int | None:
    for key in keys:
        value = _optional_int(row.get(key))
        if value is not None:
            return value
    return None


def _optional_int(value: Any) -> int | None:
    if value in (None, ""):
        return None
    try:
        number = float(value)
    except (TypeError, ValueError):
        return None
    if not number.is_integer():
        return None
    return int(number)


def _as_bool(value: Any) -> bool:
    if isinstance(value, bool):
        return value
    return str(value).strip().lower() in {"1", "true", "yes", "on"}


def _plain_value(value: Any) -> Any:
    if value is None:
        return ""
    if hasattr(value, "item"):
        try:
            return value.item()
        except ValueError:
            pass
    return value


def _excel_value(value: Any) -> Any:
    if isinstance(value, (dict, list, tuple, set)):
        return json.dumps(value, ensure_ascii=False, sort_keys=True)
    if isinstance(value, str) and value.startswith(("=", "+", "-", "@")):
        return "'" + value
    return _plain_value(value)


def _intervals_json(intervals: list[tuple[int, int]]) -> str:
    return json.dumps(
        [{"start_frame": start, "end_frame": end} for start, end in intervals]
    )


def _first_nonempty(values: Iterable[Any]) -> str:
    return next((str(value) for value in values if value not in (None, "")), "")


def _dedupe(values: Iterable[str]) -> list[str]:
    return list(dict.fromkeys(value for value in values if value))


if __name__ == "__main__":
    raise SystemExit(main())
