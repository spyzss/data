#!/usr/bin/env python3
"""Build the minimal weekly five-supplier acceptance workbook."""

from __future__ import annotations

import argparse
import csv
import json
import logging
import subprocess
from collections import Counter, defaultdict
from dataclasses import dataclass
from pathlib import Path
from typing import Any, Iterable, Sequence

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from tools.build_batch_qc_ledger import normalize_video_quality_row


LOGGER = logging.getLogger("build_weekly_supplier_acceptance_report")

SUMMARY_COLUMNS = [
    "supplier_name",
    "sample_clip_count",
    "expected_clip_count",
    "total_frame_count",
    "problem_frame_count",
    "problem_frame_ratio",
    "pass_clip_count",
    "fail_clip_count",
    "review_clip_count",
    "blocked_clip_count",
    "not_run_clip_count",
    "coverage_ratio",
    "pass_clip_ratio",
    "fail_clip_ratio",
    "main_issue_type",
    "modules_completed",
    "blocked_modules",
    "notes",
]

XJGT_DETAIL_COLUMNS = [
    "asset_id",
    "total_frames",
    "frame_count_status",
    "text_check_status",
    "skeleton_missing_status",
    "skeleton_missing_fail_intervals",
    "skeleton_missing_fail_frame_count",
    "skeleton_missing_fail_frame_ratio",
    "skeleton_morphology_status",
    "skeleton_morphology_fail_intervals",
    "skeleton_morphology_fail_frame_count",
    "skeleton_morphology_fail_frame_ratio",
    "skeleton_static_fail_frame_count",
    "skeleton_static_fail_intervals",
    "skeleton_static_fail_frame_ratio",
    "skeleton_static_status",
    "supplier_quality_signal",
    "video_quality_status",
    "video_quality_fail_frame_count",
    "video_quality_fail_frame_ratio",
    "temporal_status",
    "sam3_containment_status",
    "manual_review_status",
    "temporal_sam3_manual_status",
    "manual_problem_frame_count",
    "manual_problem_frame_ratio_of_clip",
    "manual_reviewed_frame_count",
    "manual_problem_ratio_of_reviewed",
    "abnormal_fail_frame_count",
    "abnormal_fail_frame_ratio",
    "abnormal_frame_status",
    "auto_fail_frame_count",
    "reviewed_auto_fail_frame_count",
    "reviewed_auto_fail_true_positive_frame_count",
    "reviewed_auto_fail_false_positive_frame_count",
    "unreviewed_auto_fail_frame_count",
    "auto_fail_precision_on_reviewed",
    "abnormal_status_reason",
    "fail_indicator_count",
    "final_clip_status",
    "mapped_precheck_checks",
    "missing_expected_checks",
    "precheck_mapping_status",
    "final_status_reason",
    "main_issue_type",
    "notes",
    "evidence_path",
]
DETAIL_COLUMNS = XJGT_DETAIL_COLUMNS

HARD_ISSUE_COLUMNS = [
    "issue_type",
    "current_detection_method",
    "current_threshold",
    "auto_decision_available",
    "needs_manual_review",
    "observed_count",
    "observed_unit",
    "denominator",
    "denominator_definition",
    "observed_ratio",
    "source_module",
    "why_hard",
    "next_action",
]

SHEET_NAMES = [
    "五供应商总览",
    "星际归途",
    "DeepReach",
    "供应商3",
    "供应商4",
    "供应商5",
    "人工与难测问题统计",
]

XJGT_FALLBACK_COUNTS = {
    "fail": 21,
    "review": 64,
    "pass_with_notes": 15,
}


@dataclass(frozen=True)
class WeeklyOutputPaths:
    workbook_xlsx: Path
    summary_csv: Path


@dataclass
class SourceRead:
    input_name: str
    path: Path | None
    records: list[dict[str, Any]]
    status: str


@dataclass
class PrecheckEvidence:
    by_asset: dict[str, dict[str, Any]]
    source_status: str
    unmatched_assets: list[str]


def parse_args(argv: Sequence[str] | None = None) -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        description="Build the minimal weekly five-supplier acceptance report."
    )
    parser.add_argument(
        "--run-root",
        type=Path,
        default=Path("outputs/acceptance_5x100"),
    )
    parser.add_argument(
        "--audit-inputs",
        action="store_true",
        help="Inspect candidate input schemas and mapping coverage without writing outputs.",
    )
    parser.add_argument(
        "--require-xjgt-text",
        action="store_true",
        help=(
            "Treat missing XJGT text_integrity rows as a required unresolved "
            "rule. The default is not_applicable."
        ),
    )
    parser.add_argument("--log-level", default="INFO")
    return parser.parse_args(argv)


def main(argv: Sequence[str] | None = None) -> int:
    args = parse_args(argv)
    logging.basicConfig(
        level=getattr(logging, args.log_level.upper()),
        format="%(asctime)s [%(levelname)s] %(name)s: %(message)s",
    )
    if args.audit_inputs:
        print(json.dumps(collect_input_audit(args.run_root), indent=2, ensure_ascii=False))
        return 0
    outputs = build_weekly_report(
        args.run_root,
        require_xjgt_text=args.require_xjgt_text,
    )
    LOGGER.info("Wrote %s", outputs.workbook_xlsx)
    LOGGER.info("Wrote %s", outputs.summary_csv)
    return 0


def build_weekly_report(
    run_root: Path,
    *,
    require_xjgt_text: bool = False,
) -> WeeklyOutputPaths:
    run_root.mkdir(parents=True, exist_ok=True)
    xjgt = load_xjgt(run_root, require_xjgt_text=require_xjgt_text)
    deepreach = load_deepreach(run_root)
    placeholders = [placeholder_supplier(index) for index in range(3, 6)]
    summary_rows = [xjgt["summary"], deepreach["summary"], *placeholders]

    summary_csv = run_root / "weekly_supplier_summary.csv"
    workbook_xlsx = run_root / "weekly_supplier_acceptance_report.xlsx"
    write_csv(summary_csv, SUMMARY_COLUMNS, summary_rows)
    write_workbook(
        workbook_xlsx,
        summary_rows,
        xjgt["details"],
        deepreach["details"],
        xjgt["issue_events"],
    )
    print_sanity_checks(
        xjgt["summary"],
        xjgt["details"],
        deepreach["summary"],
    )
    return WeeklyOutputPaths(workbook_xlsx, summary_csv)


def load_xjgt_source_reads(run_root: Path) -> dict[str, SourceRead]:
    precheck_dir = run_root / "xjgt" / "precheck"
    video_dir = run_root / "xjgt" / "video_quality"
    sam3_dir = run_root / "xjgt" / "sam3_containment"
    ledger_dir = run_root / "xjgt" / "ledger"
    manual_dir = run_root / "xjgt" / "manual_review"
    return {
        "precheck_check_results": read_first_records(
            "precheck_check_results",
            [
                precheck_dir / "check_results.parquet",
                precheck_dir / "check_results.csv",
                precheck_dir / "check_results.json",
            ],
        ),
        "precheck_clip_aggregates": read_first_records(
            "precheck_clip_aggregates",
            [
                precheck_dir / "clip_aggregates.parquet",
                precheck_dir / "clip_aggregates.csv",
                precheck_dir / "clip_aggregates.json",
            ],
        ),
        "precheck_candidate_windows": read_first_records(
            "precheck_candidate_windows",
            [
                precheck_dir / "candidate_windows.parquet",
                precheck_dir / "candidate_windows.csv",
                precheck_dir / "candidate_windows.json",
            ],
        ),
        "video_quality_summary": read_first_records(
            "video_quality_summary",
            [
                video_dir / "video_quality_decision_summary.csv",
                video_dir / "video_quality_acceptance_summary.csv",
            ],
        ),
        "video_quality_results": read_first_records(
            "video_quality_results",
            [video_dir / "video_quality_results.json"],
        ),
        "sam3_window_summary": read_first_records(
            "sam3_window_summary",
            [
                sam3_dir / "window_keypoint_containment_summary.parquet",
                sam3_dir / "window_keypoint_containment_summary.csv",
                sam3_dir / "window_keypoint_containment_summary.json",
            ],
        ),
        "ledger_asset_ledger": read_first_records(
            "ledger_asset_ledger",
            [ledger_dir / "xjgt_100_asset_ledger.csv"],
        ),
        "ledger_issue_events": read_first_records(
            "ledger_issue_events",
            [ledger_dir / "xjgt_100_issue_events.csv"],
        ),
        "manual_normalized_labels": read_first_records(
            "manual_normalized_labels",
            [manual_dir / "manual_labels_autosave.normalized.csv"],
        ),
        "manual_patch_labels": read_first_records(
            "manual_patch_labels",
            [manual_dir / "manual_labels_patch.json"],
        ),
    }


def build_episode_asset_map(manifest: list[dict[str, Any]]) -> dict[int, str]:
    mapping: dict[int, str] = {}
    for fallback_idx, row in enumerate(manifest):
        asset_id = normalize_asset_id(row.get("asset_id"))
        if not asset_id:
            continue
        episode_idx = integer_or_none(row.get("episode_idx"))
        mapping[episode_idx if episode_idx is not None else fallback_idx] = asset_id
    return mapping


def source_asset_id(
    row: dict[str, Any],
    episode_asset: dict[int, str],
) -> str:
    for key in ("asset_id", "clip_id"):
        asset_id = normalize_asset_id(row.get(key))
        if asset_id:
            return asset_id
    for key in ("video_path", "hdf5_path", "path"):
        asset_id = normalize_asset_id(row.get(key))
        if asset_id:
            return asset_id
    episode_idx = integer_or_none(row.get("episode_idx"))
    if episode_idx is not None:
        return episode_asset.get(episode_idx, f"episode_idx:{episode_idx}")
    return ""


def parse_metrics(row: dict[str, Any]) -> dict[str, Any]:
    raw = row.get("metrics")
    if isinstance(raw, dict):
        return raw
    if raw in (None, ""):
        return {}
    if isinstance(raw, str):
        try:
            decoded = json.loads(raw)
        except json.JSONDecodeError:
            return {}
        return decoded if isinstance(decoded, dict) else {}
    return {}


def truthy(value: Any) -> bool:
    if isinstance(value, bool):
        return value
    if value in (None, ""):
        return False
    if isinstance(value, (int, float)):
        return bool(value)
    return text(value).strip().lower() in {"true", "1", "yes", "fail", "failed"}


def boolish_or_none(value: Any) -> bool | None:
    if isinstance(value, bool):
        return value
    if value in (None, ""):
        return None
    if isinstance(value, (int, float)):
        return bool(value)
    normalized = text(value).strip().lower()
    if normalized in {"true", "1", "yes", "pass", "passed"}:
        return True
    if normalized in {"false", "0", "no", "fail", "failed"}:
        return False
    return None


def load_precheck_evidence(
    manifest: list[dict[str, Any]],
    sources: dict[str, SourceRead],
) -> PrecheckEvidence:
    episode_asset = build_episode_asset_map(manifest)
    check_rows = sources["precheck_check_results"].records
    candidate_rows = sources["precheck_candidate_windows"].records
    by_asset: dict[str, dict[str, Any]] = {}
    unmatched: set[str] = set()
    for row in manifest:
        asset_id = normalize_asset_id(row.get("asset_id"))
        if not asset_id:
            continue
        by_asset[asset_id] = {
            "checks": set(),
            "missing_intervals": [],
            "morphology_intervals": [],
            "text_status": "not_applicable",
            "morphology_status": "not_run",
            "missing_status": "not_run",
            "supplier_quality_signal": "not_provided",
            "temporal_status": "not_run",
            "temporal_seen": False,
            "quality_seen": False,
            "has_skeleton_quality": False,
            "has_morphology": False,
        }

    for row in check_rows:
        asset_id = source_asset_id(row, episode_asset)
        if not asset_id or asset_id not in by_asset:
            if asset_id:
                unmatched.add(asset_id)
            continue
        record = by_asset[asset_id]
        check = text(row.get("check"))
        metrics = parse_metrics(row)
        frame_idx = integer_or_none(row.get("frame_idx"))
        record["checks"].add(check)

        if check == "text_integrity":
            if truthy(row.get("flag")) or float_or_zero(metrics.get("missing_field_count")) > 0:
                record["text_status"] = "fail"
            elif record["text_status"] != "fail":
                record["text_status"] = "pass"
            continue

        if check == "quality_score":
            record["quality_seen"] = True
            if frame_idx == -1:
                record["supplier_quality_signal"] = supplier_quality_from_quality_score(row, metrics)
            elif float_or_zero(metrics.get("frame_score")) == 0.0:
                record["supplier_quality_signal"] = "low"
            continue

        if check == "keypoint_missing":
            # Despite its historical name, this check consumes supplier
            # quality_hand windows; it is not vendor-agnostic point existence.
            record["quality_seen"] = True
            if (
                float_or_zero(metrics.get("quality_low_left")) > 0
                or float_or_zero(metrics.get("quality_low_right")) > 0
            ):
                record["supplier_quality_signal"] = "low"
            elif record["supplier_quality_signal"] == "not_provided":
                record["supplier_quality_signal"] = "provided_ok"
            continue

        if check == "skeleton_quality_score":
            record["has_skeleton_quality"] = True
            record["temporal_seen"] = True
            if frame_idx is not None and frame_idx >= 0:
                if skeleton_presence_invalid(metrics):
                    record["missing_intervals"].append((frame_idx, frame_idx))
                verdict = text(metrics.get("skeleton_verdict")).lower()
                if verdict in {"review", "suspect"} or truthy(metrics.get("needs_visual_review")):
                    record["temporal_status"] = "review"
            continue

        if check == "keypoint_temporal":
            record["temporal_seen"] = True
            continue

        if check == "keypoint_morphology":
            record["has_morphology"] = True
            verdict = status_value(metrics.get("morphology_verdict"))
            if verdict in {"fail", "review", "pass", "not_applicable"}:
                record["morphology_status"] = worst_status(record["morphology_status"], verdict)
            if frame_idx is not None and frame_idx >= 0 and (
                verdict == "fail" or truthy(row.get("flag"))
            ):
                record["morphology_intervals"].append((frame_idx, frame_idx))

    for row in candidate_rows:
        asset_id = source_asset_id(row, episode_asset)
        if not asset_id or asset_id not in by_asset:
            if asset_id:
                unmatched.add(asset_id)
            continue
        by_asset[asset_id]["temporal_status"] = "review"
        by_asset[asset_id]["temporal_seen"] = True

    for record in by_asset.values():
        if record["has_skeleton_quality"]:
            record["missing_status"] = "fail" if record["missing_intervals"] else "pass"
        if record["has_morphology"] and record["morphology_status"] == "not_run":
            record["morphology_status"] = "pass"
        record["mapped_precheck_checks"] = "|".join(sorted(record["checks"]))
        missing_expected = []
        if not record["has_skeleton_quality"]:
            missing_expected.append("skeleton_quality_score")
        if not record["has_morphology"]:
            missing_expected.append("keypoint_morphology")
        record["missing_expected_checks"] = "|".join(missing_expected)
        if not record["checks"]:
            record["precheck_mapping_status"] = (
                "unmatched" if check_rows or candidate_rows else "source_missing"
            )
        elif missing_expected:
            record["precheck_mapping_status"] = "partial"
        else:
            record["precheck_mapping_status"] = "mapped"

    source_status = "readable" if check_rows or candidate_rows else sources["precheck_check_results"].status
    return PrecheckEvidence(
        by_asset=by_asset,
        source_status=source_status,
        unmatched_assets=sorted(unmatched),
    )


def skeleton_presence_invalid(metrics: dict[str, Any]) -> bool:
    saw_geometry_count = False
    geometry_invalid = False
    for side in ("left", "right"):
        valid = float_or_none(metrics.get(f"valid_keypoint_count_{side}"))
        missing = float_or_none(metrics.get(f"missing_keypoint_count_{side}"))
        if valid is not None:
            saw_geometry_count = True
            geometry_invalid = geometry_invalid or valid < 21
        if missing is not None:
            saw_geometry_count = True
            geometry_invalid = geometry_invalid or missing > 0
    if saw_geometry_count:
        return geometry_invalid
    return truthy(metrics.get("keypoint_presence_invalid")) and not truthy(
        metrics.get("low_quality_hand_invalid")
    )


def supplier_quality_from_quality_score(
    row: dict[str, Any],
    metrics: dict[str, Any],
) -> str:
    if truthy(row.get("flag")):
        return "provided_ok"
    pass_ratio = float_or_none(metrics.get("pass_ratio"))
    threshold = float_or_none(metrics.get("pass_threshold"))
    if pass_ratio is not None and threshold is not None:
        return "provided_ok" if pass_ratio >= threshold else "low"
    return "not_provided"


def worst_status(current: str, incoming: str) -> str:
    rank = {
        "fail": 5,
        "review": 4,
        "blocked": 3,
        "not_run": 2,
        "not_applicable": 1,
        "pass": 0,
    }
    return incoming if rank.get(incoming, 0) > rank.get(current, 0) else current


def load_xjgt(
    run_root: Path,
    *,
    require_xjgt_text: bool = False,
) -> dict[str, Any]:
    manifest_path = run_root / "manifests" / "supplier_manifest_xjgt_100.csv"
    manifest = read_csv(manifest_path)
    asset_rows, episode_asset = manifest_asset_maps(manifest)
    asset_ids = set(asset_rows)
    sources = load_xjgt_source_reads(run_root)
    precheck = load_precheck_evidence(manifest, sources)
    candidate_by_asset, candidate_unmatched = map_window_rows(
        sources["precheck_candidate_windows"].records,
        asset_ids,
        episode_asset,
    )
    video_by_asset, video_unmatched = map_video_quality_evidence(
        [
            *sources["video_quality_results"].records,
            *sources["video_quality_summary"].records,
        ],
        asset_ids,
        episode_asset,
    )
    sam3_by_asset, sam3_unmatched = map_sam3_evidence(
        sources["sam3_window_summary"].records,
        asset_ids,
        episode_asset,
    )
    manual_labels = dedupe_manual_labels(
        [
            *sources["manual_normalized_labels"].records,
            *sources["manual_patch_labels"].records,
        ]
    )
    ledger_by_asset = {
        normalize_asset_id(row.get("asset_id")): row
        for row in sources["ledger_asset_ledger"].records
    }
    frame_info: dict[str, dict[str, Any]] = {}
    for asset_id, manifest_row in asset_rows.items():
        video_path = resolve_xjgt_video_path(manifest_row, asset_id)
        probed_frames = probe_video_frame_count(video_path)
        fallback_frames = integer_or_none(
            first_value(
                manifest_row,
                ("frame_count", "total_frames", "num_frames"),
            )
        ) or 0
        frame_info[asset_id] = {
            "total_frames": probed_frames or fallback_frames,
            "frame_count_status": (
                "ok"
                if probed_frames > 0
                else "manifest_fallback"
                if fallback_frames > 0
                else "unreadable"
            ),
            "video_path": video_path,
        }
    manual_by_asset = aggregate_manual_review(
        manual_labels,
        {
            asset_id: int(info["total_frames"])
            for asset_id, info in frame_info.items()
        },
    )
    evidence_paths = [
        source.path
        for source in sources.values()
        if source.path is not None
    ]
    details = [
        build_xjgt_detail(
            asset_id=asset_id,
            total_frames=int(frame_info[asset_id]["total_frames"]),
            frame_count_status=str(frame_info[asset_id]["frame_count_status"]),
            video_path=Path(frame_info[asset_id]["video_path"]),
            precheck_row=precheck.by_asset.get(
                asset_id, empty_precheck_evidence()
            ),
            precheck_source_status=precheck.source_status,
            precheck_unmatched=precheck.unmatched_assets,
            candidate_row=candidate_by_asset.get(asset_id, {}),
            video_row=video_by_asset.get(asset_id, {}),
            sam3_row=sam3_by_asset.get(asset_id, {}),
            manual_row=manual_by_asset.get(
                asset_id,
                empty_manual_evidence(
                    int(frame_info[asset_id]["total_frames"])
                ),
            ),
            ledger_row=ledger_by_asset.get(asset_id, {}),
            require_xjgt_text=require_xjgt_text,
            evidence_paths=evidence_paths,
        )
        for asset_id in asset_rows
    ]
    issue_events = build_weekly_issue_events(
        manual_labels,
        sources["sam3_window_summary"].records,
        details,
    )
    modules_completed = []
    if details and all(
        row["precheck_mapping_status"] == "mapped" for row in details
    ):
        modules_completed.append("precheck")
    if details and all(
        row["video_quality_status"] != "no_valid_output"
        for row in details
    ):
        modules_completed.append("video_quality")
    if sources["sam3_window_summary"].status == "readable":
        modules_completed.append("sam3")
    if sources["manual_normalized_labels"].status == "readable":
        modules_completed.append("manual_review")
    blocked_modules = []
    if "precheck" not in modules_completed:
        blocked_modules.append("precheck_mapping_incomplete")
    if "video_quality" not in modules_completed:
        blocked_modules.append("video_quality_mapping_incomplete")
    summary = supplier_summary_from_details(
        supplier_name="星际归途 / XJGT",
        expected_clip_count=100,
        details=details,
        main_issue_type=most_common(
            Counter(
                text(event.get("failure_mode"))
                for event in issue_events
                if event.get("failure_mode")
            )
        ),
        modules_completed="|".join(modules_completed),
        blocked_modules="|".join(blocked_modules),
        notes=(
            "High-risk-biased manual review; not an unbiased global "
            "error-rate estimate. "
            f"unmatched_precheck={len(precheck.unmatched_assets)}; "
            f"unmatched_candidate={len(candidate_unmatched)}; "
            f"unmatched_video={len(video_unmatched)}; "
            f"unmatched_sam3={len(sam3_unmatched)}"
        ),
    )
    return {
        "summary": summary,
        "details": details,
        "issue_events": issue_events,
    }


def load_deepreach(run_root: Path) -> dict[str, Any]:
    manifest_path = (
        run_root
        / "deepreach"
        / "manifests"
        / "supplier_manifest_deepreach.csv"
    )
    manifest = read_csv(manifest_path)
    video_results = (
        run_root
        / "deepreach"
        / "video_quality"
        / "video_quality_results.json"
    )
    video_summary = (
        run_root
        / "deepreach"
        / "video_quality"
        / "video_quality_decision_summary.csv"
    )
    has_video_quality = video_results.exists() or video_summary.exists()
    details = []
    for row in manifest:
        asset_id = normalize_asset_id(row.get("asset_id"))
        video_path = Path(text(row.get("video_path")))
        probed_frames = probe_video_frame_count(video_path) if video_path else 0
        fallback_frames = integer_or_none(
            first_value(row, ("frame_count", "total_frames", "num_frames"))
        ) or 0
        total_frames = probed_frames or fallback_frames
        details.append(
            {
                "asset_id": asset_id,
                "total_frames": total_frames,
                "frame_count_status": (
                    "ok"
                    if probed_frames > 0
                    else "manifest_fallback"
                    if fallback_frames > 0
                    else "not_run"
                ),
                "text_check_status": "blocked",
                "skeleton_missing_status": "blocked",
                "skeleton_missing_fail_frame_count": 0,
                "skeleton_missing_fail_frame_ratio": 0.0,
                "skeleton_morphology_status": "blocked",
                "skeleton_morphology_fail_frame_count": 0,
                "skeleton_morphology_fail_frame_ratio": 0.0,
                "skeleton_static_fail_frame_count": 0,
                "skeleton_static_fail_frame_ratio": 0.0,
                "skeleton_static_status": "blocked",
                "supplier_quality_signal": "not_provided",
                "video_quality_status": (
                    "no_valid_output" if not has_video_quality else "review"
                ),
                "video_quality_fail_frame_count": 0,
                "video_quality_fail_frame_ratio": 0.0,
                "temporal_status": "blocked",
                "sam3_containment_status": "blocked",
                "manual_review_status": "not_run",
                "temporal_sam3_manual_status": "blocked",
                "manual_problem_frame_count": 0,
                "manual_problem_frame_ratio_of_clip": 0.0,
                "manual_reviewed_frame_count": 0,
                "manual_problem_ratio_of_reviewed": 0.0,
                "abnormal_fail_frame_count": 0,
                "abnormal_fail_frame_ratio": 0.0,
                "abnormal_frame_status": "blocked",
                "auto_fail_frame_count": 0,
                "reviewed_auto_fail_frame_count": 0,
                "reviewed_auto_fail_true_positive_frame_count": 0,
                "reviewed_auto_fail_false_positive_frame_count": 0,
                "unreviewed_auto_fail_frame_count": 0,
                "auto_fail_precision_on_reviewed": 0.0,
                "abnormal_status_reason": "blocked:sam3_projection_mapping",
                "fail_indicator_count": 0,
                "final_clip_status": "blocked",
                "mapped_precheck_checks": "",
                "missing_expected_checks": "skeleton_quality_score|keypoint_morphology",
                "precheck_mapping_status": "blocked",
                "final_status_reason": "blocked:missing_deepreach_hdf5_adapter",
                "main_issue_type": "missing_schema_adapter",
                "notes": (
                    "precheck blocked by HDF5 schema adapter; "
                    "SAM3 blocked by calibration/projection mapping"
                ),
                "evidence_path": str(manifest_path),
            }
        )
    modules_completed = ["manifest"] if manifest else []
    if has_video_quality:
        modules_completed.append("video_quality")
    counts = Counter(row["final_clip_status"] for row in details)
    covered_count = counts["pass"] + counts["fail"] + counts["review"]
    return {
        "summary": {
            "supplier_name": "DeepReach",
            "sample_clip_count": len(manifest),
            "expected_clip_count": 100,
            "total_frame_count": sum(
                int(row["total_frames"]) for row in details
            ),
            "problem_frame_count": 0,
            "problem_frame_ratio": 0.0,
            "pass_clip_count": 0,
            "fail_clip_count": 0,
            "review_clip_count": 0,
            "blocked_clip_count": counts["blocked"],
            "not_run_clip_count": counts["not_run"],
            "coverage_ratio": safe_ratio(covered_count, len(manifest)),
            "pass_clip_ratio": 0.0,
            "fail_clip_ratio": 0.0,
            "main_issue_type": "missing_schema_adapter",
            "video_quality_status": (
                "review" if has_video_quality else "no_valid_output"
            ),
            "modules_completed": "|".join(modules_completed),
            "blocked_modules": "precheck|sam3",
            "notes": (
                "video_quality not_run/no_valid_output unless result/decision summary exists"
            ),
        },
        "details": details,
    }


def placeholder_supplier(index: int) -> dict[str, Any]:
    return {
        "supplier_name": f"供应商{index}",
        "sample_clip_count": 0,
        "expected_clip_count": 100,
        "total_frame_count": 0,
        "problem_frame_count": 0,
        "problem_frame_ratio": 0.0,
        "pass_clip_count": 0,
        "fail_clip_count": 0,
        "review_clip_count": 0,
        "blocked_clip_count": 0,
        "not_run_clip_count": 0,
        "coverage_ratio": 0.0,
        "pass_clip_ratio": 0.0,
        "fail_clip_ratio": 0.0,
        "main_issue_type": "",
        "modules_completed": "",
        "blocked_modules": "missing_input",
        "notes": "input manifest not provided",
    }


def hard_issue_rows(
    events: list[dict[str, Any]],
    details: list[dict[str, Any]] | None = None,
) -> list[dict[str, Any]]:
    details = details or []
    video_fail_count = sum(
        1 for row in details if row.get("video_quality_status") == "fail"
    )
    definitions = [
        (
            "hand_out_of_frame",
            "projection/border evidence + manual review",
            "pending calibration",
            "partial",
            "yes",
            "Finite keypoints do not prove visual presence.",
            "Calibrate border and mask-truncation rules.",
        ),
        (
            "severe_keypoint_offset",
            "SAM3 containment + manual overlay",
            "inside_ratio <= 0.2 is strong mismatch evidence",
            "partial",
            "yes",
            "Occlusion and undersegmentation can look identical.",
            "Calibrate with confirmed segments.",
        ),
        (
            "skeleton_pose_hallucination",
            "static morphology + manual overlay",
            "config-driven morphology thresholds",
            "partial",
            "yes",
            "A wrong pose may remain temporally stable.",
            "Expand morphology calibration.",
        ),
        (
            "visual_skeleton_presence_mismatch",
            "projection/SAM3 + manual review",
            "pending calibration",
            "partial",
            "yes",
            "Visibility depends on framing and occlusion.",
            "Add presence/truncation evidence.",
        ),
        (
            "occlusion_or_mask_undersegmentation",
            "SAM3 diagnostics",
            "inside_ratio < 0.6 routes to review",
            "no",
            "yes",
            "Keypoint-in-mask cannot diagnose mask completeness.",
            "Retain manual arbitration.",
        ),
        (
            "projection_review",
            "3D-to-2D projection checks",
            "projected_in_image_ratio < 0.8",
            "no",
            "yes",
            "Calibration uncertainty can dominate.",
            "Validate supplier calibration.",
        ),
        (
            "video_quality_pending_colleague_thresholds",
            "acceptance_pull/video_quality.py",
            "pending colleague thresholds",
            "pending",
            "no",
            "Threshold ownership is external to this report.",
            "Consume finalized video_quality status.",
        ),
        (
            "text_check_pending_rules",
            "precheck text integrity",
            "pending rules",
            "no",
            "yes",
            "Acceptance semantics are not finalized.",
            "Define required fields and fail rules.",
        ),
    ]
    rows = []
    for issue_type, method, threshold, automatic, manual, why, action in definitions:
        source_module = (
            "sam3_containment"
            if issue_type
            in {"occlusion_or_mask_undersegmentation", "projection_review"}
            else "manual_review"
        )
        source_events = [
            event
            for event in events
            if event.get("source_module") == source_module
        ]
        count = sum(
            event.get("failure_mode") == issue_type
            for event in source_events
        )
        unit = "window" if source_module == "sam3_containment" else "segment"
        denominator = len(source_events)
        denominator_definition = (
            "SAM3 containment windows"
            if source_module == "sam3_containment"
            else "manual review segment rows"
        )
        if issue_type == "video_quality_pending_colleague_thresholds":
            count = video_fail_count
            unit = "clip"
            denominator = len(details)
            denominator_definition = "XJGT sampled clips"
            source_module = "video_quality"
        elif issue_type == "text_check_pending_rules":
            count = sum(
                row.get("text_check_status")
                in {"pending_rule", "pending_required_rule"}
                for row in details
            )
            unit = "clip"
            denominator = len(details)
            denominator_definition = "XJGT sampled clips"
            source_module = "text_integrity"
        rows.append({
            "issue_type": issue_type,
            "current_detection_method": method,
            "current_threshold": threshold,
            "auto_decision_available": automatic,
            "needs_manual_review": manual,
            "observed_count": count,
            "observed_unit": unit,
            "denominator": denominator,
            "denominator_definition": denominator_definition,
            "observed_ratio": safe_ratio(count, denominator),
            "source_module": source_module,
            "why_hard": why,
            "next_action": action,
        })
    return rows


def write_workbook(
    path: Path,
    summary_rows: list[dict[str, Any]],
    xjgt_details: list[dict[str, Any]],
    deepreach_details: list[dict[str, Any]],
    events: list[dict[str, Any]],
) -> None:
    workbook = Workbook()
    workbook.remove(workbook.active)
    add_sheet(workbook, "五供应商总览", SUMMARY_COLUMNS, summary_rows)
    add_sheet(
        workbook,
        "星际归途",
        XJGT_DETAIL_COLUMNS,
        xjgt_details,
    )
    add_sheet(
        workbook,
        "DeepReach",
        XJGT_DETAIL_COLUMNS,
        deepreach_details,
    )
    for index in range(3, 6):
        add_sheet(
            workbook,
            f"供应商{index}",
            XJGT_DETAIL_COLUMNS,
            [],
        )
    add_sheet(
        workbook,
        "人工与难测问题统计",
        HARD_ISSUE_COLUMNS,
        hard_issue_rows(events, xjgt_details),
    )
    path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(path)


def add_sheet(
    workbook: Workbook,
    name: str,
    columns: list[str],
    rows: list[dict[str, Any]],
) -> None:
    sheet = workbook.create_sheet(name)
    sheet.append(columns)
    for row in rows:
        sheet.append([row.get(column, "") for column in columns])
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    sheet.sheet_view.showGridLines = False
    for cell in sheet[1]:
        cell.fill = PatternFill("solid", fgColor="D9EAF7")
        cell.font = Font(bold=True)
        cell.alignment = Alignment(vertical="center", wrap_text=True)
    for column_index, column in enumerate(columns, 1):
        values = [len(column)]
        values.extend(
            len(str(sheet.cell(row_index, column_index).value or ""))
            for row_index in range(2, min(sheet.max_row, 101) + 1)
        )
        sheet.column_dimensions[get_column_letter(column_index)].width = min(
            max(max(values) + 2, 12), 42
        )
    percentage_columns = {
        "problem_frame_ratio",
        "pass_clip_ratio",
        "fail_clip_ratio",
        "coverage_ratio",
        "skeleton_missing_fail_frame_ratio",
        "skeleton_morphology_fail_frame_ratio",
        "skeleton_static_fail_frame_ratio",
        "video_quality_fail_frame_ratio",
        "manual_problem_frame_ratio_of_clip",
        "manual_problem_ratio_of_reviewed",
        "abnormal_fail_frame_ratio",
        "auto_fail_precision_on_reviewed",
        "observed_ratio",
    }
    for column in percentage_columns.intersection(columns):
        column_letter = get_column_letter(columns.index(column) + 1)
        for cell in sheet[column_letter][1:]:
            if isinstance(cell.value, (int, float)):
                cell.number_format = "0.0%"


def read_first_records(input_name: str, paths: Sequence[Path]) -> SourceRead:
    for path in paths:
        if not path.exists():
            continue
        try:
            records = read_table_records(path)
        except Exception as exc:  # pragma: no cover - diagnostic path
            LOGGER.warning("Could not read %s: %s", path, exc)
            return SourceRead(input_name, path, [], f"unreadable:{type(exc).__name__}")
        return SourceRead(input_name, path, records, "readable")
    return SourceRead(input_name, paths[0] if paths else None, [], "missing")


def read_table_records(path: Path) -> list[dict[str, Any]]:
    suffix = path.suffix.lower()
    if suffix == ".json":
        data = json.loads(path.read_text(encoding="utf-8"))
        if isinstance(data, list):
            return [row for row in data if isinstance(row, dict)]
        if isinstance(data, dict):
            for key in ("records", "rows", "items", "results", "segments"):
                value = data.get(key)
                if isinstance(value, list):
                    return [row for row in value if isinstance(row, dict)]
            if data and all(isinstance(value, dict) for value in data.values()):
                return [
                    {"asset_id": key, **dict(value)}
                    for key, value in data.items()
                ]
            return [data]
        return []
    if suffix == ".parquet":
        import pandas as pd

        frame = pd.read_parquet(path)
        frame = frame.where(frame.notna(), "")
        return frame.to_dict(orient="records")
    return read_csv(path)


def manifest_asset_maps(
    manifest: list[dict[str, Any]],
) -> tuple[dict[str, dict[str, Any]], dict[int, str]]:
    rows: dict[str, dict[str, Any]] = {}
    episode_asset: dict[int, str] = {}
    for fallback_idx, row in enumerate(manifest):
        asset_id = normalize_asset_id(row.get("asset_id"))
        if not asset_id:
            continue
        rows[asset_id] = row
        episode_idx = integer_or_none(row.get("episode_idx"))
        episode_asset[episode_idx if episode_idx is not None else fallback_idx] = asset_id
    return rows, episode_asset


def empty_precheck_evidence() -> dict[str, Any]:
    return {
        "checks": set(),
        "missing_intervals": [],
        "morphology_intervals": [],
        "text_status": "not_applicable",
        "supplier_quality_signal": "not_provided",
        "temporal_status": "not_run",
        "has_skeleton_quality": False,
        "has_morphology": False,
    }


def map_window_rows(
    rows: list[dict[str, Any]],
    asset_ids: set[str],
    episode_asset: dict[int, str],
) -> tuple[dict[str, dict[str, Any]], list[str]]:
    mapped: dict[str, dict[str, Any]] = defaultdict(
        lambda: {"intervals": [], "rows": []}
    )
    unmatched: set[str] = set()
    for row in rows:
        asset_id = source_asset_id(row, episode_asset)
        if not asset_id or asset_id not in asset_ids:
            if asset_id:
                unmatched.add(asset_id)
            continue
        interval = event_frame_interval(row)
        if interval is not None:
            mapped[asset_id]["intervals"].append(interval)
        mapped[asset_id]["rows"].append(row)
    return dict(mapped), sorted(unmatched)


def map_video_quality_evidence(
    rows: list[dict[str, Any]],
    asset_ids: set[str],
    episode_asset: dict[int, str],
) -> tuple[dict[str, dict[str, Any]], list[str]]:
    mapped: dict[str, dict[str, Any]] = {}
    unmatched: set[str] = set()
    for raw_row in rows:
        row = normalize_video_quality_row(raw_row)
        asset_id = source_asset_id(row, episode_asset)
        if not asset_id or asset_id not in asset_ids:
            if asset_id:
                unmatched.add(asset_id)
            continue
        raw_status = text(
            first_value(
                row,
                (
                    "status",
                    "acceptance_status",
                    "video_quality_status",
                    "decision",
                    "final_status",
                ),
            )
        ).lower()
        passed = boolish_or_none(
            first_value(row, ("passed", "pass", "video_quality_pass"))
        )
        if passed is False or raw_status in {"fail", "failed"}:
            status = "fail"
        elif raw_status in {"review", "not_run", "no_valid_output"}:
            status = raw_status
        else:
            status = "pass"
        mapped[asset_id] = {
            "status": status,
            "raw_status": raw_status or "pass",
            "fail_intervals": (
                extract_video_fail_intervals(raw_row)
                if status == "fail"
                else []
            ),
        }
    return mapped, sorted(unmatched)


def extract_video_fail_intervals(row: dict[str, Any]) -> list[tuple[int, int]]:
    intervals: list[tuple[int, int]] = []
    direct = event_frame_interval(row)
    if direct is not None:
        intervals.append(direct)
    video_quality = row.get("video_quality")
    if not isinstance(video_quality, dict):
        return intervals
    metrics = video_quality.get("metrics")
    if not isinstance(metrics, dict):
        return intervals
    freeze = metrics.get("freeze_metrics")
    if not isinstance(freeze, dict):
        return intervals
    for item in freeze.get("frozen_intervals") or []:
        if not isinstance(item, dict):
            continue
        interval = event_frame_interval(item)
        if interval is not None:
            intervals.append(interval)
    return intervals


def map_sam3_evidence(
    rows: list[dict[str, Any]],
    asset_ids: set[str],
    episode_asset: dict[int, str],
) -> tuple[dict[str, dict[str, Any]], list[str]]:
    mapped: dict[str, dict[str, Any]] = defaultdict(
        lambda: {
            "status": "pass",
            "auto_fail_intervals": [],
            "review_intervals": [],
        }
    )
    unmatched: set[str] = set()
    for row in rows:
        asset_id = source_asset_id(row, episode_asset)
        if not asset_id or asset_id not in asset_ids:
            if asset_id:
                unmatched.add(asset_id)
            continue
        verdict = text(row.get("window_containment_verdict")).lower()
        status = (
            "fail"
            if verdict == "containment_fail"
            else "review"
            if verdict
            in {
                "mixed_review",
                "projection_review",
                "review",
                "side_view_manual_review",
                "rotation_manual_review",
            }
            else "pass"
        )
        record = mapped[asset_id]
        record["status"] = worst_status(record["status"], status)
        interval = event_frame_interval(row)
        if interval is not None and status != "pass":
            record["review_intervals"].append(interval)
        if interval is not None and verdict == "containment_fail":
            record["auto_fail_intervals"].append(interval)
    return dict(mapped), sorted(unmatched)


def dedupe_manual_labels(
    labels: list[dict[str, Any]],
) -> list[dict[str, Any]]:
    result: list[dict[str, Any]] = []
    seen: set[tuple[Any, ...]] = set()
    for row in labels:
        key = (
            normalize_asset_id(row.get("asset_id")),
            text(row.get("review_id")),
            text(row.get("segment_id")),
            text(row.get("manual_outcome")),
            integer_or_none(row.get("window_start_frame")),
            integer_or_none(row.get("window_end_frame")),
            integer_or_none(row.get("affected_start_frame")),
            integer_or_none(row.get("affected_end_frame")),
        )
        if key in seen:
            continue
        seen.add(key)
        result.append(row)
    return result


def aggregate_manual_review(
    labels: list[dict[str, Any]],
    total_frames_by_asset: dict[str, int],
) -> dict[str, dict[str, Any]]:
    problems: dict[str, list[tuple[int, int]]] = defaultdict(list)
    reviewed: dict[str, list[tuple[int, int]]] = defaultdict(list)
    labeled_assets: set[str] = set()
    for row in labels:
        asset_id = normalize_asset_id(row.get("asset_id"))
        if not asset_id or asset_id not in total_frames_by_asset:
            continue
        labeled_assets.add(asset_id)
        start = integer_or_none(row.get("window_start_frame"))
        end = integer_or_none(row.get("window_end_frame"))
        if start is None or end is None:
            start = integer_or_none(row.get("affected_start_frame"))
            end = integer_or_none(row.get("affected_end_frame"))
        if start is not None and end is not None:
            reviewed[asset_id].append((start, end))
        if text(row.get("manual_outcome")).lower() != "true_positive":
            continue
        affected_start = integer_or_none(row.get("affected_start_frame"))
        affected_end = integer_or_none(row.get("affected_end_frame"))
        if affected_start is not None and affected_end is not None:
            problems[asset_id].append((affected_start, affected_end))

    output: dict[str, dict[str, Any]] = {}
    for asset_id in labeled_assets:
        total_frames = total_frames_by_asset[asset_id]
        problem_intervals = clip_intervals(problems[asset_id], total_frames)
        reviewed_intervals = clip_intervals(reviewed[asset_id], total_frames)
        problem_count = interval_frame_count(problem_intervals)
        reviewed_count = interval_frame_count(reviewed_intervals)
        problem_ratio = safe_ratio(problem_count, total_frames)
        output[asset_id] = {
            "manual_review_status": (
                "fail" if problem_ratio >= 0.10 else "pass"
            ),
            "manual_problem_frame_count": problem_count,
            "manual_problem_frame_ratio_of_clip": problem_ratio,
            "manual_reviewed_frame_count": reviewed_count,
            "manual_problem_ratio_of_reviewed": safe_ratio(
                problem_count, reviewed_count
            ),
            "problem_intervals": problem_intervals,
            "reviewed_intervals": reviewed_intervals,
        }
    return output


def empty_manual_evidence(total_frames: int) -> dict[str, Any]:
    return {
        "manual_review_status": "not_reviewed",
        "manual_problem_frame_count": 0,
        "manual_problem_frame_ratio_of_clip": safe_ratio(0, total_frames),
        "manual_reviewed_frame_count": 0,
        "manual_problem_ratio_of_reviewed": 0.0,
        "problem_intervals": [],
        "reviewed_intervals": [],
    }


def build_xjgt_detail(
    *,
    asset_id: str,
    total_frames: int,
    frame_count_status: str,
    video_path: Path,
    precheck_row: dict[str, Any],
    precheck_source_status: str,
    precheck_unmatched: list[str],
    candidate_row: dict[str, Any],
    video_row: dict[str, Any],
    sam3_row: dict[str, Any],
    manual_row: dict[str, Any],
    ledger_row: dict[str, Any],
    require_xjgt_text: bool,
    evidence_paths: list[Path],
) -> dict[str, Any]:
    missing_intervals = clip_intervals(
        precheck_row.get("missing_intervals", []), total_frames
    )
    morphology_intervals = clip_intervals(
        precheck_row.get("morphology_intervals", []), total_frames
    )
    missing_count = interval_frame_count(missing_intervals)
    morphology_count = interval_frame_count(morphology_intervals)
    missing_ratio = safe_ratio(missing_count, total_frames)
    morphology_ratio = safe_ratio(morphology_count, total_frames)
    missing_seen = bool(
        precheck_row.get("has_skeleton_quality")
        or precheck_row.get("presence_seen")
    )
    morphology_seen = bool(
        precheck_row.get("has_morphology")
        or precheck_row.get("morphology_seen")
    )
    missing_status = (
        frame_ratio_status("pass", missing_count, missing_ratio)
        if missing_seen
        else "not_run"
    )
    if not morphology_seen:
        morphology_status = "not_run"
    elif morphology_count:
        morphology_status = frame_ratio_status(
            "pass", morphology_count, morphology_ratio
        )
    else:
        morphology_status = (
            "review"
            if text(precheck_row.get("morphology_status")).lower() == "review"
            else "pass"
        )
    static_intervals = [*missing_intervals, *morphology_intervals]
    static_count = interval_frame_count(static_intervals)
    static_ratio = safe_ratio(static_count, total_frames)
    static_status = skeleton_static_status_from_parts(
        missing_status=missing_status,
        morphology_status=morphology_status,
        fail_frame_ratio=static_ratio,
    )

    mapped_checks = sorted(precheck_row.get("checks", set()))
    missing_expected = [
        check
        for check, present in (
            ("skeleton_quality_score", missing_seen),
            ("keypoint_morphology", morphology_seen),
        )
        if not present
    ]
    mapping_status = source_mapping_status(
        precheck_source_status,
        mapped_checks,
        missing_expected,
        bool(precheck_unmatched),
    )
    if "text_integrity" in mapped_checks:
        observed_text_status = text(precheck_row.get("text_status")) or "pass"
        text_status = observed_text_status if require_xjgt_text else "pending_rule"
    else:
        observed_text_status = "not_run"
        text_status = (
            "pending_required_rule"
            if require_xjgt_text
            else "not_applicable"
        )
    quality_signal = text(
        precheck_row.get("supplier_quality_signal")
    ) or "not_provided"
    candidate_intervals = clip_intervals(
        candidate_row.get("intervals", []), total_frames
    )
    temporal_seen = bool(
        precheck_row.get("has_skeleton_quality")
        or precheck_row.get("temporal_seen")
    )
    temporal_status = (
        "not_run"
        if not temporal_seen
        else "review"
        if text(precheck_row.get("temporal_status")) == "review"
        or candidate_intervals
        else "pass"
    )
    sam3_status = (
        text(sam3_row.get("status"))
        if sam3_row
        else "not_run"
        if candidate_intervals
        else "not_applicable"
    )
    auto_fail_intervals = clip_intervals(
        sam3_row.get("auto_fail_intervals", []), total_frames
    )
    review_intervals = clip_intervals(
        [
            *candidate_intervals,
            *sam3_row.get("review_intervals", []),
        ],
        total_frames,
    )
    reviewed_intervals = manual_row["reviewed_intervals"]
    manual_intervals = manual_row["problem_intervals"]
    reviewed_auto = intersect_intervals(
        auto_fail_intervals, reviewed_intervals
    )
    reviewed_auto_true_positive = intersect_intervals(
        reviewed_auto, manual_intervals
    )
    reviewed_auto_false_positive = subtract_intervals(
        reviewed_auto, manual_intervals
    )
    unreviewed_auto = subtract_intervals(
        auto_fail_intervals, reviewed_intervals
    )
    abnormal_intervals = [*unreviewed_auto, *manual_intervals]
    unresolved_review = subtract_intervals(
        review_intervals, reviewed_intervals
    )
    auto_count = interval_frame_count(auto_fail_intervals)
    reviewed_auto_count = interval_frame_count(reviewed_auto)
    reviewed_auto_true_positive_count = interval_frame_count(
        reviewed_auto_true_positive
    )
    reviewed_auto_false_positive_count = interval_frame_count(
        reviewed_auto_false_positive
    )
    unreviewed_auto_count = interval_frame_count(unreviewed_auto)
    unresolved_review_count = interval_frame_count(unresolved_review)
    abnormal_count = interval_frame_count(abnormal_intervals)
    abnormal_ratio = safe_ratio(abnormal_count, total_frames)
    if abnormal_ratio >= 0.10:
        abnormal_frame_status = "fail"
    elif (
        unresolved_review_count
        or temporal_status == "not_run"
        or (
            manual_row["manual_review_status"] == "not_reviewed"
            and (
                temporal_status == "review"
                or sam3_status == "review"
            )
        )
    ):
        abnormal_frame_status = "review"
    else:
        abnormal_frame_status = "pass"
    abnormal_reason = (
        f"abnormal_ratio={abnormal_ratio:.6f}; "
        f"manual_true_positive_frames={manual_row['manual_problem_frame_count']}; "
        f"auto_fail_frames={auto_count}; "
        f"reviewed_auto_fail_frames={reviewed_auto_count}; "
        f"unreviewed_auto_fail_frames={unreviewed_auto_count}; "
        f"unresolved_review_frames={unresolved_review_count}"
    )

    video_status = text(video_row.get("status")) or "no_valid_output"
    video_intervals = clip_intervals(
        video_row.get("fail_intervals", []), total_frames
    )
    video_count = interval_frame_count(video_intervals)
    video_ratio = safe_ratio(video_count, total_frames)
    unresolved = []
    if video_status in {"no_valid_output", "not_run", "blocked"}:
        unresolved.append(f"video_quality_status={video_status}")
    if static_status in {"not_run", "blocked", "review"}:
        unresolved.append(f"skeleton_static_status={static_status}")
    if abnormal_frame_status == "review":
        unresolved.append("abnormal_frame_status=review")
    if text_status == "pending_required_rule":
        unresolved.append("text_check_status=pending_required_rule")
    if frame_count_status == "unreadable":
        unresolved.append("frame_count_status=unreadable")
    final_status, final_reason = final_status_with_reason(
        text_check_status=text_status,
        video_quality_status=video_status,
        skeleton_static_status=static_status,
        abnormal_frame_status=abnormal_frame_status,
        unresolved=unresolved,
    )
    notes = [
        text(ledger_row.get("notes")),
        f"supplier_quality_signal={quality_signal}",
    ]
    if not require_xjgt_text and observed_text_status != "not_run":
        notes.append(
            "text_integrity_observed="
            f"{observed_text_status}; text_schema_required=false"
        )
    if frame_count_status == "unreadable":
        notes.append(f"frame_count_unreadable:{video_path}")
    if mapping_status != "mapped":
        notes.append(
            f"precheck_source={precheck_source_status}; "
            f"unmatched={','.join(precheck_unmatched[:3])}"
        )
    return {
        "asset_id": asset_id,
        "total_frames": total_frames,
        "frame_count_status": frame_count_status,
        "text_check_status": text_status,
        "skeleton_missing_status": missing_status,
        "skeleton_missing_fail_intervals": format_intervals(missing_intervals),
        "skeleton_missing_fail_frame_count": missing_count,
        "skeleton_missing_fail_frame_ratio": missing_ratio,
        "skeleton_morphology_status": morphology_status,
        "skeleton_morphology_fail_intervals": format_intervals(
            morphology_intervals
        ),
        "skeleton_morphology_fail_frame_count": morphology_count,
        "skeleton_morphology_fail_frame_ratio": morphology_ratio,
        "skeleton_static_fail_intervals": format_intervals(
            merge_intervals(static_intervals)
        ),
        "skeleton_static_fail_frame_count": static_count,
        "skeleton_static_fail_frame_ratio": static_ratio,
        "skeleton_static_status": static_status,
        "supplier_quality_signal": quality_signal,
        "video_quality_status": video_status,
        "video_quality_fail_frame_count": video_count,
        "video_quality_fail_frame_ratio": video_ratio,
        "temporal_status": temporal_status,
        "sam3_containment_status": sam3_status,
        "manual_review_status": manual_row["manual_review_status"],
        "temporal_sam3_manual_status": abnormal_frame_status,
        "manual_problem_frame_count": manual_row[
            "manual_problem_frame_count"
        ],
        "manual_problem_frame_ratio_of_clip": manual_row[
            "manual_problem_frame_ratio_of_clip"
        ],
        "manual_reviewed_frame_count": manual_row[
            "manual_reviewed_frame_count"
        ],
        "manual_problem_ratio_of_reviewed": manual_row[
            "manual_problem_ratio_of_reviewed"
        ],
        "abnormal_fail_frame_count": abnormal_count,
        "abnormal_fail_frame_ratio": abnormal_ratio,
        "abnormal_frame_status": abnormal_frame_status,
        "auto_fail_frame_count": auto_count,
        "reviewed_auto_fail_frame_count": reviewed_auto_count,
        "reviewed_auto_fail_true_positive_frame_count": (
            reviewed_auto_true_positive_count
        ),
        "reviewed_auto_fail_false_positive_frame_count": (
            reviewed_auto_false_positive_count
        ),
        "unreviewed_auto_fail_frame_count": unreviewed_auto_count,
        "auto_fail_precision_on_reviewed": (
            safe_ratio(
                reviewed_auto_true_positive_count,
                reviewed_auto_count,
            )
            if reviewed_auto_count
            else ""
        ),
        "abnormal_status_reason": abnormal_reason,
        "fail_indicator_count": count_fail_indicators(
            text_check_status=text_status,
            video_quality_status=video_status,
            skeleton_static_status=static_status,
            abnormal_frame_status=abnormal_frame_status,
        ),
        "final_clip_status": final_status,
        "mapped_precheck_checks": "|".join(mapped_checks),
        "missing_expected_checks": "|".join(missing_expected),
        "precheck_mapping_status": mapping_status,
        "final_status_reason": final_reason,
        "main_issue_type": first_issue(ledger_row.get("top_issue_types")),
        "notes": "; ".join(note for note in notes if note),
        "evidence_path": "|".join(str(path) for path in evidence_paths),
        "_problem_intervals": [
            *static_intervals,
            *abnormal_intervals,
            *video_intervals,
        ],
    }


def source_mapping_status(
    source_status: str,
    mapped_checks: list[str],
    missing_expected: list[str],
    has_unmatched_source_rows: bool,
) -> str:
    if source_status == "missing":
        return "source_missing"
    if source_status.startswith("unreadable"):
        return "source_unreadable"
    if not mapped_checks:
        return "unmatched" if has_unmatched_source_rows else "check_absent"
    if missing_expected:
        return "partial"
    return "mapped"


def final_status_with_reason(
    *,
    text_check_status: str,
    video_quality_status: str,
    skeleton_static_status: str,
    abnormal_frame_status: str,
    unresolved: list[str],
) -> tuple[str, str]:
    failed = [
        name
        for name, status in (
            ("text_check_status", text_check_status),
            ("video_quality_status", video_quality_status),
            ("skeleton_static_status", skeleton_static_status),
            ("abnormal_frame_status", abnormal_frame_status),
        )
        if status == "fail"
    ]
    if len(failed) >= 2:
        return "fail", f"fail: fail_indicators={','.join(failed)}"
    if unresolved:
        return (
            "review",
            "review: fail_indicators="
            f"{','.join(failed) or 'none'}; unresolved={','.join(unresolved)}",
        )
    return (
        "pass",
        "pass: fail_indicators="
        f"{','.join(failed) or 'none'}; below_two_fail_threshold",
    )


def supplier_summary_from_details(
    *,
    supplier_name: str,
    expected_clip_count: int,
    details: list[dict[str, Any]],
    main_issue_type: str,
    modules_completed: str,
    blocked_modules: str,
    notes: str,
) -> dict[str, Any]:
    counts = Counter(text(row.get("final_clip_status")) for row in details)
    sample_count = len(details)
    total_frames = sum(int(row.get("total_frames") or 0) for row in details)
    problem_frames = sum(
        interval_frame_count(row.get("_problem_intervals", []))
        for row in details
    )
    covered = counts["pass"] + counts["fail"] + counts["review"]
    return {
        "supplier_name": supplier_name,
        "sample_clip_count": sample_count,
        "expected_clip_count": expected_clip_count,
        "total_frame_count": total_frames,
        "problem_frame_count": problem_frames,
        "problem_frame_ratio": safe_ratio(problem_frames, total_frames),
        "pass_clip_count": counts["pass"],
        "fail_clip_count": counts["fail"],
        "review_clip_count": counts["review"],
        "blocked_clip_count": counts["blocked"],
        "not_run_clip_count": counts["not_run"],
        "coverage_ratio": safe_ratio(covered, sample_count),
        "pass_clip_ratio": safe_ratio(counts["pass"], sample_count),
        "fail_clip_ratio": safe_ratio(counts["fail"], sample_count),
        "main_issue_type": main_issue_type,
        "modules_completed": modules_completed,
        "blocked_modules": blocked_modules,
        "notes": notes,
    }


def build_weekly_issue_events(
    manual_labels: list[dict[str, Any]],
    sam3_rows: list[dict[str, Any]],
    details: list[dict[str, Any]],
) -> list[dict[str, Any]]:
    events: list[dict[str, Any]] = []
    for row in manual_labels:
        events.append(
            {
                "source_module": "manual_review",
                "asset_id": normalize_asset_id(row.get("asset_id")),
                "failure_mode": text(row.get("failure_mode")),
                "manual_outcome": text(row.get("manual_outcome")),
            }
        )
    for row in sam3_rows:
        verdict = text(row.get("window_containment_verdict"))
        events.append(
            {
                "source_module": "sam3_containment",
                "asset_id": normalize_asset_id(row.get("asset_id")),
                "failure_mode": {
                    "mixed_review": "occlusion_or_mask_undersegmentation",
                    "projection_review": "projection_review",
                }.get(verdict, verdict),
                "status": verdict,
            }
        )
    for row in details:
        events.append(
            {
                "source_module": "video_quality",
                "asset_id": row["asset_id"],
                "failure_mode": (
                    "video_quality_pending_colleague_thresholds"
                    if row["video_quality_status"] == "fail"
                    else ""
                ),
                "status": row["video_quality_status"],
            }
        )
        events.append(
            {
                "source_module": "text_integrity",
                "asset_id": row["asset_id"],
                "failure_mode": (
                    "text_check_pending_rules"
                    if row["text_check_status"]
                    in {"pending_rule", "pending_required_rule"}
                    else ""
                ),
                "status": row["text_check_status"],
            }
        )
    return events


def collect_input_audit(run_root: Path) -> dict[str, Any]:
    manifest_path = run_root / "manifests" / "supplier_manifest_xjgt_100.csv"
    manifest_source = read_first_records("xjgt_manifest", [manifest_path])
    manifest = manifest_source.records
    asset_rows, episode_asset = manifest_asset_maps(manifest)
    asset_ids = set(asset_rows)
    sources = load_xjgt_source_reads(run_root)
    inputs = []
    mapped_fields: dict[str, set[str]] = defaultdict(set)
    for source in sources.values():
        normalized_ids: list[str] = []
        unmapped: list[str] = []
        checks: set[str] = set()
        statuses: set[str] = set()
        for row in source.records:
            asset_id = source_asset_id(row, episode_asset)
            if asset_id:
                normalized_ids.append(asset_id)
                if asset_id not in asset_ids:
                    unmapped.append(asset_id)
            check = text(row.get("check"))
            if check:
                checks.add(check)
                target = {
                    "text_integrity": "text_check_status",
                    "quality_score": "supplier_quality_signal",
                    "skeleton_quality_score": "skeleton_missing_status",
                    "keypoint_morphology": "skeleton_morphology_status",
                    "keypoint_temporal": "temporal_status",
                }.get(check)
                if target and asset_id in asset_ids:
                    mapped_fields[target].add(asset_id)
            for key in (
                "status",
                "decision",
                "final_status",
                "window_containment_verdict",
                "manual_outcome",
            ):
                value = text(row.get(key))
                if value:
                    statuses.add(value)
        overlap = len(set(normalized_ids) & asset_ids)
        read_status = source.status
        if source.status == "readable" and source.records and overlap == 0:
            read_status = "readable_unmatched"
        inputs.append(
            {
                "input_name": source.input_name,
                "path": str(source.path) if source.path else "",
                "read_status": read_status,
                "row_count": len(source.records),
                "columns_or_keys": sorted(
                    {key for row in source.records for key in row}
                ),
                "distinct_checks": sorted(checks),
                "distinct_status_values": sorted(statuses),
                "first_normalized_asset_ids": list(
                    dict.fromkeys(normalized_ids)
                )[:3],
                "asset_overlap_count": overlap,
                "unmapped_asset_examples": list(
                    dict.fromkeys(unmapped)
                )[:3],
            }
        )
    inputs.insert(
        0,
        {
            "input_name": "xjgt_manifest",
            "path": str(manifest_path),
            "read_status": manifest_source.status,
            "row_count": len(manifest),
            "columns_or_keys": sorted(
                {key for row in manifest for key in row}
            ),
            "distinct_checks": [],
            "distinct_status_values": [],
            "first_normalized_asset_ids": list(asset_rows)[:3],
            "asset_overlap_count": len(asset_rows),
            "unmapped_asset_examples": [],
        },
    )
    observed_checks = {
        check for item in inputs for check in item["distinct_checks"]
    }
    return {
        "run_root": str(run_root),
        "manifest_asset_count": len(asset_rows),
        "inputs": inputs,
        "mapped_weekly_fields": {
            field: len(ids) for field, ids in sorted(mapped_fields.items())
        },
        "missing_expected_check_names": sorted(
            {"skeleton_quality_score", "keypoint_morphology"}
            - observed_checks
        ),
    }


def read_csv(path: Path) -> list[dict[str, str]]:
    if not path.exists():
        return []
    with path.open(newline="", encoding="utf-8-sig") as handle:
        return list(csv.DictReader(handle))




def write_csv(
    path: Path,
    columns: list[str],
    rows: Iterable[dict[str, Any]],
) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", newline="", encoding="utf-8") as handle:
        writer = csv.DictWriter(handle, fieldnames=columns)
        writer.writeheader()
        writer.writerows(
            {column: row.get(column, "") for column in columns}
            for row in rows
        )


def interval_frame_count(intervals: list[tuple[int, int]]) -> int:
    if not intervals:
        return 0
    normalized = sorted(
        (min(start, end), max(start, end)) for start, end in intervals
    )
    merged: list[list[int]] = []
    for start, end in normalized:
        if not merged or start > merged[-1][1] + 1:
            merged.append([start, end])
        else:
            merged[-1][1] = max(merged[-1][1], end)
    return sum(end - start + 1 for start, end in merged)


def clip_intervals(
    intervals: list[tuple[int, int]],
    total_frames: int,
) -> list[tuple[int, int]]:
    if total_frames <= 0:
        return []
    clipped: list[tuple[int, int]] = []
    max_frame = total_frames - 1
    for start, end in intervals:
        normalized_start = max(0, min(start, end))
        normalized_end = min(max_frame, max(start, end))
        if normalized_start <= normalized_end:
            clipped.append((normalized_start, normalized_end))
    return merge_intervals(clipped)


def merge_intervals(
    intervals: list[tuple[int, int]],
) -> list[tuple[int, int]]:
    if not intervals:
        return []
    normalized = sorted(
        (min(start, end), max(start, end)) for start, end in intervals
    )
    merged: list[list[int]] = []
    for start, end in normalized:
        if not merged or start > merged[-1][1] + 1:
            merged.append([start, end])
        else:
            merged[-1][1] = max(merged[-1][1], end)
    return [(start, end) for start, end in merged]


def format_intervals(intervals: list[tuple[int, int]]) -> str:
    return "|".join(
        f"{start}-{end}" for start, end in merge_intervals(intervals)
    )


def subtract_intervals(
    intervals: list[tuple[int, int]],
    excluded: list[tuple[int, int]],
) -> list[tuple[int, int]]:
    if not intervals or not excluded:
        return intervals
    excluded_normalized = sorted(
        (min(start, end), max(start, end)) for start, end in excluded
    )
    remaining: list[tuple[int, int]] = []
    for interval_start, interval_end in (
        (min(start, end), max(start, end)) for start, end in intervals
    ):
        pieces = [(interval_start, interval_end)]
        for exclude_start, exclude_end in excluded_normalized:
            next_pieces: list[tuple[int, int]] = []
            for piece_start, piece_end in pieces:
                if exclude_end < piece_start or exclude_start > piece_end:
                    next_pieces.append((piece_start, piece_end))
                    continue
                if exclude_start > piece_start:
                    next_pieces.append((piece_start, exclude_start - 1))
                if exclude_end < piece_end:
                    next_pieces.append((exclude_end + 1, piece_end))
            pieces = next_pieces
            if not pieces:
                break
        remaining.extend(pieces)
    return remaining


def intersect_intervals(
    left: list[tuple[int, int]],
    right: list[tuple[int, int]],
) -> list[tuple[int, int]]:
    intersections: list[tuple[int, int]] = []
    for left_start, left_end in (
        (min(start, end), max(start, end)) for start, end in left
    ):
        for right_start, right_end in (
            (min(start, end), max(start, end)) for start, end in right
        ):
            start = max(left_start, right_start)
            end = min(left_end, right_end)
            if start <= end:
                intersections.append((start, end))
    return intersections


def skeleton_static_status_from_parts(
    *,
    missing_status: str,
    morphology_status: str,
    fail_frame_ratio: float,
) -> str:
    if any(
        status in {"not_run", "blocked", "no_valid_output", "source_missing", "unmatched"}
        for status in (missing_status, morphology_status)
    ):
        return "not_run"
    return "fail" if fail_frame_ratio >= 0.10 else "pass"








def resolve_xjgt_video_path(
    manifest_row: dict[str, Any],
    asset_id: str,
) -> Path:
    video_path = text(manifest_row.get("video_path")).strip()
    if video_path:
        return Path(video_path)
    return Path(
        f"/mnt/oss/egodata/XJGT_20260629/video/{asset_id}_video.mp4"
    )


def probe_video_frame_count(path: Path) -> int:
    try:
        import cv2

        capture = cv2.VideoCapture(str(path))
        try:
            if capture.isOpened():
                frame_count = int(capture.get(cv2.CAP_PROP_FRAME_COUNT))
                if frame_count > 0:
                    return frame_count
        finally:
            capture.release()
    except (ImportError, OSError, ValueError):
        pass

    try:
        completed = subprocess.run(
            [
                "ffprobe",
                "-v",
                "error",
                "-select_streams",
                "v:0",
                "-count_frames",
                "-show_entries",
                "stream=nb_read_frames,nb_frames",
                "-of",
                "json",
                str(path),
            ],
            check=False,
            capture_output=True,
            text=True,
            timeout=30,
        )
    except (FileNotFoundError, OSError, subprocess.TimeoutExpired):
        return 0
    if completed.returncode != 0:
        return 0
    try:
        payload = json.loads(completed.stdout)
    except json.JSONDecodeError:
        return 0
    streams = payload.get("streams")
    if not isinstance(streams, list) or not streams:
        return 0
    stream = streams[0] if isinstance(streams[0], dict) else {}
    return (
        integer_or_none(stream.get("nb_read_frames"))
        or integer_or_none(stream.get("nb_frames"))
        or 0
    )


def frame_ratio_status(
    base_status: str,
    fail_frame_count: int,
    fail_frame_ratio: float,
) -> str:
    if fail_frame_ratio >= 0.10:
        return "fail"
    if fail_frame_count > 0 or base_status == "fail":
        return "review"
    return base_status






def count_fail_indicators(
    *,
    text_check_status: str,
    video_quality_status: str,
    skeleton_static_status: str,
    abnormal_frame_status: str,
) -> int:
    return sum(
        status == "fail"
        for status in (
            text_check_status,
            video_quality_status,
            skeleton_static_status,
            abnormal_frame_status,
        )
    )


def recompute_xjgt_final_status(
    *,
    text_check_status: str,
    video_quality_status: str,
    skeleton_static_status: str,
    abnormal_frame_status: str,
    expected_module_missing: bool,
) -> str:
    # quality_hand is intentionally absent: it is supplier evidence, not a
    # vendor-agnostic skeleton validity or final acceptance condition.
    fail_indicator_count = count_fail_indicators(
        text_check_status=text_check_status,
        video_quality_status=video_quality_status,
        skeleton_static_status=skeleton_static_status,
        abnormal_frame_status=abnormal_frame_status,
    )
    if fail_indicator_count >= 2:
        return "fail"
    if abnormal_frame_status == "review" or expected_module_missing:
        return "review"
    return "pass"


def event_frame_interval(
    event: dict[str, Any],
) -> tuple[int, int] | None:
    start = integer_or_none(
        first_value(event, ("start_frame", "window_start_frame", "frame_idx"))
    )
    end = integer_or_none(
        first_value(
            event,
            ("end_frame", "window_end_frame", "frame_idx"),
        )
    )
    if start is None or end is None:
        return None
    return (start, end)










def print_sanity_checks(
    xjgt_summary: dict[str, Any],
    xjgt_details: list[dict[str, Any]],
    deepreach_summary: dict[str, Any],
) -> None:
    total_frames = int(xjgt_summary["total_frame_count"])
    problem_frames = int(xjgt_summary["problem_frame_count"])
    has_nonzero_manual_ratio = any(
        float(row["manual_problem_frame_ratio_of_clip"]) > 0
        for row in xjgt_details
    )
    final_counts = dict(
        sorted(Counter(row["final_clip_status"] for row in xjgt_details).items())
    )
    print(f"XJGT total_frame_count={total_frames} ok={total_frames > 0}")
    print(
        f"XJGT problem_frame_count={problem_frames} "
        f"ok={problem_frames > 0}"
    )
    print(
        "XJGT manual_problem_ratio_nonzero="
        f"{has_nonzero_manual_ratio}"
    )
    print(f"XJGT final_clip_status counts={final_counts}")
    print(
        "DeepReach sample_clip_count="
        f"{deepreach_summary['sample_clip_count']} "
        "video_quality="
        f"{deepreach_summary['video_quality_status']}"
    )


def first_value(
    row: dict[str, Any],
    keys: tuple[str, ...],
) -> Any:
    for key in keys:
        value = row.get(key)
        if value not in (None, ""):
            return value
    return None


def integer_or_none(value: Any) -> int | None:
    if value in (None, ""):
        return None
    try:
        return int(float(value))
    except (TypeError, ValueError):
        return None


def float_or_none(value: Any) -> float | None:
    if value in (None, ""):
        return None
    try:
        return float(value)
    except (TypeError, ValueError):
        return None


def float_or_zero(value: Any) -> float:
    parsed = float_or_none(value)
    return parsed if parsed is not None else 0.0


def normalize_asset_id(value: Any) -> str:
    result = text(value).strip()
    if not result:
        return ""
    if "/" in result or "\\" in result:
        result = Path(result).name
    lower = result.lower()
    for suffix in (".hdf5", ".h5", ".mp4", ".mov", ".mkv", ".avi", ".json", ".parquet", ".csv"):
        if lower.endswith(suffix):
            result = result[: -len(suffix)]
            lower = result.lower()
            break
    for suffix in ("_video", "_hdf5"):
        if result.lower().endswith(suffix):
            result = result[: -len(suffix)]
    return result[:-2] if result.endswith(".0") else result


def status_value(value: Any) -> str:
    status = text(value).lower()
    return {
        "pass_with_notes": "pass",
        "risk": "review",
        "warn": "review",
        "warning": "review",
        "input_missing": "blocked",
        "adapter_missing": "blocked",
    }.get(status, status or "not_run")




def first_issue(value: Any) -> str:
    raw = text(value)
    return raw.split("|", 1)[0] if raw else ""


def most_common(counts: Counter[str]) -> str:
    return counts.most_common(1)[0][0] if counts else ""


def safe_ratio(numerator: int | float, denominator: int | float) -> float:
    return float(numerator) / float(denominator) if denominator else 0.0


def text(value: Any) -> str:
    return "" if value is None else str(value)


if __name__ == "__main__":
    raise SystemExit(main())
