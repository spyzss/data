"""Weekly-template mode for the reusable acceptance ledger builder."""

from __future__ import annotations

import json
import os
import tempfile
from collections import Counter
from copy import copy
from pathlib import Path
from typing import Any, Iterable

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

try:
    from tools import build_weekly_supplier_acceptance_report as weekly
except ModuleNotFoundError as exc:  # Direct execution: python tools/<script>.py
    if exc.name != "tools":
        raise
    import build_weekly_supplier_acceptance_report as weekly


WEEKLY_SHEET_NAMES = (
    "五供应商总览",
    "星际归途",
    "DeepReach",
    "京东JDT",
    "供应商4",
    "供应商5",
    "异常检测拆解",
    "人工问题与阈值",
)
WEEKLY_OVERVIEW_COLUMNS = (
    *weekly.SUMMARY_COLUMNS,
    "review_clip_count",
    "review_clip_ratio",
)
WEEKLY_ABNORMAL_SOURCE_COLUMNS = (
    *weekly.ABNORMAL_SOURCE_COLUMNS,
    "abnormal_fail_frame_count",
    "abnormal_review_frame_count",
)
SUPPLIER_NAMES = {
    "xjgt": "星际归途 / XJGT",
    "deepreach": "DeepReach",
    "jdt": "京东JDT / JDT",
    "supplier4": "供应商4",
    "supplier5": "供应商5",
}


def build_weekly_template_workbook(
    *,
    config: dict[str, Any],
    config_path: Path,
    output_path: Path,
    existing_workbook: Path | None,
) -> Path:
    workbook = _load_template(existing_workbook)
    _ensure_weekly_sheets(workbook)
    canonical_sheet = workbook["星际归途"]
    canonical_header = [cell.value for cell in canonical_sheet[1] if cell.value]
    if not canonical_header:
        canonical_header = weekly.detail_columns_for_rows([])
        _write_detail_sheet(
            canonical_sheet,
            canonical_header,
            [],
            canonical_sheet=None,
        )

    xjgt_details = _sheet_records(canonical_sheet)
    details_by_supplier: dict[str, list[dict[str, Any]]] = {"xjgt": xjgt_details}
    labels_by_supplier: dict[str, list[dict[str, Any]]] = {}
    review_by_supplier: dict[str, list[dict[str, Any]]] = {}
    for raw_config in config.get("suppliers", []):
        supplier_config = dict(raw_config)
        supplier_id = str(supplier_config.get("supplier_id") or "").lower()
        if supplier_id == "xjgt" and xjgt_details:
            continue
        details, labels, review_rows = _load_supplier_details(
            supplier_config,
            config_dir=config_path.parent,
        )
        details_by_supplier[supplier_id] = details
        labels_by_supplier[supplier_id] = labels
        review_by_supplier[supplier_id] = review_rows

    if not xjgt_details and details_by_supplier.get("xjgt"):
        _write_detail_sheet(
            canonical_sheet,
            canonical_header,
            details_by_supplier["xjgt"],
            canonical_sheet=None,
        )

    _write_detail_sheet(
        workbook["DeepReach"],
        canonical_header,
        details_by_supplier.get("deepreach", []),
        canonical_sheet=canonical_sheet,
    )
    _write_detail_sheet(
        workbook["京东JDT"],
        canonical_header,
        details_by_supplier.get("jdt", []),
        canonical_sheet=canonical_sheet,
    )

    summary_rows = [
        _summary_from_details(
            SUPPLIER_NAMES["xjgt"], details_by_supplier.get("xjgt", [])
        ),
        _summary_from_details(
            SUPPLIER_NAMES["deepreach"], details_by_supplier.get("deepreach", [])
        ),
        _summary_from_details(
            SUPPLIER_NAMES["jdt"], details_by_supplier.get("jdt", [])
        ),
        _summary_from_details("供应商4", []),
        _summary_from_details("供应商5", []),
    ]
    _rewrite_tabular_sheet(
        workbook["五供应商总览"], WEEKLY_OVERVIEW_COLUMNS, summary_rows
    )
    decomposition_rows = [
        _abnormal_source_row(name, details_by_supplier.get(supplier_id, []))
        for supplier_id, name in (
            ("xjgt", SUPPLIER_NAMES["xjgt"]),
            ("deepreach", SUPPLIER_NAMES["deepreach"]),
            ("jdt", SUPPLIER_NAMES["jdt"]),
            ("supplier4", "供应商4"),
            ("supplier5", "供应商5"),
        )
    ]
    _rewrite_tabular_sheet(
        workbook["异常检测拆解"],
        WEEKLY_ABNORMAL_SOURCE_COLUMNS,
        decomposition_rows,
    )
    _append_manual_issue_rows(
        workbook["人工问题与阈值"],
        labels_by_supplier=labels_by_supplier,
        review_by_supplier=review_by_supplier,
    )
    workbook._sheets = [workbook[name] for name in WEEKLY_SHEET_NAMES]
    workbook.properties.title = str(
        config.get("output_label") or "Weekly Supplier Acceptance"
    )
    _save_atomic(workbook, output_path)
    return output_path


def _load_supplier_details(
    config: dict[str, Any],
    *,
    config_dir: Path,
) -> tuple[list[dict[str, Any]], list[dict[str, Any]], list[dict[str, Any]]]:
    supplier_id = str(config.get("supplier_id") or "").lower()
    manifest_source = _read_source(config, "manifest", config_dir, required=True)
    manifest = manifest_source.records
    asset_rows, episode_asset = weekly.manifest_asset_maps(manifest)
    check_source = _read_source(config, "precheck_check_results", config_dir)
    aggregate_source = _read_source(config, "precheck_clip_aggregates", config_dir)
    ledger_source = _read_source(config, "ledger_asset_ledger", config_dir)
    candidate_source = _read_source(config, "candidate_windows", config_dir)
    video_source = _read_source(config, "video_quality", config_dir)
    sam3_source = _read_source(config, "sam3_summary", config_dir)
    review_source = _read_source(config, "review_queue", config_dir)
    manual_source = _read_source(config, "manual_labels", config_dir)
    labels = weekly.dedupe_manual_labels(
        [_manual_fallback_bounds(row) for row in manual_source.records]
    )
    review_rows = [_window_review_row(row) for row in review_source.records]
    sources = {
        "precheck_check_results": check_source,
        "precheck_candidate_windows": candidate_source,
    }
    precheck = weekly.load_precheck_evidence(manifest, sources)
    asset_ids = set(asset_rows)
    candidate_by_asset, _ = weekly.map_window_rows(
        candidate_source.records, asset_ids, episode_asset
    )
    video_by_asset, _ = weekly.map_video_quality_evidence(
        video_source.records, asset_ids, episode_asset
    )
    sam3_by_asset, _ = weekly.map_sam3_evidence(
        sam3_source.records, asset_ids, episode_asset
    )
    submitted_by_asset, _ = weekly.map_window_rows(
        review_rows, asset_ids, episode_asset
    )
    frame_info = {
        asset_id: _manifest_frame_info(row)
        for asset_id, row in asset_rows.items()
    }
    manual_by_asset = weekly.aggregate_manual_review(
        labels,
        {
            asset_id: frame["total_frames"]
            for asset_id, frame in frame_info.items()
        },
    )
    ledger_by_asset = {
        weekly.normalize_asset_id(row.get("asset_id")): row
        for row in ledger_source.records
        if weekly.normalize_asset_id(row.get("asset_id"))
    }
    blocked = _truthy(config.get("blocked"))
    blocker_source = _read_source(config, "blocker", config_dir)
    blocker_reason = next(
        (
            str(row.get("reason"))
            for row in blocker_source.records
            if row.get("reason")
        ),
        "",
    )
    evidence_paths = [
        source.path
        for source in (
            check_source,
            aggregate_source,
            ledger_source,
            candidate_source,
            video_source,
            sam3_source,
            review_source,
            manual_source,
            blocker_source,
        )
        if source.path is not None
    ]
    details = []
    for asset_id, manifest_row in asset_rows.items():
        frame = frame_info[asset_id]
        sam3_row = sam3_by_asset.get(asset_id, {})
        if blocked:
            sam3_row = {
                "status": "blocked",
                "auto_fail_intervals": [],
                "review_intervals": [],
                "rows": [],
            }
        detail = weekly.build_xjgt_detail(
            asset_id=asset_id,
            total_frames=frame["total_frames"],
            frame_count_status=frame["frame_count_status"],
            video_path=Path(
                str(
                    manifest_row.get("primary_video_path")
                    or manifest_row.get("video_path")
                    or ""
                )
            ),
            precheck_row=precheck.by_asset.get(
                asset_id, weekly.empty_precheck_evidence()
            ),
            precheck_source_status=precheck.source_status,
            precheck_unmatched=precheck.unmatched_assets,
            candidate_row=candidate_by_asset.get(asset_id, {}),
            video_row=video_by_asset.get(asset_id, {}),
            sam3_row=sam3_row,
            submitted_review_row=submitted_by_asset.get(asset_id, {}),
            submitted_review_source_status=review_source.status,
            manual_row=manual_by_asset.get(
                asset_id, weekly.empty_manual_evidence(frame["total_frames"])
            ),
            ledger_row=ledger_by_asset.get(asset_id, {}),
            require_xjgt_text=_truthy(config.get("text_required")),
            evidence_paths=[Path(path) for path in evidence_paths],
        )
        _apply_final_priority(
            detail,
            blocked=blocked,
            blocker_reason=blocker_reason,
        )
        details.append(detail)
    return details, labels, review_rows


def _apply_final_priority(
    detail: dict[str, Any],
    *,
    blocked: bool,
    blocker_reason: str,
) -> None:
    manual_fail_frames = int(detail.get("manual_true_positive_frame_count") or 0)
    unresolved_frames = int(detail.get("unreviewed_submitted_frame_count") or 0)
    submitted_count = int(detail.get("submitted_review_interval_count") or 0)
    temporal_candidate_frames = int(detail.get("temporal_detected_frame_count") or 0)
    sam3_detected_frames = int(detail.get("sam3_detected_frame_count") or 0)
    if manual_fail_frames:
        abnormal_status = "fail"
        reason = f"manual_true_positive_frames={manual_fail_frames}"
    elif blocked:
        abnormal_status = "review"
        reason = blocker_reason or "required abnormal-frame module blocked"
    elif unresolved_frames:
        abnormal_status = "review"
        reason = f"unresolved_review_frames={unresolved_frames}"
    elif submitted_count == 0 and (temporal_candidate_frames or sam3_detected_frames):
        abnormal_status = "review"
        reason = "automatic abnormal evidence has no submitted manual resolution"
    else:
        abnormal_status = "pass"
        reason = "no unresolved or manually confirmed abnormal frames"
    detail["abnormal_frame_status"] = abnormal_status
    detail["abnormal_frame_status_v2"] = abnormal_status
    detail["abnormal_status_reason"] = reason
    detail["abnormal_status_reason_v2"] = reason
    if blocked:
        detail["sam3_containment_status"] = "blocked"

    status_fields = (
        "text_check_status",
        "skeleton_static_status",
        "video_quality_status",
        "abnormal_frame_status",
    )
    statuses = [str(detail.get(field) or "not_evaluated") for field in status_fields]
    fail_count = sum(status == "fail" for status in statuses)
    unresolved = any(
        status in {
            "review",
            "blocked",
            "not_evaluated",
            "not_ready",
            "not_run",
            "no_valid_output",
            "source_missing",
        }
        for status in statuses
    )
    final = "fail" if fail_count else "review" if unresolved else "pass"
    detail["fail_indicator_count"] = fail_count
    detail["acceptance_status"] = final
    detail["acceptance_status_v2"] = final
    detail["review_status"] = "review" if final == "review" else "completed"
    detail["review_status_v2"] = detail["review_status"]
    detail["final_status_reason"] = (
        f"priority=fail>review>pass; statuses={','.join(statuses)}"
    )
    detail["final_status_reason_v2"] = detail["final_status_reason"]


def _summary_from_details(
    supplier_name: str, details: list[dict[str, Any]]
) -> dict[str, Any]:
    counts = Counter(str(row.get("acceptance_status") or "review") for row in details)
    total_frames = sum(int(row.get("total_frames") or 0) for row in details)
    problem_frames = sum(int(row.get("problem_frame_count") or 0) for row in details)
    sample_count = len(details)
    return {
        "supplier_name": supplier_name,
        "sample_clip_count": sample_count,
        "total_frame_count": total_frames,
        "problem_frame_count": problem_frames,
        "problem_frame_ratio": _ratio(problem_frames, total_frames),
        "pass_clip_count": counts["pass"],
        "pass_clip_ratio": _ratio(counts["pass"], sample_count),
        "fail_clip_count": counts["fail"],
        "fail_clip_ratio": _ratio(counts["fail"], sample_count),
        "review_clip_count": counts["review"],
        "review_clip_ratio": _ratio(counts["review"], sample_count),
    }


def _abnormal_source_row(
    supplier_name: str, details: list[dict[str, Any]]
) -> dict[str, Any]:
    row = {column: 0 for column in WEEKLY_ABNORMAL_SOURCE_COLUMNS}
    row.update(
        {
            "supplier_name": supplier_name,
            "evaluation_scope": "auto_triggered_review_only" if details else "not_ready",
            "measurement_note": weekly.ABNORMAL_COVERAGE_NOTE,
        }
    )
    sum_fields = {
        column
        for column in weekly.ABNORMAL_SOURCE_COLUMNS
        if column.endswith("_count")
    }
    for column in sum_fields:
        row[column] = sum(int(detail.get(column) or 0) for detail in details)
    row["abnormal_review_frame_count"] = sum(
        int(detail.get("abnormal_v2_unreviewed_as_review_frame_count") or 0)
        for detail in details
    )
    for numerator, denominator, ratio_name in (
        (
            "precheck_to_sam3_frame_count",
            "precheck_temporal_candidate_frame_count",
            "precheck_to_sam3_frame_ratio",
        ),
        (
            "sam3_to_manual_frame_count",
            "sam3_processed_frame_count",
            "sam3_to_manual_frame_ratio",
        ),
        (
            "manual_true_positive_frame_count",
            "manual_submitted_frame_count",
            "manual_true_problem_to_submitted_ratio",
        ),
        (
            "manual_true_positive_frame_count",
            "manual_reviewed_frame_count",
            "manual_true_problem_to_reviewed_ratio",
        ),
    ):
        row[ratio_name] = _ratio(row[numerator], row[denominator])
    for name in (
        "temporal_manual_tp_coverage_rate",
        "sam3_manual_tp_coverage_rate",
        "auto_union_manual_tp_coverage_rate",
        "reviewed_auto_confirmation_rate",
        "blind_manual_auto_recall",
    ):
        row[name] = "not_measurable"
    return row


def _append_manual_issue_rows(
    sheet,
    *,
    labels_by_supplier: dict[str, list[dict[str, Any]]],
    review_by_supplier: dict[str, list[dict[str, Any]]],
) -> None:
    marker = "新增供应商人工问题统计"
    for row_index in range(1, sheet.max_row + 1):
        if sheet.cell(row_index, 1).value == marker:
            sheet.delete_rows(row_index, sheet.max_row - row_index + 1)
            break
    issue_rows = []
    for supplier_id, labels in labels_by_supplier.items():
        issue_rows.extend(
            weekly.build_manual_issue_rows(
                supplier_name=SUPPLIER_NAMES.get(supplier_id, supplier_id),
                labels=labels,
            )
        )
        reviewed_ids = {
            str(row.get("review_id"))
            for row in labels
            if row.get("review_id")
        }
        unresolved = [
            row
            for row in review_by_supplier.get(supplier_id, [])
            if str(row.get("review_id") or "") not in reviewed_ids
        ]
        if unresolved:
            issue_rows.append(
                {
                    "supplier_name": SUPPLIER_NAMES.get(supplier_id, supplier_id),
                    "issue_type": "unresolved_review",
                    "manual_label_count": len(unresolved),
                    "true_positive_count": 0,
                    "false_positive_count": 0,
                    "acceptable_flagged_count": 0,
                    "review_count": len(unresolved),
                }
            )
    if not issue_rows:
        return
    start = sheet.max_row + 2
    sheet.cell(start, 1, marker).font = Font(bold=True)
    for column_index, column in enumerate(weekly.MANUAL_ISSUE_COLUMNS, 1):
        cell = sheet.cell(start + 1, column_index, column)
        cell.font = Font(bold=True)
        cell.fill = PatternFill("solid", fgColor="D9EAF7")
    for row_index, row in enumerate(issue_rows, start + 2):
        for column_index, column in enumerate(weekly.MANUAL_ISSUE_COLUMNS, 1):
            sheet.cell(row_index, column_index, row.get(column, ""))
    caveat_row = start + 2 + len(issue_rows)
    sheet.cell(
        caveat_row,
        1,
        "定向复核不是无偏的全局召回率估计；需要独立随机抽检评估召回率。",
    )


def _write_detail_sheet(
    sheet,
    header: list[str],
    rows: list[dict[str, Any]],
    *,
    canonical_sheet,
) -> None:
    sheet.delete_rows(1, sheet.max_row)
    sheet.append(header)
    for row in rows:
        sheet.append([_excel_value(row.get(column, "")) for column in header])
    if canonical_sheet is not None:
        _clone_sheet_format(canonical_sheet, sheet, len(header), len(rows))
    else:
        _default_format(sheet, header)
    sheet.auto_filter.ref = sheet.dimensions


def _clone_sheet_format(source, target, column_count: int, row_count: int) -> None:
    target.freeze_panes = source.freeze_panes
    target.sheet_view.showGridLines = source.sheet_view.showGridLines
    target.row_dimensions[1].height = source.row_dimensions[1].height
    for column_index in range(1, column_count + 1):
        source_cell = source.cell(1, column_index)
        target_cell = target.cell(1, column_index)
        _copy_cell_style(source_cell, target_cell)
        letter = get_column_letter(column_index)
        target.column_dimensions[letter].width = source.column_dimensions[letter].width
    prototype_row = 2 if source.max_row >= 2 else 1
    for row_index in range(2, row_count + 2):
        target.row_dimensions[row_index].height = source.row_dimensions[prototype_row].height
        for column_index in range(1, column_count + 1):
            _copy_cell_style(
                source.cell(prototype_row, column_index),
                target.cell(row_index, column_index),
            )


def _copy_cell_style(source, target) -> None:
    target.font = copy(source.font)
    target.fill = copy(source.fill)
    target.border = copy(source.border)
    target.alignment = copy(source.alignment)
    target.number_format = source.number_format
    target.protection = copy(source.protection)


def _default_format(sheet, header: Iterable[str]) -> None:
    sheet.freeze_panes = "A2"
    sheet.sheet_view.showGridLines = False
    for cell in sheet[1]:
        cell.font = Font(bold=True)
        cell.fill = PatternFill("solid", fgColor="D9EAF7")
    for index, column in enumerate(header, 1):
        sheet.column_dimensions[get_column_letter(index)].width = min(
            max(len(str(column)) + 2, 12), 42
        )


def _rewrite_tabular_sheet(sheet, columns: Iterable[str], rows: list[dict[str, Any]]) -> None:
    columns = list(columns)
    source_header = [copy(cell._style) for cell in sheet[1]] if sheet.max_row else []
    sheet.delete_rows(1, sheet.max_row)
    sheet.append(columns)
    for row in rows:
        sheet.append([_excel_value(row.get(column, "")) for column in columns])
    for index, cell in enumerate(sheet[1]):
        if index < len(source_header):
            cell._style = copy(source_header[index])
        else:
            cell.font = Font(bold=True)
            cell.fill = PatternFill("solid", fgColor="D9EAF7")
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions


def _ensure_weekly_sheets(workbook) -> None:
    if "供应商3" in workbook.sheetnames and "京东JDT" not in workbook.sheetnames:
        workbook["供应商3"].title = "京东JDT"
    for name in list(workbook.sheetnames):
        if name not in WEEKLY_SHEET_NAMES:
            workbook.remove(workbook[name])
    for name in WEEKLY_SHEET_NAMES:
        if name not in workbook.sheetnames:
            workbook.create_sheet(name)


def _load_template(existing_workbook: Path | None):
    if existing_workbook is not None:
        return load_workbook(existing_workbook)
    workbook = Workbook()
    workbook.active.title = "五供应商总览"
    for name in WEEKLY_SHEET_NAMES[1:]:
        workbook.create_sheet(name)
    return workbook


def _sheet_records(sheet) -> list[dict[str, Any]]:
    if sheet.max_row < 2:
        return []
    header = [cell.value for cell in sheet[1]]
    return [
        dict(zip(header, row))
        for row in sheet.iter_rows(min_row=2, values_only=True)
        if any(value not in (None, "") for value in row)
    ]


def _read_source(
    config: dict[str, Any], key: str, config_dir: Path, required: bool = False
) -> weekly.SourceRead:
    value = config.get(key)
    if not value:
        if required:
            raise FileNotFoundError(f"required input {key} is not configured")
        return weekly.SourceRead(key, None, [], "missing")
    path = _resolve_path(value, config_dir)
    if not path.exists():
        if required:
            raise FileNotFoundError(f"required input {key} not found: {path}")
        return weekly.SourceRead(key, path, [], "missing")
    try:
        return weekly.SourceRead(
            key, path, weekly.read_table_records(path), "readable"
        )
    except (OSError, ValueError) as exc:
        if required:
            raise
        return weekly.SourceRead(key, path, [], f"unreadable:{exc}")


def _resolve_path(value: Any, config_dir: Path) -> Path:
    path = Path(str(value)).expanduser()
    if path.is_absolute():
        return path
    cwd = (Path.cwd() / path).resolve()
    local = (config_dir / path).resolve()
    return cwd if cwd.exists() or not local.exists() else local


def _manifest_frame_info(row: dict[str, Any]) -> dict[str, Any]:
    start = weekly.integer_or_none(
        weekly.first_value(row, ("start_frame", "source_start_frame", "clip_start_frame"))
    )
    end = weekly.integer_or_none(
        weekly.first_value(row, ("end_frame", "source_end_frame", "clip_end_frame"))
    )
    if start is not None and end is not None and end >= start:
        return {"total_frames": end - start + 1, "frame_count_status": "manifest_range"}
    count = weekly.integer_or_none(
        weekly.first_value(
            row,
            ("clip_frame_count", "total_frame_count", "total_frames", "num_frames", "frame_count"),
        )
    ) or 0
    return {
        "total_frames": count,
        "frame_count_status": "manifest_fallback" if count else "unreadable",
    }


def _manual_fallback_bounds(row: dict[str, Any]) -> dict[str, Any]:
    output = dict(row)
    if weekly.is_confirmed_manual_outcome(output) and (
        weekly.integer_or_none(output.get("affected_start_frame")) is None
        or weekly.integer_or_none(output.get("affected_end_frame")) is None
    ):
        output["affected_start_frame"] = output.get("window_start_frame", "")
        output["affected_end_frame"] = output.get("window_end_frame", "")
    return output


def _window_review_row(row: dict[str, Any]) -> dict[str, Any]:
    output = dict(row)
    output.setdefault("source_level", "window")
    return output


def _truthy(value: Any) -> bool:
    if isinstance(value, bool):
        return value
    return str(value or "").lower() in {"1", "true", "yes", "on"}


def _ratio(numerator: int, denominator: int) -> float:
    return numerator / denominator if denominator else 0.0


def _excel_value(value: Any) -> Any:
    if isinstance(value, (dict, list, tuple, set)):
        return json.dumps(value, ensure_ascii=False, sort_keys=True)
    if isinstance(value, str) and value.startswith(("=", "+", "-", "@")):
        return "'" + value
    return value


def _save_atomic(workbook, output_path: Path) -> None:
    output_path.parent.mkdir(parents=True, exist_ok=True)
    with tempfile.NamedTemporaryFile(
        suffix=".xlsx", dir=output_path.parent, delete=False
    ) as handle:
        temporary = Path(handle.name)
    try:
        workbook.save(temporary)
        os.replace(temporary, output_path)
    finally:
        temporary.unlink(missing_ok=True)
