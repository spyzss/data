#!/usr/bin/env python3
"""Build the final XJGT acceptance ledger and Excel workbook."""

from __future__ import annotations

import argparse
import csv
import json
import logging
import shutil
from collections import Counter, defaultdict
from dataclasses import dataclass
from pathlib import Path
from typing import Any, Iterable, Sequence

import pandas as pd
from openpyxl import Workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from tools.build_batch_qc_ledger import normalize_video_quality_row


LOGGER = logging.getLogger("build_xjgt_acceptance_report")

ASSET_LEDGER_COLUMNS = [
    "supplier_id",
    "asset_id",
    "hdf5_path",
    "video_path",
    "precheck_status",
    "video_quality_raw_decision",
    "video_quality_status",
    "sam3_evidence_status",
    "manual_review_status",
    "final_verdict",
    "risk_level",
    "top_issue_types",
    "true_positive_segment_count",
    "false_positive_count",
    "acceptable_flagged_count",
    "key_metrics_json",
    "evidence_paths",
    "notes",
]

ISSUE_EVENT_COLUMNS = [
    "supplier_id",
    "asset_id",
    "source_module",
    "source_verdict",
    "manual_outcome",
    "algorithm_outcome",
    "failure_mode",
    "severity",
    "confidence",
    "start_frame",
    "end_frame",
    "representative_frame",
    "reason",
    "key_metrics_json",
    "evidence_path",
    "reviewer",
]

METHOD_THRESHOLDS = [
    (
        "video_quality_fail_mapping",
        "video_quality",
        "raw decision = fail",
        "video_quality_status=fail; final asset fail",
    ),
    (
        "video_quality_warn_mapping",
        "video_quality",
        "raw decision in warn, warning, pass_with_notes",
        "video_quality_status=pass_with_notes; not a hard fail",
    ),
    (
        "video_quality_pass_mapping",
        "video_quality",
        "raw decision = pass",
        "video_quality_status=pass",
    ),
    (
        "sam3_projected_in_image_ratio_projection_review",
        "sam3_containment",
        "projected_in_image_ratio < 0.8",
        "projection_review",
    ),
    (
        "sam3_inside_ratio_strong_mismatch",
        "sam3_containment",
        "inside_ratio <= 0.2",
        "strong_containment_mismatch",
    ),
    (
        "sam3_inside_ratio_review",
        "sam3_containment",
        "inside_ratio < 0.6",
        "containment_review",
    ),
    (
        "sam3_inside_ratio_likely_visible",
        "sam3_containment",
        "inside_ratio >= 0.6",
        "likely_visible_ok",
    ),
    (
        "sam3_window_strong_fail_frame_count",
        "sam3_containment",
        "strong_fail_frame_count >= 3",
        "containment_fail candidate",
    ),
    (
        "sam3_window_strong_fail_frame_ratio",
        "sam3_containment",
        "strong_fail_frame_ratio >= 0.6",
        "containment_fail candidate",
    ),
    (
        "sam3_window_mean_containment_fail",
        "sam3_containment",
        "projected_in_image_ratio_mean >= 0.8 and inside_ratio_mean <= 0.2 "
        "and projection_review_frame_count == 0",
        "containment_fail candidate",
    ),
]

PIPELINE_NOTES = [
    "quality_hand is a supplier quality signal, not universal skeleton quality.",
    "Temporal metrics are maxima over hand keypoints, so one jittering point can trigger a window.",
    "SAM3 containment measures keypoint-in-mask evidence, not full mask-region validity.",
    "hand_out_of_frame remains a manual/heuristic judgment and is not fully automated.",
    "The Top-60 manual review queue is high-priority biased and is not an unbiased overall error-rate estimate.",
    "SAM3 containment_fail is risk evidence until manually confirmed; projection_review is not a hard fail.",
    "All SAM3 thresholds listed here are v0 heuristics pending calibration.",
]

STATUS_RANK = {
    "not_run": 0,
    "pass": 1,
    "pass_with_notes": 2,
    "review": 3,
    "fail": 4,
}


@dataclass(frozen=True)
class OutputPaths:
    asset_ledger_csv: Path
    issue_events_csv: Path
    workbook_xlsx: Path
    summary_json: Path


def parse_args(argv: Sequence[str] | None = None) -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        description="Build the final XJGT asset ledger, issue events, summary, and Excel report."
    )
    parser.add_argument("--quality-archive", required=True, type=Path)
    parser.add_argument("--legacy-reconciliation-manifest", "--manifest", dest="legacy_manifest", type=Path)
    parser.add_argument("--legacy-reconciliation-precheck-check-results", "--precheck-check-results", dest="legacy_precheck_check_results", type=Path)
    parser.add_argument("--legacy-reconciliation-precheck-clip-aggregates", "--precheck-clip-aggregates", dest="legacy_precheck_clip_aggregates", type=Path)
    parser.add_argument("--legacy-reconciliation-precheck-candidate-windows", "--precheck-candidate-windows", dest="legacy_precheck_candidate_windows", type=Path)
    parser.add_argument("--legacy-reconciliation-video-quality-results", "--video-quality-results", dest="legacy_video_quality_results", type=Path)
    parser.add_argument("--video-quality-summary", type=Path)
    parser.add_argument("--sam3-window-summary", type=Path)
    parser.add_argument("--sam3-clip-summary", type=Path)
    parser.add_argument("--sam3-frame-results", type=Path)
    parser.add_argument("--manual-review-labels", type=Path)
    parser.add_argument("--manual-review-csv", type=Path)
    parser.add_argument("--output-dir", required=True, type=Path)
    parser.add_argument("--log-level", default="INFO")
    return parser.parse_args(argv)


def main(argv: Sequence[str] | None = None) -> int:
    args = parse_args(argv)
    logging.basicConfig(
        level=getattr(logging, args.log_level.upper()),
        format="%(asctime)s [%(levelname)s] %(name)s: %(message)s",
    )
    if args.quality_archive is not None:
        from tools.build_qc_json_projection import run_projection_cli

        projected = run_projection_cli(
            args.quality_archive,
            args.output_dir,
            formats=("csv", "parquet", "xlsx", "markdown"),
            legacy_candidate_windows=args.legacy_precheck_candidate_windows,
            legacy_sam3_window_summary=args.sam3_window_summary,
            legacy_video_quality=args.legacy_video_quality_results,
            legacy_issue_events=args.manual_review_labels,
        )
        LOGGER.info("Wrote %s", projected.get("xlsx", args.output_dir / "qc_projection.xlsx"))
        LOGGER.info("Wrote %s", projected.get("asset_csv", args.output_dir / "assets.csv"))
        return 0
    required_legacy = {
        "--manifest": args.legacy_manifest,
        "--precheck-clip-aggregates": args.legacy_precheck_clip_aggregates,
        "--precheck-candidate-windows": args.legacy_precheck_candidate_windows,
        "--video-quality-results": args.legacy_video_quality_results,
        "--sam3-window-summary": args.sam3_window_summary,
        "--manual-review-labels": args.manual_review_labels,
    }
    missing = [name for name, value in required_legacy.items() if value is None]
    if missing:
        raise SystemExit(
            "--quality-archive is required for formal output; missing legacy arguments: "
            + ", ".join(missing)
        )
    outputs = build_acceptance_outputs(
        manifest=args.legacy_manifest,
        precheck_clip_aggregates=args.legacy_precheck_clip_aggregates,
        precheck_candidate_windows=args.legacy_precheck_candidate_windows,
        video_quality_results=args.legacy_video_quality_results,
        sam3_window_summary=args.sam3_window_summary,
        manual_review_labels=args.manual_review_labels,
        output_dir=args.output_dir,
        precheck_check_results=args.legacy_precheck_check_results,
        video_quality_summary=args.video_quality_summary,
        sam3_clip_summary=args.sam3_clip_summary,
        sam3_frame_results=args.sam3_frame_results,
        manual_review_csv=args.manual_review_csv,
    )
    for path in (
        outputs.asset_ledger_csv,
        outputs.issue_events_csv,
        outputs.workbook_xlsx,
        outputs.summary_json,
    ):
        LOGGER.info("Wrote %s", path)
    return 0


def build_acceptance_outputs(
    *,
    manifest: Path | None = None,
    precheck_clip_aggregates: Path | None = None,
    precheck_candidate_windows: Path | None = None,
    video_quality_results: Path | None = None,
    sam3_window_summary: Path | None = None,
    manual_review_labels: Path | None = None,
    output_dir: Path,
    precheck_check_results: Path | None = None,
    video_quality_summary: Path | None = None,
    sam3_clip_summary: Path | None = None,
    sam3_frame_results: Path | None = None,
    manual_review_csv: Path | None = None,
    quality_archive: Path | None = None,
) -> OutputPaths:
    output_dir.mkdir(parents=True, exist_ok=True)
    if quality_archive is not None:
        from tools.build_qc_json_projection import run_projection_cli

        projected = run_projection_cli(
            quality_archive,
            output_dir,
            formats=("csv", "parquet", "xlsx", "markdown"),
            legacy_candidate_windows=precheck_candidate_windows,
            legacy_sam3_window_summary=sam3_window_summary,
            legacy_video_quality=video_quality_results,
            legacy_issue_events=manual_review_labels,
        )
        outputs = OutputPaths(
            asset_ledger_csv=output_dir / "xjgt_100_asset_ledger.csv",
            issue_events_csv=output_dir / "xjgt_100_issue_events.csv",
            workbook_xlsx=output_dir / "xjgt_100_acceptance_report.xlsx",
            summary_json=output_dir / "xjgt_100_summary.json",
        )
        if projected.get("asset_csv") is not None:
            shutil.copyfile(projected["asset_csv"], outputs.asset_ledger_csv)
        if projected.get("issue_csv") is not None:
            shutil.copyfile(projected["issue_csv"], outputs.issue_events_csv)
        if projected.get("xlsx") is not None:
            shutil.copyfile(projected["xlsx"], outputs.workbook_xlsx)
        if projected.get("statistics_json") is not None:
            shutil.copyfile(projected["statistics_json"], outputs.summary_json)
        return outputs
    required_inputs = {
        "manifest": manifest,
        "precheck_clip_aggregates": precheck_clip_aggregates,
        "precheck_candidate_windows": precheck_candidate_windows,
        "video_quality_results": video_quality_results,
        "sam3_window_summary": sam3_window_summary,
        "manual_review_labels": manual_review_labels,
    }
    missing_inputs = [name for name, value in required_inputs.items() if value is None]
    if missing_inputs:
        raise TypeError(
            "legacy XJGT output requires inputs: " + ", ".join(missing_inputs)
        )
    manifest_rows = read_records(manifest)
    assets, episode_to_asset = manifest_assets(manifest_rows)
    events: list[dict[str, Any]] = []
    module_details: dict[str, dict[str, Any]] = {
        asset_id: {
            "precheck_status": "pass",
            "video_quality_raw_decision": "not_run",
            "video_quality_status": "not_run",
            "sam3_evidence_status": "not_run",
            "manual_review_status": "not_run",
            "metrics": {},
        }
        for asset_id in assets
    }

    add_precheck_evidence(
        read_records(precheck_clip_aggregates),
        precheck_clip_aggregates,
        assets,
        episode_to_asset,
        module_details,
        events,
    )
    add_candidate_evidence(
        read_records(precheck_candidate_windows),
        precheck_candidate_windows,
        assets,
        episode_to_asset,
        module_details,
        events,
    )
    add_video_quality_evidence(
        read_records(video_quality_results),
        video_quality_results,
        assets,
        episode_to_asset,
        module_details,
        events,
    )
    sam3_rows = read_records(sam3_window_summary)
    add_sam3_evidence(
        sam3_rows,
        sam3_window_summary,
        assets,
        episode_to_asset,
        module_details,
        events,
    )
    manual_rows = load_manual_rows(manual_review_labels)
    add_manual_evidence(
        manual_rows,
        manual_review_labels,
        assets,
        episode_to_asset,
        module_details,
        events,
    )

    extra_evidence_paths = [
        path
        for path in (
            precheck_check_results,
            video_quality_summary,
            sam3_clip_summary,
            sam3_frame_results,
            manual_review_csv,
        )
        if path is not None
    ]
    ledger_rows = build_asset_ledger(
        assets,
        module_details,
        events,
        extra_evidence_paths,
    )
    summary = build_summary(ledger_rows, events, sam3_rows)

    outputs = OutputPaths(
        asset_ledger_csv=output_dir / "xjgt_100_asset_ledger.csv",
        issue_events_csv=output_dir / "xjgt_100_issue_events.csv",
        workbook_xlsx=output_dir / "xjgt_100_acceptance_report.xlsx",
        summary_json=output_dir / "xjgt_100_summary.json",
    )
    write_csv(outputs.asset_ledger_csv, ASSET_LEDGER_COLUMNS, ledger_rows)
    write_csv(outputs.issue_events_csv, ISSUE_EVENT_COLUMNS, events)
    outputs.summary_json.write_text(
        json.dumps(summary, ensure_ascii=False, indent=2, sort_keys=True),
        encoding="utf-8",
    )
    write_workbook(
        outputs.workbook_xlsx,
        ledger_rows,
        events,
        manual_rows,
        sam3_rows,
        read_records(video_quality_results),
        summary,
    )
    return outputs


def read_records(path: Path) -> list[dict[str, Any]]:
    if not path.exists():
        raise FileNotFoundError(path)
    suffix = path.suffix.lower()
    if suffix == ".csv":
        frame = pd.read_csv(path, dtype={"asset_id": str})
        return dataframe_records(frame)
    if suffix == ".parquet":
        return dataframe_records(pd.read_parquet(path))
    if suffix != ".json":
        raise ValueError(f"unsupported input extension: {path}")
    data = json.loads(path.read_text(encoding="utf-8"))
    return records_from_json_value(data)


def records_from_json_value(data: Any) -> list[dict[str, Any]]:
    if isinstance(data, list):
        return [dict(row) for row in data if isinstance(row, dict)]
    if not isinstance(data, dict):
        return []
    for key in (
        "rows",
        "results",
        "assets",
        "records",
        "items",
        "segments",
        "windows",
    ):
        value = data.get(key)
        if isinstance(value, list):
            return [dict(row) for row in value if isinstance(row, dict)]
    if all(isinstance(value, dict) for value in data.values()):
        rows = []
        for key, value in data.items():
            row = dict(value)
            row.setdefault("asset_id", key)
            rows.append(row)
        return rows
    return [dict(data)]


def dataframe_records(frame: pd.DataFrame) -> list[dict[str, Any]]:
    frame = frame.where(pd.notna(frame), None)
    return [dict(row) for row in frame.to_dict(orient="records")]


def load_manual_rows(path: Path) -> list[dict[str, Any]]:
    if path.suffix.lower() != ".json":
        return read_records(path)
    data = json.loads(path.read_text(encoding="utf-8"))
    if isinstance(data, dict) and isinstance(data.get("segments"), list):
        return [dict(row) for row in data["segments"] if isinstance(row, dict)]
    return records_from_json_value(data)


def manifest_assets(
    rows: list[dict[str, Any]],
) -> tuple[dict[str, dict[str, Any]], dict[int, str]]:
    assets: dict[str, dict[str, Any]] = {}
    episode_to_asset: dict[int, str] = {}
    for index, row in enumerate(rows):
        asset_id = normalize_asset_id(
            first_present(row, ("asset_id", "content_id", "clip_id"), str(index))
        )
        episode_idx = integer_or_none(row.get("episode_idx"))
        if episode_idx is None:
            episode_idx = index
        assets[asset_id] = {
            "supplier_id": str(
                first_present(row, ("supplier_id", "supplier", "vendor_id"), "unknown")
            ),
            "asset_id": asset_id,
            "hdf5_path": text(first_present(row, ("hdf5_path", "hdf5"), "")),
            "video_path": text(first_present(row, ("video_path", "video"), "")),
        }
        episode_to_asset[episode_idx] = asset_id
    return assets, episode_to_asset


def add_precheck_evidence(
    rows: list[dict[str, Any]],
    path: Path,
    assets: dict[str, dict[str, Any]],
    episode_to_asset: dict[int, str],
    details: dict[str, dict[str, Any]],
    events: list[dict[str, Any]],
) -> None:
    for row in rows:
        asset_id = resolve_asset_id(row, assets, episode_to_asset)
        if asset_id is None:
            continue
        check = text(row.get("check"))
        flagged_frames = integer_or_none(row.get("flagged_frames")) or 0
        checked_frames = integer_or_none(row.get("checked_frames")) or 0
        clip_flag = boolish(row.get("clip_flag"))
        has_issue = flagged_frames > 0 or clip_flag is True
        if not has_issue:
            continue
        raw_invalid = check in {"keypoint_missing", "skeleton_quality_score"} and any(
            number(row.get(name), 0) > 0
            for name in ("nan_count", "inf_count", "invalid_point_count")
        )
        hard_check = check in {"text_integrity", "hdf5_text_integrity"} or raw_invalid
        status = "fail" if hard_check else "review"
        details[asset_id]["precheck_status"] = max_status(
            details[asset_id]["precheck_status"], status
        )
        metrics = {
            "check": check,
            "checked_frames": checked_frames,
            "flagged_frames": flagged_frames,
            "flagged_ratio": flagged_frames / checked_frames if checked_frames else None,
        }
        merge_metric(details[asset_id], "precheck", metrics)
        add_event(
            events,
            assets[asset_id],
            source_module="precheck",
            source_verdict=status,
            failure_mode=precheck_issue_type(check, raw_invalid),
            severity="high" if hard_check else "medium",
            reason=text(row.get("reason")) or f"{check} aggregate flagged",
            key_metrics=metrics,
            evidence_path=path,
        )


def add_candidate_evidence(
    rows: list[dict[str, Any]],
    path: Path,
    assets: dict[str, dict[str, Any]],
    episode_to_asset: dict[int, str],
    details: dict[str, dict[str, Any]],
    events: list[dict[str, Any]],
) -> None:
    for row in rows:
        asset_id = resolve_asset_id(row, assets, episode_to_asset)
        if asset_id is None:
            continue
        details[asset_id]["precheck_status"] = max_status(
            details[asset_id]["precheck_status"], "review"
        )
        metrics = parse_metrics(row.get("trigger_metrics"))
        add_event(
            events,
            assets[asset_id],
            source_module="precheck_candidate_window",
            source_verdict="review",
            failure_mode=joined_value(
                first_present(row, ("review_type", "trigger_reason"), "projection_review")
            ),
            severity=text(row.get("priority")) or "medium",
            start_frame=first_present(row, ("start_frame", "window_start_frame")),
            end_frame=first_present(row, ("end_frame", "window_end_frame")),
            representative_frame=first_present(
                row, ("peak_frame", "representative_frame")
            ),
            reason=joined_value(row.get("trigger_reason")) or "precheck candidate window",
            key_metrics=metrics,
            evidence_path=path,
        )


def add_video_quality_evidence(
    rows: list[dict[str, Any]],
    path: Path,
    assets: dict[str, dict[str, Any]],
    episode_to_asset: dict[int, str],
    details: dict[str, dict[str, Any]],
    events: list[dict[str, Any]],
) -> None:
    for raw_row in rows:
        row = normalize_video_quality_row(raw_row)
        asset_id = resolve_asset_id(row, assets, episode_to_asset)
        if asset_id is None:
            continue
        raw_decision = text(
            first_present(row, ("status", "decision", "final_status"), "pass")
        ).lower()
        if boolish(row.get("passed")) is False:
            raw_decision = "fail"
        if raw_decision in {"fail", "failed"}:
            status = "fail"
        elif raw_decision in {"warn", "warning", "risk", "pass_with_notes"}:
            status = "pass_with_notes"
        elif raw_decision == "review":
            status = "review"
        else:
            status = "pass"
        details[asset_id]["video_quality_raw_decision"] = raw_decision
        details[asset_id]["video_quality_status"] = status
        metrics = compact_dict(
            {
                "fps": row.get("fps"),
                "frame_count": row.get("frame_count"),
                "duration_sec": row.get("duration_sec"),
                "bad_frame_count": row.get("bad_frame_count"),
                "bad_duration_sec": row.get("bad_duration_sec"),
                text(row.get("metric_name")): row.get("metric_value"),
            }
        )
        merge_metric(details[asset_id], "video_quality", metrics)
        if status != "pass":
            add_event(
                events,
                assets[asset_id],
                source_module="video_quality",
                source_verdict=status,
                failure_mode=video_quality_issue_type(row),
                severity="high" if status == "fail" else "low",
                reason=text(first_present(row, ("reason", "notes"), raw_decision)),
                key_metrics=metrics,
                evidence_path=path,
            )


def add_sam3_evidence(
    rows: list[dict[str, Any]],
    path: Path,
    assets: dict[str, dict[str, Any]],
    episode_to_asset: dict[int, str],
    details: dict[str, dict[str, Any]],
    events: list[dict[str, Any]],
) -> None:
    for row in rows:
        asset_id = resolve_asset_id(row, assets, episode_to_asset)
        if asset_id is None:
            continue
        verdict = text(row.get("window_containment_verdict")) or "review"
        if verdict == "containment_fail":
            status, severity = "fail", "high"
        elif verdict in {
            "mixed_review",
            "projection_review",
            "review",
            "side_view_manual_review",
            "rotation_manual_review",
        }:
            status, severity = "review", "medium"
        elif verdict == "acceptable_flagged":
            status, severity = "pass_with_notes", "low"
        else:
            status, severity = "pass", "low"
        details[asset_id]["sam3_evidence_status"] = max_status(
            details[asset_id]["sam3_evidence_status"], status
        )
        metrics = compact_dict(
            {
                key: row.get(key)
                for key in (
                    "sampled_frame_count",
                    "strong_fail_frame_count",
                    "strong_fail_frame_ratio",
                    "inside_ratio_mean",
                    "projected_in_image_ratio_mean",
                    "projection_review_frame_count",
                )
            }
        )
        merge_metric(details[asset_id], "sam3_containment", metrics)
        if status != "pass":
            add_event(
                events,
                assets[asset_id],
                source_module="sam3_containment",
                source_verdict=status,
                failure_mode=sam3_issue_type(verdict),
                severity=severity,
                start_frame=first_present(
                    row, ("window_start_frame", "start_frame")
                ),
                end_frame=first_present(row, ("window_end_frame", "end_frame")),
                representative_frame=first_present(
                    row, ("representative_frame", "peak_frame")
                ),
                reason=text(row.get("reason")) or verdict,
                key_metrics=metrics,
                evidence_path=path,
            )


def add_manual_evidence(
    rows: list[dict[str, Any]],
    path: Path,
    assets: dict[str, dict[str, Any]],
    episode_to_asset: dict[int, str],
    details: dict[str, dict[str, Any]],
    events: list[dict[str, Any]],
) -> None:
    for row in rows:
        asset_id = resolve_asset_id(row, assets, episode_to_asset)
        if asset_id is None:
            continue
        outcome = text(
            first_present(row, ("manual_outcome", "algorithm_outcome", "label"), "review")
        )
        acceptance_status = text(row.get("acceptance_status"))
        if outcome in {"true_positive", "partial", "positive"} or acceptance_status == "rejected":
            status, severity = "fail", text(row.get("severity")) or "high"
        elif outcome in {"review", "uncertain", "needs_review"} or acceptance_status == "review":
            status, severity = "review", text(row.get("severity")) or "medium"
        elif outcome in {"false_positive", "acceptable_flagged"}:
            status, severity = "pass_with_notes", text(row.get("severity")) or "low"
        else:
            status, severity = "review", text(row.get("severity")) or "medium"
        details[asset_id]["manual_review_status"] = max_status(
            details[asset_id]["manual_review_status"], status
        )
        metrics = parse_metrics(row.get("key_metrics_json"))
        add_event(
            events,
            assets[asset_id],
            source_module="manual_review",
            source_verdict=status,
            manual_outcome=outcome,
            algorithm_outcome=text(row.get("algorithm_outcome")) or outcome,
            failure_mode=text(row.get("failure_mode")) or "unknown",
            severity=severity,
            confidence=text(row.get("confidence")),
            start_frame=first_present(
                row, ("affected_start_frame", "start", "start_frame")
            ),
            end_frame=first_present(row, ("affected_end_frame", "end", "end_frame")),
            representative_frame=row.get("representative_frame"),
            reason=text(first_present(row, ("comment", "note", "reason"), outcome)),
            key_metrics=metrics,
            evidence_path=path,
            reviewer=text(row.get("reviewer")),
        )


def build_asset_ledger(
    assets: dict[str, dict[str, Any]],
    details: dict[str, dict[str, Any]],
    events: list[dict[str, Any]],
    extra_evidence_paths: list[Path],
) -> list[dict[str, Any]]:
    by_asset: dict[str, list[dict[str, Any]]] = defaultdict(list)
    for event in events:
        by_asset[event["asset_id"]].append(event)
    rows = []
    for asset_id in sorted(assets):
        asset = assets[asset_id]
        detail = details[asset_id]
        asset_events = by_asset.get(asset_id, [])
        manual_outcomes = Counter(
            event["manual_outcome"]
            for event in asset_events
            if event["source_module"] == "manual_review" and event["manual_outcome"]
        )
        final_verdict = final_asset_verdict(detail)
        issue_counts = Counter(
            event["failure_mode"] for event in asset_events if event["failure_mode"]
        )
        evidence_paths = sorted(
            {
                event["evidence_path"]
                for event in asset_events
                if event.get("evidence_path")
            }
            | {str(path) for path in extra_evidence_paths}
        )
        rows.append(
            {
                "supplier_id": asset["supplier_id"],
                "asset_id": asset_id,
                "hdf5_path": asset["hdf5_path"],
                "video_path": asset["video_path"],
                "precheck_status": detail["precheck_status"],
                "video_quality_raw_decision": detail["video_quality_raw_decision"],
                "video_quality_status": detail["video_quality_status"],
                "sam3_evidence_status": detail["sam3_evidence_status"],
                "manual_review_status": detail["manual_review_status"],
                "final_verdict": final_verdict,
                "risk_level": risk_level(final_verdict, detail),
                "top_issue_types": "|".join(
                    name for name, _ in issue_counts.most_common(5)
                ),
                "true_positive_segment_count": manual_outcomes["true_positive"]
                + manual_outcomes["partial"]
                + manual_outcomes["positive"],
                "false_positive_count": manual_outcomes["false_positive"],
                "acceptable_flagged_count": manual_outcomes["acceptable_flagged"],
                "key_metrics_json": json_text(detail["metrics"]),
                "evidence_paths": "|".join(evidence_paths),
                "notes": asset_notes(asset_events),
            }
        )
    return rows


def final_asset_verdict(detail: dict[str, Any]) -> str:
    if (
        detail["precheck_status"] == "fail"
        or detail["video_quality_status"] == "fail"
        or detail["manual_review_status"] == "fail"
    ):
        return "fail"
    if (
        detail["video_quality_status"] == "review"
        or detail["sam3_evidence_status"] in {"fail", "review"}
        or detail["manual_review_status"] == "review"
    ):
        return "review"
    if (
        detail["video_quality_status"] == "pass_with_notes"
        or detail["sam3_evidence_status"] == "pass_with_notes"
        or detail["manual_review_status"] == "pass_with_notes"
    ):
        return "pass_with_notes"
    return "pass"


def risk_level(final_verdict: str, detail: dict[str, Any]) -> str:
    if final_verdict == "fail" or detail["sam3_evidence_status"] == "fail":
        return "high"
    if final_verdict == "review":
        return "medium"
    return "low"


def build_summary(
    ledger_rows: list[dict[str, Any]],
    events: list[dict[str, Any]],
    sam3_rows: list[dict[str, Any]],
) -> dict[str, Any]:
    manual_events = [
        event for event in events if event["source_module"] == "manual_review"
    ]
    return {
        "schema_version": "xjgt_acceptance_summary.v1",
        "asset_count": len(ledger_rows),
        "final_verdict_counts": sorted_counter(
            row["final_verdict"] for row in ledger_rows
        ),
        "risk_level_counts": sorted_counter(row["risk_level"] for row in ledger_rows),
        "video_quality_status_counts": sorted_counter(
            row["video_quality_status"] for row in ledger_rows
        ),
        "sam3_evidence_status_counts": sorted_counter(
            row["sam3_evidence_status"] for row in ledger_rows
        ),
        "manual_review_status_counts": sorted_counter(
            row["manual_review_status"] for row in ledger_rows
        ),
        "manual_outcome_counts": sorted_counter(
            event["manual_outcome"]
            for event in manual_events
            if event["manual_outcome"]
        ),
        "manual_failure_mode_counts": sorted_counter(
            event["failure_mode"] for event in manual_events
        ),
        "sam3_window_verdict_counts": sorted_counter(
            text(row.get("window_containment_verdict")) or "unknown"
            for row in sam3_rows
        ),
        "issue_event_count": len(events),
        "notes": PIPELINE_NOTES,
    }


def write_workbook(
    path: Path,
    ledger_rows: list[dict[str, Any]],
    events: list[dict[str, Any]],
    manual_rows: list[dict[str, Any]],
    sam3_rows: list[dict[str, Any]],
    video_quality_rows: list[dict[str, Any]],
    summary: dict[str, Any],
) -> None:
    workbook = Workbook()
    workbook.remove(workbook.active)
    supplier_rows = counter_table(
        "final_verdict", summary["final_verdict_counts"]
    ) + counter_table("risk_level", summary["risk_level_counts"])
    add_sheet(workbook, "supplier_summary", supplier_rows)
    add_sheet(workbook, "asset_ledger", ledger_rows, ASSET_LEDGER_COLUMNS)
    add_sheet(workbook, "issue_events", events, ISSUE_EVENT_COLUMNS)
    manual_summary = (
        counter_table("manual_outcome", summary["manual_outcome_counts"])
        + counter_table("failure_mode", summary["manual_failure_mode_counts"])
    )
    add_sheet(workbook, "manual_review_summary", manual_summary)
    add_sheet(
        workbook,
        "sam3_summary",
        counter_table("window_containment_verdict", summary["sam3_window_verdict_counts"]),
    )
    video_decisions = []
    for row in video_quality_rows:
        normalized = normalize_video_quality_row(row)
        video_decisions.append(
            text(first_present(normalized, ("status", "decision", "final_status"), "pass"))
        )
    add_sheet(
        workbook,
        "video_quality_summary",
        counter_table("raw_decision", sorted_counter(video_decisions)),
    )
    add_sheet(
        workbook,
        "method_thresholds",
        [
            {
                "threshold_name": name,
                "module": module,
                "condition": condition,
                "interpretation": interpretation,
                "calibration_status": "v0 heuristic pending calibration",
            }
            for name, module, condition, interpretation in METHOD_THRESHOLDS
        ],
    )
    add_sheet(
        workbook,
        "pipeline_notes",
        [{"note_id": index, "note": note} for index, note in enumerate(PIPELINE_NOTES, 1)],
    )
    path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(path)


def add_sheet(
    workbook: Workbook,
    name: str,
    rows: list[dict[str, Any]],
    columns: list[str] | None = None,
) -> None:
    sheet = workbook.create_sheet(name)
    columns = columns or (list(rows[0]) if rows else ["category", "value", "count"])
    sheet.append(columns)
    for row in rows:
        sheet.append([excel_value(row.get(column)) for column in columns])
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    sheet.sheet_view.showGridLines = False
    header_fill = PatternFill("solid", fgColor="1F4E78")
    for cell in sheet[1]:
        cell.fill = header_fill
        cell.font = Font(color="FFFFFF", bold=True)
        cell.alignment = Alignment(vertical="center")
    sheet.row_dimensions[1].height = 24
    for index, column in enumerate(columns, start=1):
        width = max(
            len(column),
            *(
                len(text(sheet.cell(row=row, column=index).value))
                for row in range(2, min(sheet.max_row, 101) + 1)
            ),
        )
        sheet.column_dimensions[get_column_letter(index)].width = min(max(width + 2, 12), 48)
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    if "final_verdict" in columns and sheet.max_row > 1:
        verdict_column = get_column_letter(columns.index("final_verdict") + 1)
        target = f"{verdict_column}2:{verdict_column}{sheet.max_row}"
        sheet.conditional_formatting.add(
            target,
            FormulaRule(
                formula=[f'{verdict_column}2="fail"'],
                fill=PatternFill("solid", fgColor="F4CCCC"),
            ),
        )
        sheet.conditional_formatting.add(
            target,
            FormulaRule(
                formula=[f'{verdict_column}2="review"'],
                fill=PatternFill("solid", fgColor="FFF2CC"),
            ),
        )


def add_event(
    events: list[dict[str, Any]],
    asset: dict[str, Any],
    *,
    source_module: str,
    source_verdict: str,
    failure_mode: str,
    severity: str,
    reason: str,
    key_metrics: dict[str, Any],
    evidence_path: Path,
    manual_outcome: str = "",
    algorithm_outcome: str = "",
    confidence: str = "",
    start_frame: Any = None,
    end_frame: Any = None,
    representative_frame: Any = None,
    reviewer: str = "",
) -> None:
    events.append(
        {
            "supplier_id": asset["supplier_id"],
            "asset_id": asset["asset_id"],
            "source_module": source_module,
            "source_verdict": source_verdict,
            "manual_outcome": manual_outcome,
            "algorithm_outcome": algorithm_outcome,
            "failure_mode": failure_mode,
            "severity": severity,
            "confidence": confidence,
            "start_frame": integer_or_none(start_frame),
            "end_frame": integer_or_none(end_frame),
            "representative_frame": integer_or_none(representative_frame),
            "reason": reason,
            "key_metrics_json": json_text(key_metrics),
            "evidence_path": str(evidence_path),
            "reviewer": reviewer,
        }
    )


def resolve_asset_id(
    row: dict[str, Any],
    assets: dict[str, dict[str, Any]],
    episode_to_asset: dict[int, str],
) -> str | None:
    value = first_present(row, ("asset_id", "content_id", "clip_id"))
    if value is not None:
        asset_id = normalize_asset_id(value)
        return asset_id if asset_id in assets else None
    episode_idx = integer_or_none(row.get("episode_idx"))
    return episode_to_asset.get(episode_idx) if episode_idx is not None else None


def precheck_issue_type(check: str, raw_invalid: bool) -> str:
    if raw_invalid:
        return "keypoint_raw_invalid"
    return {
        "text_integrity": "hdf5_text_invalid",
        "quality_score": "quality_hand_low",
        "keypoint_missing": "keypoint_low_quality_window",
        "keypoint_temporal": "temporal_jump",
        "keypoint_morphology": "keypoint_morphology_review",
    }.get(check, check or "precheck_review")


def sam3_issue_type(verdict: str) -> str:
    return {
        "containment_fail": "strong_containment_mismatch",
        "mixed_review": "occlusion_or_mask_undersegmentation",
    }.get(verdict, verdict)


def video_quality_issue_type(row: dict[str, Any]) -> str:
    reason = text(first_present(row, ("reason", "notes"), "video_quality"))
    lowered = reason.lower()
    for name in ("black_screen", "blur", "exposure", "stutter", "freeze"):
        if name in lowered:
            return f"video_{name}"
    return "video_quality"


def asset_notes(events: list[dict[str, Any]]) -> str:
    notes = []
    for event in events:
        if event["source_module"] == "manual_review":
            notes.append(
                f"manual:{event['manual_outcome']}:{event['failure_mode']}"
            )
    return "; ".join(notes[:8])


def merge_metric(detail: dict[str, Any], module: str, metrics: dict[str, Any]) -> None:
    if not metrics:
        return
    current = detail["metrics"].setdefault(module, [])
    current.append(metrics)


def parse_metrics(value: Any) -> dict[str, Any]:
    if isinstance(value, dict):
        return compact_dict(value)
    if isinstance(value, str) and value.strip():
        try:
            loaded = json.loads(value)
            return compact_dict(loaded) if isinstance(loaded, dict) else {"value": loaded}
        except json.JSONDecodeError:
            return {"value": value}
    return {}


def compact_dict(value: dict[str, Any]) -> dict[str, Any]:
    return {
        str(key): scalar(item)
        for key, item in value.items()
        if item is not None and str(key)
    }


def scalar(value: Any) -> Any:
    if hasattr(value, "item"):
        value = value.item()
    if isinstance(value, dict):
        return compact_dict(value)
    if isinstance(value, (list, tuple, set)):
        return [scalar(item) for item in value]
    if pd.isna(value) if not isinstance(value, (list, dict)) else False:
        return None
    return value


def write_csv(
    path: Path, columns: list[str], rows: Iterable[dict[str, Any]]
) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", newline="", encoding="utf-8") as handle:
        writer = csv.DictWriter(handle, fieldnames=columns)
        writer.writeheader()
        writer.writerows(rows)


def counter_table(dimension: str, counts: dict[str, int]) -> list[dict[str, Any]]:
    return [
        {"dimension": dimension, "value": value, "count": count}
        for value, count in sorted(counts.items())
    ]


def sorted_counter(values: Iterable[str]) -> dict[str, int]:
    return dict(sorted(Counter(value for value in values if value).items()))


def json_text(value: Any) -> str:
    return json.dumps(scalar(value), ensure_ascii=False, sort_keys=True)


def excel_value(value: Any) -> Any:
    if isinstance(value, (dict, list, tuple, set)):
        return json_text(value)
    return scalar(value)


def first_present(
    row: dict[str, Any], keys: tuple[str, ...], default: Any = None
) -> Any:
    for key in keys:
        value = row.get(key)
        if value is not None and value != "":
            return value
    return default


def max_status(current: str, candidate: str) -> str:
    return candidate if STATUS_RANK[candidate] > STATUS_RANK[current] else current


def normalize_asset_id(value: Any) -> str:
    value = scalar(value)
    text_value = str(value)
    return text_value[:-2] if text_value.endswith(".0") else text_value


def integer_or_none(value: Any) -> int | None:
    if value is None or value == "":
        return None
    try:
        return int(float(value))
    except (TypeError, ValueError):
        return None


def boolish(value: Any) -> bool | None:
    if isinstance(value, bool):
        return value
    if value is None or value == "":
        return None
    lowered = str(value).strip().lower()
    if lowered in {"1", "true", "yes", "pass"}:
        return True
    if lowered in {"0", "false", "no", "fail"}:
        return False
    return None


def number(value: Any, default: float = 0.0) -> float:
    try:
        return float(value)
    except (TypeError, ValueError):
        return default


def joined_value(value: Any) -> str:
    if isinstance(value, str):
        return value
    if isinstance(value, (list, tuple, set)):
        return "|".join(str(item) for item in value)
    return text(value)


def text(value: Any) -> str:
    return "" if value is None else str(value)


if __name__ == "__main__":
    raise SystemExit(main())
