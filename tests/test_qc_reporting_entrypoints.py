from __future__ import annotations

import json
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook

from tools.build_qc_json_projection import main as projection_main
from tests.test_qc_reporting_projection import _write_report


def test_projection_cli_prefers_qc_json_over_conflicting_legacy_sidecar(
    tmp_path: Path,
) -> None:
    archive = tmp_path / "quality_archive"
    _write_report(archive, "a", "acceptance", "pass", issues=[])
    sidecar = tmp_path / "candidate_windows.json"
    sidecar.write_text(
        json.dumps([{"asset_id": "a", "auto_verdict": "fail"}]),
        encoding="utf-8",
    )
    output_dir = tmp_path / "out"

    assert (
        projection_main(
            [
                "--quality-archive",
                str(archive),
                "--output-dir",
                str(output_dir),
                "--formats",
                "csv",
                "xlsx",
                "markdown",
                "--legacy-reconciliation-candidate-windows",
                str(sidecar),
            ]
        )
        == 0
    )

    ledger = pd.read_csv(output_dir / "assets.csv")
    assert ledger.loc[0, "overall_decision"] == "pass"
    reconciliation = pd.read_csv(output_dir / "reconciliation.csv")
    assert reconciliation.loc[0, "difference_type"] == "legacy_conflicts_with_qc_json"
    workbook = load_workbook(output_dir / "qc_projection.xlsx", read_only=True)
    assert workbook.sheetnames == [
        "Summary",
        "Assets",
        "Issues",
        "Execution",
        "Data_Dictionary",
    ]


def test_projection_cli_writes_parquet_tables_from_same_projection(
    tmp_path: Path,
) -> None:
    archive = tmp_path / "quality_archive"
    _write_report(archive, "a", "acceptance", "pass", issues=[])
    output_dir = tmp_path / "out"

    assert (
        projection_main(
            [
                "--quality-archive",
                str(archive),
                "--output-dir",
                str(output_dir),
                "--formats",
                "parquet",
            ]
        )
        == 0
    )

    parquet = pd.read_parquet(output_dir / "assets.parquet")
    assert parquet.loc[0, "overall_decision"] == "pass"
