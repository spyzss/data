from __future__ import annotations

import json
import sys
from pathlib import Path

import pandas as pd
import pytest
import yaml
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Border, Font, PatternFill, Side

from tools.build_acceptance_ledger import (
    GENERATED_SHEETS,
    build_acceptance_ledger,
    build_supplier_ledger,
    merge_intervals,
)
import tools.build_acceptance_ledger as acceptance_ledger_module
import tools.acceptance_ledger_weekly as weekly_ledger
from tests.qc_report_fixtures import make_v2_report
WEEKLY_SHEETS = [
    "五供应商总览",
    "星际归途",
    "DeepReach",
    "京东JDT",
    "供应商4",
    "供应商5",
    "人工问题与阈值",
]

WEEKLY_DETAIL_COLUMNS = [
    "asset_id",
    "total_frames",
    "text_check_status",
    "skeleton_static_status",
    "video_quality_status",
    "abnormal_frame_status",
    "final_acceptance_status",
    "precheck_window_count",
    "precheck_fail_window_count",
    "precheck_to_sam3_window_count",
    "precheck_to_sam3_ratio",
    "sam3_processed_window_count",
    "sam3_fail_window_count",
    "sam3_to_manual_window_count",
    "sam3_to_manual_ratio",
    "manual_submitted_window_count",
    "manual_reviewed_window_count",
    "manual_fail_window_count",
    "manual_pass_window_count",
    "manual_pending_window_count",
    "manual_completion_ratio",
    "final_fail_window_count",
    "final_review_window_count",
    "main_reason",
    "evidence_path",
]

REAL_WEEKLY_CONFIG = (
    Path(__file__).resolve().parents[1]
    / "configs"
    / "acceptance_ledger_xjgt_jdt_dr_weekly.yaml"
)
WEEKLY_SUPPLIER_CONFIG_KEYS = {
    "supplier_id",
    "manifest",
    "ledger_asset_ledger",
    "precheck_check_results",
    "precheck_clip_aggregates",
    "candidate_windows",
    "video_quality",
    "sam3_summary",
    "review_queue",
    "manual_labels",
    "blocker",
    "blocked",
    "blocked_policy",
    "text_required",
}


def test_real_xjgt_jdt_deepreach_weekly_config_loads() -> None:
    config = yaml.safe_load(REAL_WEEKLY_CONFIG.read_text(encoding="utf-8"))

    assert config["workbook_mode"] == "weekly_template"
    suppliers = {
        supplier["supplier_id"]: supplier for supplier in config["suppliers"]
    }
    assert set(suppliers) == {"xjgt", "jdt", "deepreach"}
    assert all(
        set(supplier) <= WEEKLY_SUPPLIER_CONFIG_KEYS
        for supplier in suppliers.values()
    )
    assert suppliers["xjgt"]["manifest"] == (
        "outputs/acceptance_5x100/xjgt/ledger/xjgt_video_frame_counts.csv"
    )
    assert suppliers["xjgt"]["ledger_asset_ledger"] == (
        "outputs/acceptance_5x100/xjgt/ledger/xjgt_100_asset_ledger.csv"
    )
    assert suppliers["xjgt"]["precheck_check_results"] == (
        "outputs/acceptance_5x100/xjgt/precheck/check_results.parquet"
    )
    assert suppliers["xjgt"]["video_quality"] == (
        "outputs/acceptance_5x100/xjgt/video_quality/"
        "video_quality_acceptance_summary.csv"
    )
    assert suppliers["xjgt"]["sam3_summary"] == (
        "outputs/acceptance_5x100/xjgt/sam3_containment/"
        "window_keypoint_containment_summary.parquet"
    )
    assert suppliers["xjgt"]["review_queue"] == (
        "outputs/acceptance_5x100/xjgt/review/review_queue.csv"
    )
    assert suppliers["xjgt"]["manual_labels"] == (
        "outputs/acceptance_5x100/xjgt/manual_review/"
        "manual_labels_autosave.normalized.csv"
    )
    assert suppliers["jdt"]["manifest"].endswith(
        "/jdt/manifest/jdt_sample_100_manifest.csv"
    )
    assert suppliers["jdt"]["precheck_check_results"].endswith(
        "/jdt/precheck/check_results.parquet"
    )
    assert suppliers["jdt"]["video_quality"].endswith(
        "/jdt/video_quality/video_quality_results.json"
    )
    assert suppliers["jdt"]["sam3_summary"].endswith(
        "/jdt/sam3_containment/window_keypoint_containment_summary.parquet"
    )
    assert suppliers["jdt"]["review_queue"].endswith(
        "/jdt/review/review_queue.csv"
    )
    assert suppliers["jdt"]["manual_labels"].endswith(
        "/jdt/manual_review/manual_labels.csv"
    )
    assert suppliers["deepreach"]["manifest"].endswith(
        "/deepreach/manifest/deepreach_sample_100_manifest.csv"
    )
    assert suppliers["deepreach"]["precheck_check_results"].endswith(
        "/deepreach/precheck/check_results.parquet"
    )
    assert suppliers["deepreach"]["video_quality"].endswith(
        "/deepreach/video_quality/video_quality_results.json"
    )
    assert suppliers["deepreach"]["blocked"] is True
    assert suppliers["deepreach"]["blocked_policy"] == "candidate_windows_review"
    assert suppliers["deepreach"]["blocker"].endswith(
        "/deepreach/blockers/sam3_containment_blocker.json"
    )


@pytest.mark.parametrize(
    ("row", "expected"),
    [
        ({"manual_outcome": " true_positive "}, "fail"),
        ({"manual_outcome": "positive"}, "fail"),
        ({"acceptance_status": "rejected"}, "fail"),
        ({"manual_outcome": "false_positive"}, "pass"),
        ({"manual_outcome": "acceptable_flagged"}, "pass"),
        ({"acceptance_status": "accepted"}, "pass"),
        ({"manual_outcome": "review"}, "review"),
    ],
)
def test_generic_and_weekly_manual_outcome_normalization_agree(
    row: dict[str, str], expected: str
) -> None:
    assert acceptance_ledger_module.normalize_manual_outcome(row) == expected
    assert weekly_ledger._manual_window_kind([row]) == expected


def test_weekly_review_id_normalization_and_exact_window_fallback() -> None:
    result = weekly_ledger._resolve_windows(
        asset_id="asset",
        total_frames=10,
        precheck_row={"temporal_seen": True},
        precheck_source_status="readable",
        candidate_rows=[
            {
                "review_id": "candidate-a",
                "asset_id": "asset",
                "window_start_frame": 0,
                "window_end_frame": 2,
            },
            {
                "review_id": "candidate-b",
                "asset_id": "asset",
                "window_start_frame": 3,
                "window_end_frame": 5,
            },
        ],
        sam3_rows=[
            {
                "asset_id": "asset",
                "window_start_frame": 0,
                "window_end_frame": 2,
                "window_containment_verdict": "mixed_review",
            },
            {
                "asset_id": "asset",
                "window_start_frame": 3,
                "window_end_frame": 5,
                "window_containment_verdict": "mixed_review",
            },
        ],
        review_rows=[
            {
                "review_id": " 42 ",
                "asset_id": "asset",
                "window_start_frame": 0,
                "window_end_frame": 2,
            },
            {
                "review_id": "queue-b",
                "asset_id": "asset",
                "window_start_frame": 3,
                "window_end_frame": 5,
            },
        ],
        manual_rows=[
            {
                "review_id": 42,
                "asset_id": "asset",
                "window_start_frame": 0,
                "window_end_frame": 2,
                "manual_outcome": "false_positive",
            },
            {
                "review_id": "manual-b",
                "asset_id": "asset",
                "window_start_frame": 3,
                "window_end_frame": 5,
                "manual_outcome": "true_positive",
            },
        ],
        sam3_blocked=False,
        blocker_reason="",
    )

    assert result["manual_submitted_window_count"] == 2
    assert result["manual_reviewed_window_count"] == 2
    assert result["manual_pending_window_count"] == 0
    assert result["manual_pass_window_count"] == 1
    assert result["manual_fail_window_count"] == 1
    assert result["final_review_window_count"] == 0


def test_manual_pending_is_sam3_arbitration_minus_completed_manual() -> None:
    result = weekly_ledger._resolve_windows(
        asset_id="asset",
        total_frames=10,
        precheck_row={"temporal_seen": True},
        precheck_source_status="readable",
        candidate_rows=[
            {
                "asset_id": "asset",
                "window_start_frame": 1,
                "window_end_frame": 4,
            }
        ],
        sam3_rows=[
            {
                "asset_id": "asset",
                "window_start_frame": 1,
                "window_end_frame": 4,
                "window_containment_verdict": "mixed_review",
            }
        ],
        review_rows=[],
        manual_rows=[],
        sam3_blocked=False,
        blocker_reason="",
    )

    assert result["sam3_to_manual_window_count"] == 1
    assert result["manual_submitted_window_count"] == 0
    assert result["manual_reviewed_window_count"] == 0
    assert result["manual_pending_window_count"] == 1


def test_incomplete_manual_row_remains_unresolved_in_generic_ledger(
    tmp_path: Path,
) -> None:
    manifest = tmp_path / "manifest.csv"
    review_queue = tmp_path / "review_queue.csv"
    manual_labels = tmp_path / "manual_labels.csv"
    pd.DataFrame(
        [{"asset_id": "asset", "start_frame": 0, "end_frame": 9}]
    ).to_csv(manifest, index=False)
    pd.DataFrame(
        [
            {
                "review_id": "review-1",
                "asset_id": "asset",
                "window_start_frame": 2,
                "window_end_frame": 4,
            }
        ]
    ).to_csv(review_queue, index=False)
    pd.DataFrame(
        [
            {
                "review_id": "review-1",
                "asset_id": "asset",
                "window_start_frame": 2,
                "window_end_frame": 4,
                "manual_outcome": "review",
            }
        ]
    ).to_csv(manual_labels, index=False)

    result = build_supplier_ledger(
        {
            "supplier_id": "supplier",
            "manifest": str(manifest),
            "review_queue": str(review_queue),
            "manual_labels": str(manual_labels),
            "required_inputs": ["manifest"],
        },
        config_dir=tmp_path,
    )

    assert result.asset_rows[0]["final_status"] == "review"
    assert result.asset_rows[0]["unresolved_review_window_count"] == 1


def test_missing_optional_morphology_does_not_block_skeleton() -> None:
    status, reasons = weekly_ledger._skeleton_status(
        source_status="readable",
        precheck_row={
            "checks": {"skeleton_quality_score", "keypoint_temporal"},
            "missing_intervals": [],
        },
    )
    assert status == "pass"
    assert "optional keypoint_morphology not executed" in reasons


def test_applicable_text_check_skipped_is_not_run() -> None:
    status, _ = weekly_ledger._text_status(
        config={"text_required": True},
        source_status="readable",
        precheck_row={"checks": {"skeleton_quality_score"}},
    )
    assert status == "not_run"


def _canonical_weekly_header() -> list[str]:
    return WEEKLY_DETAIL_COLUMNS


def _write_weekly_template(tmp_path: Path) -> Path:
    path = tmp_path / "weekly_template.xlsx"
    workbook = Workbook()
    workbook.remove(workbook.active)
    old_names = [
        "五供应商总览",
        "星际归途",
        "DeepReach",
        "供应商3",
        "供应商4",
        "供应商5",
        "人工问题与阈值",
    ]
    header = _canonical_weekly_header()
    border = Border(bottom=Side(style="thin", color="44546A"))
    for name in old_names:
        sheet = workbook.create_sheet(name)
        if name in {"星际归途", "DeepReach", "供应商3", "供应商4", "供应商5"}:
            sheet.append(["旧模板表头"])
            sheet.append(header)
            for cell in sheet[1]:
                cell.fill = PatternFill("solid", fgColor="B4C6E7")
                cell.font = Font(bold=True, color="1F1F1F")
                cell.border = border
            sheet.freeze_panes = "A3"
            sheet.auto_filter.ref = sheet.dimensions
            sheet.row_dimensions[1].height = 37
            sheet.column_dimensions["A"].width = 25
    workbook["供应商4"]["A3"] = "placeholder-four"
    workbook["供应商5"]["A3"] = "placeholder-five"
    workbook["人工问题与阈值"]["A1"] = (
        "定向复核不是无偏全局召回率估计"
    )
    workbook.save(path)
    return path


def _pass_check_rows(asset_ids: list[str]) -> list[dict[str, object]]:
    rows: list[dict[str, object]] = []
    for asset_id in asset_ids:
        rows.extend(
            [
                {
                    "asset_id": asset_id,
                    "check": "text_integrity",
                    "frame_idx": -1,
                    "flag": False,
                    "metrics": json.dumps({"missing_field_count": 0}),
                },
                {
                    "asset_id": asset_id,
                    "check": "skeleton_quality_score",
                    "frame_idx": -1,
                    "flag": False,
                    "metrics": "{}",
                },
                {
                    "asset_id": asset_id,
                    "check": "keypoint_morphology",
                    "frame_idx": -1,
                    "flag": False,
                    "metrics": json.dumps({"morphology_verdict": "pass"}),
                },
                {
                    "asset_id": asset_id,
                    "check": "keypoint_temporal",
                    "frame_idx": -1,
                    "flag": None,
                    "metrics": "{}",
                },
            ]
        )
    return rows


def _write_weekly_supplier_inputs(tmp_path: Path) -> tuple[Path, Path, Path]:
    xjgt_dir = tmp_path / "xjgt"
    xjgt_dir.mkdir()
    pd.DataFrame(
        [{"asset_id": "xjgt-000", "frame_count": 10}]
    ).to_csv(xjgt_dir / "manifest.csv", index=False)
    pd.DataFrame(_pass_check_rows(["xjgt-000"])).to_parquet(
        xjgt_dir / "check_results.parquet", index=False
    )
    pd.DataFrame([{"asset_id": "xjgt-000", "status": "pass"}]).to_json(
        xjgt_dir / "video_quality.json", orient="records"
    )
    xjgt_candidate = {
        "review_id": "xjgt-review-0",
        "asset_id": "xjgt-000",
        "window_start_frame": 0,
        "window_end_frame": 4,
        "module": "precheck",
        "source_verdict": "review",
    }
    pd.DataFrame([xjgt_candidate]).to_parquet(
        xjgt_dir / "candidate_windows.parquet", index=False
    )
    pd.DataFrame([xjgt_candidate]).to_csv(
        xjgt_dir / "review_queue.csv", index=False
    )

    jdt_dir = tmp_path / "jdt"
    jdt_dir.mkdir()
    asset_ids = [f"jdt-{index:03d}" for index in range(100)]
    pd.DataFrame(
        [
            {
                "asset_id": asset_id,
                "start_frame": 0,
                "end_frame": 9,
                "primary_video_path": f"/{asset_id}.mp4",
            }
            for asset_id in asset_ids
        ]
    ).to_csv(jdt_dir / "manifest.csv", index=False)
    check_rows = _pass_check_rows(asset_ids)
    pd.DataFrame(check_rows).to_parquet(jdt_dir / "check_results.parquet", index=False)
    pd.DataFrame(
        [{"asset_id": asset_id, "check_count": 3} for asset_id in asset_ids]
    ).to_parquet(jdt_dir / "clip_aggregates.parquet", index=False)
    pd.DataFrame(
        [{"asset_id": asset_id, "status": "pass"} for asset_id in asset_ids]
    ).to_json(jdt_dir / "video_quality.json", orient="records")
    candidate_rows = [
        {
            "review_id": f"review-{index}",
            "asset_id": asset_ids[index],
            "window_start_frame": 0,
            "window_end_frame": 4,
            "source_level": "window",
            "module": "precheck",
        }
        for index in range(12)
    ]
    pd.DataFrame(candidate_rows).to_parquet(
        jdt_dir / "candidate_windows.parquet", index=False
    )
    pd.DataFrame(candidate_rows[:11]).to_csv(
        jdt_dir / "review_queue.csv", index=False
    )
    pd.DataFrame(
        [
            {
                "asset_id": asset_ids[index],
                "window_start_frame": 0,
                "window_end_frame": 4,
                "window_containment_verdict": (
                    "mixed_review" if index < 11 else "containment_pass"
                ),
            }
            for index in range(12)
        ]
    ).to_parquet(jdt_dir / "sam3.parquet", index=False)
    pd.DataFrame(
        [
            {
                "review_id": "review-0",
                "asset_id": asset_ids[0],
                "window_start_frame": 0,
                "window_end_frame": 4,
                "affected_start_frame": 1,
                "affected_end_frame": 2,
                "manual_outcome": "true_positive",
                "acceptance_status": "rejected",
                "failure_mode": "severe_keypoint_offset",
            },
            {
                "review_id": "review-1",
                "asset_id": asset_ids[1],
                "window_start_frame": 0,
                "window_end_frame": 4,
                "affected_start_frame": 1,
                "affected_end_frame": 2,
                "manual_outcome": "true_positive",
                "acceptance_status": "rejected",
                "failure_mode": "severe_keypoint_offset",
            },
            *[
                {
                    "review_id": f"review-{index}",
                    "asset_id": asset_ids[index],
                    "window_start_frame": 0,
                    "window_end_frame": 4,
                    "affected_start_frame": "",
                    "affected_end_frame": "",
                    "manual_outcome": "false_positive",
                    "acceptance_status": "accepted",
                    "failure_mode": "unknown",
                }
                for index in range(2, 11)
            ],
        ]
    ).to_csv(jdt_dir / "manual_labels.csv", index=False)

    dr_dir = tmp_path / "deepreach"
    dr_dir.mkdir()
    pd.DataFrame(
        [
            {"asset_id": "dr-000", "start_frame": 0, "end_frame": 9},
            {"asset_id": "dr-001", "start_frame": 0, "end_frame": 19},
        ]
    ).to_csv(dr_dir / "manifest.csv", index=False)
    dr_checks = _pass_check_rows(["dr-000", "dr-001"])
    dr_checks.append(
        {
            "asset_id": "dr-001",
            "check": "skeleton_quality_score",
            "frame_idx": 3,
            "flag": True,
            "metrics": json.dumps(
                {
                    "keypoint_presence_invalid": 1,
                    "valid_keypoint_count_left": 20,
                    "missing_keypoint_count_left": 1,
                    "valid_keypoint_count_right": 21,
                    "missing_keypoint_count_right": 0,
                }
            ),
        }
    )
    pd.DataFrame(dr_checks).to_parquet(
        dr_dir / "check_results.parquet", index=False
    )
    pd.DataFrame(
        [
            {"asset_id": "dr-000", "status": "pass"},
            {"asset_id": "dr-001", "status": "pass"},
        ]
    ).to_json(dr_dir / "video_quality.json", orient="records")
    pd.DataFrame(
        [
            {
                "review_id": "dr-review-0",
                "asset_id": "dr-000",
                "window_start_frame": 1,
                "window_end_frame": 4,
                "module": "precheck",
                "source_verdict": "review",
            }
        ]
    ).to_parquet(dr_dir / "candidate_windows.parquet", index=False)
    (dr_dir / "blocker.json").write_text(
        json.dumps(
            {
                "status": "blocked",
                "reason": "missing_head_calibration_lineage",
            }
        ),
        encoding="utf-8",
    )
    return xjgt_dir, jdt_dir, dr_dir


def _weekly_config(tmp_path: Path) -> Path:
    xjgt_dir, jdt_dir, dr_dir = _write_weekly_supplier_inputs(tmp_path)
    config = tmp_path / "weekly_ledger.yaml"
    config.write_text(
        yaml.safe_dump(
            {
                "workbook_mode": "weekly_template",
                "output_label": "weekly_test",
                "suppliers": [
                    {
                        "supplier_id": "xjgt",
                        "manifest": str(xjgt_dir / "manifest.csv"),
                        "precheck_check_results": str(
                            xjgt_dir / "check_results.parquet"
                        ),
                        "candidate_windows": str(
                            xjgt_dir / "candidate_windows.parquet"
                        ),
                        "video_quality": str(xjgt_dir / "video_quality.json"),
                        "review_queue": str(xjgt_dir / "review_queue.csv"),
                        "text_required": True,
                    },
                    {
                        "supplier_id": "jdt",
                        "manifest": str(jdt_dir / "manifest.csv"),
                        "precheck_check_results": str(
                            jdt_dir / "check_results.parquet"
                        ),
                        "precheck_clip_aggregates": str(
                            jdt_dir / "clip_aggregates.parquet"
                        ),
                        "candidate_windows": str(
                            jdt_dir / "candidate_windows.parquet"
                        ),
                        "video_quality": str(jdt_dir / "video_quality.json"),
                        "sam3_summary": str(jdt_dir / "sam3.parquet"),
                        "review_queue": str(jdt_dir / "review_queue.csv"),
                        "manual_labels": str(jdt_dir / "manual_labels.csv"),
                        "text_required": False,
                        "required_inputs": ["manifest"],
                    },
                    {
                        "supplier_id": "deepreach",
                        "manifest": str(dr_dir / "manifest.csv"),
                        "precheck_check_results": str(
                            dr_dir / "check_results.parquet"
                        ),
                        "candidate_windows": str(
                            dr_dir / "candidate_windows.parquet"
                        ),
                        "video_quality": str(dr_dir / "video_quality.json"),
                        "blocker": str(dr_dir / "blocker.json"),
                        "blocked": True,
                        "blocked_policy": "candidate_windows_review",
                        "text_required": False,
                        "required_inputs": ["manifest"],
                    },
                ],
            },
            sort_keys=False,
        ),
        encoding="utf-8",
    )
    return config


def _write_jdt_inputs(tmp_path: Path) -> dict[str, Path]:
    manifest = tmp_path / "jdt_manifest.csv"
    pd.DataFrame(
        [
            {"asset_id": "asset-a", "start_frame": 0, "end_frame": 99},
            {"asset_id": "asset-b", "start_frame": 0, "end_frame": 49},
            {"asset_id": "asset-c", "start_frame": 0, "end_frame": 29},
            {"asset_id": "asset-d", "start_frame": 0, "end_frame": 19},
        ]
    ).to_csv(manifest, index=False)

    manual_labels = tmp_path / "manual_labels.csv"
    pd.DataFrame(
        [
            {
                "review_id": "review-a",
                "asset_id": "asset-a",
                "window_start_frame": 0,
                "window_end_frame": 29,
                "affected_start_frame": 10,
                "affected_end_frame": 19,
                "manual_outcome": "true_positive",
                "acceptance_status": "rejected",
            },
            {
                "review_id": "review-a",
                "asset_id": "asset-a",
                "window_start_frame": 0,
                "window_end_frame": 29,
                "affected_start_frame": 15,
                "affected_end_frame": 25,
                "manual_outcome": "true_positive",
                "acceptance_status": "rejected",
            },
            {
                "review_id": "review-b",
                "asset_id": "asset-b",
                "window_start_frame": 5,
                "window_end_frame": 9,
                "affected_start_frame": "",
                "affected_end_frame": "",
                "manual_outcome": "true_positive",
                "acceptance_status": "rejected",
            },
        ]
    ).to_csv(manual_labels, index=False)

    review_queue = tmp_path / "review_queue.csv"
    pd.DataFrame(
        [
            {
                "review_id": "review-a",
                "asset_id": "asset-a",
                "window_start_frame": 0,
                "window_end_frame": 29,
            },
            {
                "review_id": "review-a-unresolved",
                "asset_id": "asset-a",
                "window_start_frame": 50,
                "window_end_frame": 59,
            },
            {
                "review_id": "review-b",
                "asset_id": "asset-b",
                "window_start_frame": 5,
                "window_end_frame": 9,
            },
            {
                "review_id": "review-c",
                "asset_id": "asset-c",
                "window_start_frame": 0,
                "window_end_frame": 9,
            },
        ]
    ).to_csv(review_queue, index=False)

    sam3_summary = tmp_path / "sam3_summary.parquet"
    pd.DataFrame(
        [
            {"asset_id": "asset-a", "window_containment_verdict": "review"},
            {"asset_id": "asset-c", "window_containment_verdict": "review"},
        ]
    ).to_parquet(sam3_summary, index=False)
    return {
        "manifest": manifest,
        "manual_labels": manual_labels,
        "review_queue": review_queue,
        "sam3_summary": sam3_summary,
    }


def _write_deepreach_inputs(tmp_path: Path) -> dict[str, Path]:
    manifest = tmp_path / "deepreach_manifest.csv"
    pd.DataFrame(
        [
            {"asset_id": "dr-a", "frame_count": 10},
            {"asset_id": "dr-b", "frame_count": 20},
        ]
    ).to_csv(manifest, index=False)
    blocker = tmp_path / "sam3_containment_blocker.json"
    blocker.write_text(
        json.dumps(
            {
                "status": "blocked",
                "reason": "head calibration lineage missing",
            }
        ),
        encoding="utf-8",
    )
    return {"manifest": manifest, "blocker": blocker}


def _config(tmp_path: Path) -> Path:
    jdt = _write_jdt_inputs(tmp_path)
    deepreach = _write_deepreach_inputs(tmp_path)
    config_path = tmp_path / "ledger.yaml"
    config_path.write_text(
        yaml.safe_dump(
            {
                "output_label": "test_jdt_dr",
                "suppliers": [
                    {
                        "supplier_id": "jdt",
                        **{key: str(value) for key, value in jdt.items()},
                        "blocked": False,
                        "required_inputs": ["manifest"],
                    },
                    {
                        "supplier_id": "deepreach",
                        **{key: str(value) for key, value in deepreach.items()},
                        "blocked": True,
                        "blocked_policy": "all_frames_review",
                        "required_inputs": ["manifest"],
                    },
                ],
            },
            sort_keys=False,
        ),
        encoding="utf-8",
    )
    return config_path


def test_merge_intervals_unions_overlapping_inclusive_ranges() -> None:
    assert merge_intervals([(10, 19), (15, 25), (30, 30)]) == [
        (10, 25),
        (30, 30),
    ]


def test_jdt_manual_labels_compute_fail_review_pass_and_priority(tmp_path: Path) -> None:
    inputs = _write_jdt_inputs(tmp_path)
    result = build_supplier_ledger(
        {
            "supplier_id": "jdt",
            **{key: str(value) for key, value in inputs.items()},
            "blocked": False,
            "required_inputs": ["manifest"],
        },
        config_dir=tmp_path,
    )
    assets = {row["asset_id"]: row for row in result.asset_rows}

    assert assets["asset-a"]["fail_frame_count"] == 16
    assert assets["asset-a"]["review_frame_count"] == 10
    assert assets["asset-a"]["pass_frame_count"] == 74
    assert assets["asset-a"]["final_status"] == "fail"
    assert assets["asset-b"]["fail_frame_count"] == 5
    assert assets["asset-b"]["final_status"] == "fail"
    assert assets["asset-c"]["review_frame_count"] == 10
    assert assets["asset-c"]["final_status"] == "review"
    assert assets["asset-d"]["final_status"] == "pass"

    overview = result.overview
    assert overview["total_clip_count"] == 4
    assert overview["total_frame_count"] == 200
    assert overview["fail_frame_count"] == 21
    assert overview["review_frame_count"] == 20
    assert overview["pass_frame_count"] == 159
    assert overview["fail_clip_count"] == 2
    assert overview["review_clip_count"] == 1
    assert overview["pass_clip_count"] == 1
    assert overview["manual_review_count"] == 2
    assert overview["manual_review_window_count"] == 2
    assert overview["manual_review_asset_count"] == 2
    assert overview["manual_review_label_rows"] == 3
    assert overview["sam3_status"] == "evaluated"
    assert "blocker=not_evaluated" not in overview["note"]


def test_manual_review_counts_windows_assets_and_rows_independently(
    tmp_path: Path,
) -> None:
    manifest = tmp_path / "manifest.csv"
    pd.DataFrame([{"asset_id": "asset-a", "frame_count": 20}]).to_csv(
        manifest, index=False
    )
    manual_labels = tmp_path / "manual_labels.csv"
    pd.DataFrame(
        [
            {
                "review_id": "review-001",
                "asset_id": "asset-a",
                "affected_start_frame": 1,
                "affected_end_frame": 2,
                "manual_outcome": "true_positive",
            },
            {
                "review_id": "review-002",
                "asset_id": "asset-a",
                "affected_start_frame": 10,
                "affected_end_frame": 11,
                "manual_outcome": "true_positive",
            },
        ]
    ).to_csv(manual_labels, index=False)

    result = build_supplier_ledger(
        {
            "supplier_id": "jdt",
            "manifest": str(manifest),
            "manual_labels": str(manual_labels),
            "required_inputs": ["manifest"],
        },
        config_dir=tmp_path,
    )

    assert result.overview["manual_review_window_count"] == 2
    assert result.overview["manual_review_asset_count"] == 1
    assert result.overview["manual_review_label_rows"] == 2
    assert result.overview["manual_review_count"] == 2


def test_blocked_all_frames_review_marks_every_clip_and_frame_review(
    tmp_path: Path,
) -> None:
    inputs = _write_deepreach_inputs(tmp_path)
    result = build_supplier_ledger(
        {
            "supplier_id": "deepreach",
            **{key: str(value) for key, value in inputs.items()},
            "blocked": True,
            "blocked_policy": "all_frames_review",
            "required_inputs": ["manifest"],
        },
        config_dir=tmp_path,
    )

    assert result.overview["total_clip_count"] == 2
    assert result.overview["total_frame_count"] == 30
    assert result.overview["review_frame_count"] == 30
    assert result.overview["pass_frame_count"] == 0
    assert result.overview["fail_frame_count"] == 0
    assert result.overview["review_clip_count"] == 2
    assert result.overview["sam3_status"] == "blocked"
    assert "head calibration lineage missing" in result.overview["note"]
    assert all(row["final_status"] == "review" for row in result.asset_rows)


def test_configured_missing_blocker_is_explicit_but_unconfigured_is_silent(
    tmp_path: Path,
) -> None:
    manifest = tmp_path / "manifest.csv"
    pd.DataFrame([{"asset_id": "asset", "frame_count": 5}]).to_csv(
        manifest, index=False
    )
    base = {
        "supplier_id": "supplier",
        "manifest": str(manifest),
        "required_inputs": ["manifest"],
    }

    without_blocker = build_supplier_ledger(base, config_dir=tmp_path)
    with_missing_blocker = build_supplier_ledger(
        {**base, "blocker": str(tmp_path / "missing_blocker.json")},
        config_dir=tmp_path,
    )

    assert "blocker=not_evaluated" not in without_blocker.overview["note"]
    assert "blocker=not_evaluated" in with_missing_blocker.overview["note"]


def test_overview_totals_balance_and_workbook_contains_generated_sheets(
    tmp_path: Path,
) -> None:
    output = tmp_path / "ledger.xlsx"
    build_acceptance_ledger(
        config_path=_config(tmp_path),
        output_path=output,
        overwrite=True,
    )

    workbook = load_workbook(output, read_only=True)
    assert workbook.sheetnames == list(GENERATED_SHEETS)
    overview = pd.read_excel(output, sheet_name="Overview")
    assert (
        overview["pass_frame_count"]
        + overview["fail_frame_count"]
        + overview["review_frame_count"]
    ).equals(overview["total_frame_count"])
    assert (
        overview["pass_clip_count"]
        + overview["fail_clip_count"]
        + overview["review_clip_count"]
    ).equals(overview["total_clip_count"])


def test_formal_cli_does_not_require_legacy_config(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    archive = tmp_path / "quality_archive"
    report = make_v2_report(status="completed", overall_decision="pass")
    report["asset_id"] = "asset-a"
    report["execution"]["profile"] = "acceptance"
    report["execution"]["module_states"] = {
        "hdf5_text_info": {"state": "completed"},
    }
    archive.mkdir()
    (archive / "asset-a.json").write_text(json.dumps(report), encoding="utf-8")
    output = tmp_path / "ledger.xlsx"
    monkeypatch.setattr(
        sys,
        "argv",
        [
            "build_acceptance_ledger.py",
            "--quality-archive",
            str(archive),
            "--output",
            str(output),
            "--overwrite",
        ],
    )

    assert acceptance_ledger_module.main() == 0
    assert output.exists()


def test_legacy_build_requires_config_explicitly(tmp_path: Path) -> None:
    with pytest.raises(ValueError, match="config_path is required"):
        build_acceptance_ledger(config_path=None, output_path=tmp_path / "ledger.xlsx")


def test_existing_workbook_preserves_unrelated_sheets_and_replaces_generated(
    tmp_path: Path,
) -> None:
    existing = tmp_path / "existing.xlsx"
    workbook = Workbook()
    notes = workbook.active
    notes.title = "Keep_Me"
    notes["A1"] = "preserved"
    old_overview = workbook.create_sheet("Overview")
    old_overview["A1"] = "stale"
    workbook.save(existing)
    output = tmp_path / "updated.xlsx"

    build_acceptance_ledger(
        config_path=_config(tmp_path),
        output_path=output,
        existing_workbook=existing,
        overwrite=True,
    )

    updated = load_workbook(output)
    assert updated["Keep_Me"]["A1"].value == "preserved"
    assert updated["Overview"]["A1"].value == "supplier_id"
    assert updated.sheetnames == ["Keep_Me", *GENERATED_SHEETS]


def test_missing_required_manifest_raises_but_optional_modules_do_not(
    tmp_path: Path,
) -> None:
    with pytest.raises(FileNotFoundError, match="required input manifest"):
        build_supplier_ledger(
            {
                "supplier_id": "missing",
                "manifest": str(tmp_path / "missing.csv"),
                "required_inputs": ["manifest"],
            },
            config_dir=tmp_path,
        )

    manifest = tmp_path / "manifest.csv"
    pd.DataFrame([{"asset_id": "asset", "frame_count": 5}]).to_csv(
        manifest, index=False
    )
    result = build_supplier_ledger(
        {
            "supplier_id": "optional",
            "manifest": str(manifest),
            "sam3_summary": str(tmp_path / "missing.parquet"),
            "required_inputs": ["manifest"],
        },
        config_dir=tmp_path,
    )

    assert result.overview["sam3_status"] == "not_evaluated"
    assert result.overview["pass_clip_count"] == 1


def test_weekly_template_uses_concise_schema_and_correct_window_semantics(
    tmp_path: Path,
) -> None:
    template = _write_weekly_template(tmp_path)
    output = tmp_path / "weekly_output.xlsx"
    config_path = _weekly_config(tmp_path)

    build_acceptance_ledger(
        config_path=config_path,
        output_path=output,
        existing_workbook=template,
        overwrite=True,
    )

    workbook = load_workbook(output)
    assert workbook.sheetnames == WEEKLY_SHEETS
    xjgt = workbook["星际归途"]
    jdt = workbook["京东JDT"]
    deepreach = workbook["DeepReach"]
    canonical_header = [cell.value for cell in xjgt[2]]
    assert canonical_header == _canonical_weekly_header()
    assert [cell.value for cell in jdt[2]] == canonical_header
    assert [cell.value for cell in deepreach[2]] == canonical_header
    assert [xjgt.cell(1, column).value for column in (1, 8, 12, 16, 22)] == [
        "基础与四大项",
        "Precheck",
        "SAM3",
        "人工复核",
        "最终结果",
    ]
    assert jdt.max_row == 108
    assert jdt.freeze_panes == xjgt.freeze_panes == "A3"
    assert jdt.row_dimensions[1].height == xjgt.row_dimensions[1].height
    assert jdt.column_dimensions["A"].width == xjgt.column_dimensions["A"].width
    assert jdt["A1"].fill.fgColor.rgb == xjgt["A1"].fill.fgColor.rgb
    assert jdt["A1"].font.bold == xjgt["A1"].font.bold
    assert jdt["A1"].border.bottom.style == xjgt["A1"].border.bottom.style
    assert "无偏全局召回率" in workbook["人工问题与阈值"]["A1"].value

    columns = {name: index for index, name in enumerate(canonical_header)}
    jdt_rows = {
        row[columns["asset_id"]]: row
        for row in jdt.iter_rows(min_row=3, max_row=102, values_only=True)
    }
    for row in jdt_rows.values():
        assert row[columns["text_check_status"]] == "not_applicable"
        assert row[columns["skeleton_static_status"]] == "pass"
        assert row[columns["video_quality_status"]] == "pass"
        assert row[columns["abnormal_frame_status"]] in {
            "pass",
            "fail",
            "review",
        }

    assert jdt_rows["jdt-000"][columns["final_acceptance_status"]] == "fail"
    assert jdt_rows["jdt-000"][columns["manual_fail_window_count"]] == 1
    assert jdt_rows["jdt-001"][columns["final_acceptance_status"]] == "fail"
    assert jdt_rows["jdt-001"][columns["manual_fail_window_count"]] == 1
    assert jdt_rows["jdt-002"][columns["final_acceptance_status"]] == "pass"
    assert jdt_rows["jdt-002"][columns["manual_pass_window_count"]] == 1
    assert jdt_rows["jdt-011"][columns["final_acceptance_status"]] == "pass"
    assert jdt_rows["jdt-011"][columns["sam3_processed_window_count"]] == 1
    assert jdt_rows["jdt-011"][columns["sam3_to_manual_window_count"]] == 0
    assert sum(row[columns["manual_submitted_window_count"]] for row in jdt_rows.values()) == 11
    assert sum(row[columns["manual_reviewed_window_count"]] for row in jdt_rows.values()) == 11
    assert sum(row[columns["manual_pending_window_count"]] for row in jdt_rows.values()) == 0
    assert sum(row[columns["manual_pass_window_count"]] for row in jdt_rows.values()) == 9
    assert sum(row[columns["manual_fail_window_count"]] for row in jdt_rows.values()) == 2
    assert all(row[columns["abnormal_frame_status"]] != "review" for row in jdt_rows.values())
    config_payload = yaml.safe_load(config_path.read_text(encoding="utf-8"))
    jdt_config = next(
        row for row in config_payload["suppliers"] if row["supplier_id"] == "jdt"
    )
    jdt_details, _, _ = weekly_ledger._load_supplier_details(
        jdt_config, config_dir=config_path.parent
    )
    assert sum(len(row["_unmatched_queue_review_ids"]) for row in jdt_details) == 0
    assert sum(len(row["_unmatched_manual_review_ids"]) for row in jdt_details) == 0

    xjgt_row = next(xjgt.iter_rows(min_row=3, values_only=True))
    assert xjgt_row[columns["text_check_status"]] == "pass"
    assert xjgt_row[columns["skeleton_static_status"]] == "pass"
    assert xjgt_row[columns["video_quality_status"]] == "pass"
    assert xjgt_row[columns["abnormal_frame_status"]] == "review"
    assert xjgt_row[columns["final_acceptance_status"]] == "review"

    dr_rows = {
        row[columns["asset_id"]]: row
        for row in deepreach.iter_rows(min_row=3, values_only=True)
    }
    assert dr_rows["dr-000"][columns["abnormal_frame_status"]] == "review"
    assert dr_rows["dr-000"][columns["final_review_window_count"]] == 1
    assert dr_rows["dr-000"][columns["final_acceptance_status"]] == "review"
    assert dr_rows["dr-001"][columns["skeleton_static_status"]] == "fail"
    assert dr_rows["dr-001"][columns["precheck_fail_window_count"]] == 1
    assert dr_rows["dr-001"][columns["final_acceptance_status"]] == "fail"
    assert "keypoint_morphology" not in dr_rows["dr-001"][columns["main_reason"]]

    overview = workbook["五供应商总览"]
    overview_rows = list(overview.iter_rows(values_only=True))
    overview_header = list(overview_rows[0])
    overview_by_name = {
        row[overview_header.index("supplier_name")]: dict(zip(overview_header, row))
        for row in overview_rows[1:]
    }
    jdt_summary = overview_by_name["京东JDT / JDT"]
    assert jdt_summary["sample_clip_count"] == 100
    assert jdt_summary["pass_clip_count"] == 98
    assert jdt_summary["fail_clip_count"] == 2
    assert jdt_summary["review_clip_count"] == 0
    assert (
        jdt_summary["pass_clip_count"]
        + jdt_summary["fail_clip_count"]
        + jdt_summary["review_clip_count"]
        == jdt_summary["sample_clip_count"]
    )
    assert overview_by_name["供应商4"]["sample_clip_count"] == 0
    assert overview_by_name["供应商5"]["sample_clip_count"] == 0
    assert "异常检测拆解" not in workbook.sheetnames
    assert jdt["K3"].number_format == "0.0%"
    assert jdt["O3"].number_format == "0.0%"
    assert jdt["U3"].number_format == "0.0%"

    detail_counts = {
        "星际归途": 1,
        "DeepReach": 2,
        "京东JDT": 100,
        "供应商4": 0,
        "供应商5": 0,
    }
    for sheet_name, total in detail_counts.items():
        sheet = workbook[sheet_name]
        data_end = 2 + total
        summary_header = data_end + 2
        assert [sheet.cell(summary_header, column).value for column in range(1, 5)] == [
            "module",
            "pass_count",
            "total_clip_count",
            "pass_ratio",
        ]
        assert int(str(sheet.auto_filter.ref).split(":")[-1][1:]) == data_end
        for offset, module in enumerate(
            (
                "text_check_status",
                "skeleton_static_status",
                "video_quality_status",
                "abnormal_frame_status",
            ),
            1,
        ):
            values = [
                str(sheet.cell(row, WEEKLY_DETAIL_COLUMNS.index(module) + 1).value or "")
                .strip()
                .lower()
                for row in range(3, data_end + 1)
            ]
            assert sheet.cell(summary_header + offset, 1).value == module
            assert sheet.cell(summary_header + offset, 2).value == values.count("pass")
            assert sheet.cell(summary_header + offset, 3).value == total
            assert sheet.cell(summary_header + offset, 4).value == (
                values.count("pass") / total if total else 0
            )
            assert sheet.cell(summary_header + offset, 4).number_format == "0.0%"
