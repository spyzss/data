import csv
import json
import subprocess
import sys
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook

from tools.build_xjgt_acceptance_report import (
    ASSET_LEDGER_COLUMNS,
    ISSUE_EVENT_COLUMNS,
    build_acceptance_outputs,
    main,
)
from tests.test_qc_reporting_projection import _write_report


def _write_csv(path: Path, rows: list[dict[str, object]]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", newline="", encoding="utf-8") as handle:
        writer = csv.DictWriter(handle, fieldnames=list(rows[0]))
        writer.writeheader()
        writer.writerows(rows)


def _write_json(path: Path, value: object) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(value), encoding="utf-8")


def _fixture_inputs(tmp_path: Path) -> dict[str, Path]:
    manifest = tmp_path / "manifest.csv"
    _write_csv(
        manifest,
        [
            {
                "supplier_id": "xjgt",
                "asset_id": asset_id,
                "hdf5_path": f"/data/{asset_id}.hdf5",
                "video_path": f"/data/{asset_id}.mp4",
            }
            for asset_id in ("video_fail", "manual_fail", "sam3_only", "accepted")
        ],
    )
    precheck = tmp_path / "clip_aggregates.json"
    _write_json(
        precheck,
        [
            {
                "asset_id": asset_id,
                "check": "keypoint_temporal",
                "checked_frames": 100,
                "flagged_frames": 0,
                "clip_flag": False,
            }
            for asset_id in ("video_fail", "manual_fail", "sam3_only", "accepted")
        ],
    )
    candidate_windows = tmp_path / "candidate_windows.json"
    _write_json(
        candidate_windows,
        [
            {
                "asset_id": "accepted",
                "start_frame": 40,
                "end_frame": 50,
                "peak_frame": 45,
                "review_type": ["projection_review"],
                "trigger_reason": ["rotation_edge_risk"],
                "priority": "medium",
                "trigger_metrics": {"rotation_delta_max": 0.7},
            }
        ],
    )
    video_quality = tmp_path / "video_quality.json"
    _write_json(
        video_quality,
        [
            {
                "asset_id": "video_fail",
                "qc_summary": {
                    "status": "fail",
                    "passed": False,
                    "reasons": ["black_screen_ratio_above_max"],
                },
                "video_quality": {
                    "metadata": {
                        "fps": 30.0,
                        "frame_count": 300,
                        "duration_seconds": 10.0,
                    }
                },
            },
            {
                "asset_id": "accepted",
                "qc_summary": {
                    "status": "warn",
                    "passed": True,
                    "warn_reasons": ["laplacian_under_100_ratio_warn"],
                },
            },
        ],
    )
    sam3 = tmp_path / "sam3.json"
    _write_json(
        sam3,
        [
            {
                "asset_id": "sam3_only",
                "window_start_frame": 10,
                "window_end_frame": 30,
                "representative_frame": 20,
                "window_containment_verdict": "containment_fail",
                "strong_fail_frame_count": 4,
                "strong_fail_frame_ratio": 0.8,
                "inside_ratio_mean": 0.1,
                "projected_in_image_ratio_mean": 0.95,
                "projection_review_frame_count": 0,
                "reason": "strong containment mismatch",
            },
            {
                "asset_id": "accepted",
                "window_start_frame": 40,
                "window_end_frame": 50,
                "window_containment_verdict": "acceptable_flagged",
                "reason": "minor side-view projection",
            },
        ],
    )
    manual = tmp_path / "manual.json"
    _write_json(
        manual,
        {
            "schema_version": "skeleton_qc_manual_patch.v1",
            "segments": [
                {
                    "asset_id": "manual_fail",
                    "affected_start_frame": 20,
                    "affected_end_frame": 40,
                    "manual_outcome": "true_positive",
                    "algorithm_outcome": "true_positive",
                    "failure_mode": "severe_keypoint_offset",
                    "severity": "high",
                    "confidence": "medium",
                    "acceptance_status": "rejected",
                    "reviewer": "nathan",
                    "reason": "confirmed offset",
                },
                {
                    "asset_id": "accepted",
                    "manual_outcome": "false_positive",
                    "algorithm_outcome": "false_positive",
                    "failure_mode": "unknown",
                    "severity": "low",
                    "confidence": "medium",
                    "acceptance_status": "accepted",
                    "reviewer": "nathan",
                },
                {
                    "asset_id": "accepted",
                    "manual_outcome": "acceptable_flagged",
                    "algorithm_outcome": "acceptable_flagged",
                    "failure_mode": "acceptable_minor_misalignment",
                    "severity": "low",
                    "confidence": "medium",
                    "acceptance_status": "accepted",
                    "reviewer": "nathan",
                },
            ],
        },
    )
    return {
        "manifest": manifest,
        "precheck_clip_aggregates": precheck,
        "precheck_candidate_windows": candidate_windows,
        "video_quality_results": video_quality,
        "sam3_window_summary": sam3,
        "manual_review_labels": manual,
    }


def test_final_policy_and_output_schemas(tmp_path: Path) -> None:
    paths = _fixture_inputs(tmp_path)
    output_dir = tmp_path / "ledger"

    result = build_acceptance_outputs(output_dir=output_dir, **paths)

    with result.asset_ledger_csv.open(newline="", encoding="utf-8") as handle:
        ledger_rows = list(csv.DictReader(handle))
    ledger = {row["asset_id"]: row for row in ledger_rows}
    assert list(ledger_rows[0]) == ASSET_LEDGER_COLUMNS
    assert ledger["video_fail"]["final_verdict"] == "fail"
    assert ledger["manual_fail"]["final_verdict"] == "fail"
    assert ledger["sam3_only"]["final_verdict"] == "review"
    assert ledger["sam3_only"]["sam3_evidence_status"] == "fail"
    assert ledger["accepted"]["final_verdict"] == "pass_with_notes"
    assert ledger["accepted"]["false_positive_count"] == "1"
    assert ledger["accepted"]["acceptable_flagged_count"] == "1"

    with result.issue_events_csv.open(newline="", encoding="utf-8") as handle:
        event_rows = list(csv.DictReader(handle))
    assert list(event_rows[0]) == ISSUE_EVENT_COLUMNS
    manual_event = next(
        row
        for row in event_rows
        if row["asset_id"] == "manual_fail" and row["source_module"] == "manual_review"
    )
    assert manual_event["start_frame"] == "20"
    assert manual_event["end_frame"] == "40"
    assert manual_event["failure_mode"] == "severe_keypoint_offset"


def test_formal_cli_uses_xjgt_wrapper_and_stable_output_names(tmp_path: Path) -> None:
    archive = tmp_path / "quality_archive"
    _write_report(archive, "asset-a", "acceptance", "pass", issues=[])
    output_dir = tmp_path / "formal"

    assert (
        main(
            [
                "--quality-archive",
                str(archive),
                "--output-dir",
                str(output_dir),
            ]
        )
        == 0
    )

    assert (output_dir / "xjgt_100_asset_ledger.csv").exists()
    assert (output_dir / "xjgt_100_issue_events.csv").exists()
    assert (output_dir / "xjgt_100_acceptance_report.xlsx").exists()
    assert (output_dir / "xjgt_100_summary.json").exists()


def test_formal_xjgt_cli_reconciles_all_legacy_sidecars(tmp_path: Path) -> None:
    archive = tmp_path / "quality_archive"
    _write_report(archive, "asset-a", "acceptance", "pass", issues=[])
    output_dir = tmp_path / "formal"
    sidecars = {
        "manifest": "manifest.json",
        "precheck_check_results": "check.json",
        "precheck_clip_aggregates": "clip.json",
        "precheck_candidate_windows": "candidate.json",
        "video_quality_results": "video.json",
        "sam3_window_summary": "sam3.json",
        "manual_review_labels": "manual.json",
        "video_quality_summary": "video-summary.json",
        "sam3_clip_summary": "sam3-clip.json",
        "sam3_frame_results": "sam3-frame.json",
        "manual_review_csv": "manual.csv",
    }
    options = {
        "manifest": "--legacy-reconciliation-manifest",
        "precheck_check_results": "--legacy-reconciliation-precheck-check-results",
        "precheck_clip_aggregates": "--legacy-reconciliation-precheck-clip-aggregates",
        "precheck_candidate_windows": "--legacy-reconciliation-precheck-candidate-windows",
        "video_quality_results": "--legacy-reconciliation-video-quality-results",
        "sam3_window_summary": "--legacy-reconciliation-sam3-window-summary",
        "manual_review_labels": "--legacy-reconciliation-manual-review-labels",
        "video_quality_summary": "--legacy-reconciliation-video-quality-summary",
        "sam3_clip_summary": "--legacy-reconciliation-sam3-clip-summary",
        "sam3_frame_results": "--legacy-reconciliation-sam3-frame-results",
        "manual_review_csv": "--legacy-reconciliation-manual-review-csv",
    }
    args = ["--quality-archive", str(archive), "--output-dir", str(output_dir)]
    for source, name in sidecars.items():
        path = tmp_path / name
        if path.suffix == ".csv":
            path.write_text("asset_id,verdict\nasset-a,pass\n", encoding="utf-8")
        else:
            path.write_text(
                json.dumps([{"asset_id": "asset-a", "verdict": "pass"}]),
                encoding="utf-8",
            )
        args.extend([options[source], str(path)])

    assert main(args) == 0
    reconciliation = pd.read_csv(output_dir / "reconciliation.csv")
    assert set(reconciliation["source"]) == set(sidecars)
    assert (output_dir / "xjgt_100_acceptance_report.xlsx").exists()


def test_formal_xjgt_script_cli_runs_from_repository_root(tmp_path: Path) -> None:
    archive = tmp_path / "quality_archive"
    _write_report(archive, "asset-a", "acceptance", "pass", issues=[])
    output_dir = tmp_path / "formal"
    repo_root = Path(__file__).resolve().parents[1]
    result = subprocess.run(
        [
            sys.executable,
            str(repo_root / "tools" / "build_xjgt_acceptance_report.py"),
            "--quality-archive",
            str(archive),
            "--output-dir",
            str(output_dir),
        ],
        cwd=repo_root,
        capture_output=True,
        text=True,
        check=False,
    )

    assert result.returncode == 0, result.stderr
    assert (output_dir / "xjgt_100_acceptance_report.xlsx").exists()


def test_workbook_and_summary_contain_required_sections(tmp_path: Path) -> None:
    paths = _fixture_inputs(tmp_path)
    result = build_acceptance_outputs(output_dir=tmp_path / "ledger", **paths)

    workbook = load_workbook(result.workbook_xlsx, read_only=True, data_only=False)
    assert workbook.sheetnames == [
        "supplier_summary",
        "asset_ledger",
        "issue_events",
        "manual_review_summary",
        "sam3_summary",
        "video_quality_summary",
        "method_thresholds",
        "pipeline_notes",
    ]
    thresholds = [
        row[0]
        for row in workbook["method_thresholds"].iter_rows(min_row=2, values_only=True)
    ]
    assert "sam3_projected_in_image_ratio_projection_review" in thresholds
    assert "sam3_window_strong_fail_frame_count" in thresholds

    summary = json.loads(result.summary_json.read_text(encoding="utf-8"))
    assert summary["asset_count"] == 4
    assert summary["final_verdict_counts"] == {
        "fail": 2,
        "pass_with_notes": 1,
        "review": 1,
    }
    assert summary["manual_outcome_counts"] == {
        "acceptable_flagged": 1,
        "false_positive": 1,
        "true_positive": 1,
    }
