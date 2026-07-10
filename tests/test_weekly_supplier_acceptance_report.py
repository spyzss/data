import ast
import csv
import inspect
import json
import subprocess
import sys
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook

from tools.build_weekly_supplier_acceptance_report import (
    CORE_DETAIL_COLUMNS,
    HARD_ISSUE_COLUMNS,
    MANUAL_ISSUE_COLUMNS,
    SUMMARY_COLUMNS,
    THRESHOLD_RULE_COLUMNS,
    acceptance_and_review_status,
    aggregate_manual_review,
    build_weekly_report,
    collect_input_audit,
    count_fail_indicators,
    detail_columns_for_rows,
    main,
    normalize_asset_id,
    recompute_xjgt_final_status,
    resolve_xjgt_text_status,
    skeleton_static_status_from_parts,
    supplier_summary_from_details,
    video_quality_atomic_statuses,
)


def test_weekly_builder_has_no_duplicate_top_level_definitions() -> None:
    module_path = (
        Path(__file__).parents[1]
        / "tools"
        / "build_weekly_supplier_acceptance_report.py"
    )
    tree = ast.parse(module_path.read_text(encoding="utf-8"))
    definitions: dict[str, list[int]] = {}
    for node in tree.body:
        if isinstance(node, (ast.FunctionDef, ast.AsyncFunctionDef, ast.ClassDef)):
            definitions.setdefault(node.name, []).append(node.lineno)

    duplicates = {
        name: lines for name, lines in definitions.items() if len(lines) > 1
    }
    assert duplicates == {}


def test_weekly_builder_cli_supports_script_and_module_invocation(
    tmp_path: Path,
) -> None:
    repo_root = Path(__file__).parents[1]
    script = repo_root / "tools" / "build_weekly_supplier_acceptance_report.py"
    commands = [
        [sys.executable, str(script), "--help"],
        [sys.executable, "-m", "tools.build_weekly_supplier_acceptance_report", "--help"],
    ]
    for command in commands:
        completed = subprocess.run(
            command,
            cwd=repo_root,
            check=False,
            capture_output=True,
            text=True,
        )
        assert completed.returncode == 0, completed.stderr
        assert "--audit-inputs" in completed.stdout
        assert "--skip-xjgt-text" in completed.stdout

    run_root = tmp_path / "acceptance_5x100"
    completed = subprocess.run(
        [
            sys.executable,
            str(script),
            "--run-root",
            str(run_root),
            "--audit-inputs",
        ],
        cwd=repo_root,
        check=False,
        capture_output=True,
        text=True,
    )
    assert completed.returncode == 0, completed.stderr
    audit = json.loads(completed.stdout)
    assert audit["run_root"] == str(run_root)
    assert not (run_root / "weekly_supplier_acceptance_report.xlsx").exists()
    assert not (run_root / "weekly_supplier_summary.csv").exists()


def _write_csv(path: Path, rows: list[dict[str, object]]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    fieldnames = list(dict.fromkeys(key for row in rows for key in row))
    with path.open("w", newline="", encoding="utf-8") as handle:
        writer = csv.DictWriter(handle, fieldnames=fieldnames)
        writer.writeheader()
        writer.writerows(rows)


def _write_parquet(path: Path, rows: list[dict[str, object]]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    pd.DataFrame(rows).to_parquet(path, index=False)


def _check_result(
    episode_idx: int,
    check: str,
    frame_idx: int,
    metrics: dict[str, object],
    flag: bool | None,
    reason: str = "fixture",
) -> dict[str, object]:
    return {
        "check": check,
        "episode_idx": episode_idx,
        "frame_idx": frame_idx,
        "metrics": json.dumps(metrics),
        "flag": flag,
        "reason": reason,
    }


def _text_metrics(*, text_en_valid: bool = True) -> dict[str, float]:
    return {
        "field_present_scene": 1.0,
        "field_nonempty_scene": 1.0,
        "field_present_task": 1.0,
        "field_nonempty_task": 1.0,
        "field_present_text_en": float(text_en_valid),
        "field_nonempty_text_en": float(text_en_valid),
        "missing_field_count": float(not text_en_valid),
        "empty_field_count": 0.0,
    }


def _video_quality_result(asset_id: str, frame_count: int) -> dict[str, object]:
    return {
        "asset_id": asset_id,
        "qc_summary": {
            "status": "pass",
            "passed": True,
            "reasons": [],
            "warn_reasons": [],
        },
        "hdf5_text_info": {"alignment": {"status": "matched"}},
        "video_quality": {
            "evaluation": {
                "decision": "pass",
                "passed": True,
                "reasons": [],
                "warn_reasons": [],
            },
            "metadata": {"frame_count": frame_count, "fps": 30.0},
            "thresholds": {
                "fps": {
                    "min_fps_pass": 24.0,
                    "min_fps_fail": 20.0,
                },
                "exposure": {
                    "black": {
                        "mean_y_max": 10.0,
                        "ratio_pass": 0.01,
                        "ratio_warn": 0.90,
                    },
                    "over_dark": {
                        "mean_y_max": 35.0,
                        "ratio_pass": 0.05,
                    },
                    "over_exposed": {
                        "mean_y_min": 235.0,
                        "ratio_pass": 0.05,
                    },
                },
                "sharpness_global": {
                    "laplacian_p10_pass": 35.0,
                    "tenengrad_p10_pass": 12.0,
                },
                "freeze": {
                    "frozen_frame_ratio_warn": 0.10,
                    "max_consecutive_frozen_sec_fail": 1.0,
                },
                "timeline": {"drop_frame_ratio_warn": 0.10},
                "hdf5_alignment": {"max_delta_frames_warn": 5},
            },
            "metrics": {
                "exposure_metrics": {
                    "black_frame_ratio": 0.0,
                    "black_frame_count_estimate": 0,
                    "mean_over_dark_ratio": 0.0,
                    "mean_over_exposed_ratio": 0.0,
                },
                "sharpness_global": {
                    "laplacian_p10": 120.0,
                    "tenengrad_p10": 20.0,
                },
                "freeze_metrics": {
                    "frozen_frame_ratio": 0.0,
                    "frozen_intervals": [],
                },
                "timeline_metrics": {"drop_frame_ratio": 0.0},
                "hdf5_alignment": {"status": "matched"},
            },
        },
    }


def _prepare_inputs(run_root: Path) -> None:
    _write_csv(
        run_root / "manifests" / "supplier_manifest_xjgt_100.csv",
        [
            {"supplier_id": "xjgt", "asset_id": "1001", "video_path": "/videos/1001.mp4"},
            {"supplier_id": "xjgt", "asset_id": "1002", "video_path": "/videos/1002.mp4"},
            {"supplier_id": "xjgt", "asset_id": "1003", "video_path": ""},
        ],
    )
    # The old ledger is deliberately stale. The weekly report must map the
    # raw module outputs below instead of copying these values.
    common = {
        "hdf5_text_status": "not_run",
        "keypoint_missing_status": "not_run",
        "video_quality_status": "not_run",
        "temporal_status": "not_run",
        "sam3_evidence_status": "not_run",
        "manual_review_status": "not_run",
        "top_issue_types": "old_stale_issue",
        "evidence_paths": "old_ledger.csv",
    }
    _write_csv(
        run_root / "xjgt" / "ledger" / "xjgt_100_asset_ledger.csv",
        [
            {
                **common,
                "asset_id": "1001",
                "quality_hand_status": "not_run",
                "final_verdict": "review",
            },
            {
                **common,
                "asset_id": "1002",
                "quality_hand_status": "not_run",
                "final_verdict": "review",
            },
            {**common, "asset_id": "1003", "final_verdict": "review"},
        ],
    )
    _write_csv(
        run_root / "xjgt" / "ledger" / "xjgt_100_issue_events.csv",
        [
            {
                "asset_id": "1001",
                "source_module": "precheck",
                "source_verdict": "fail",
                "failure_mode": "keypoint_missing",
                "start_frame": 80,
                "end_frame": 89,
            },
            {
                "asset_id": "1002",
                "source_module": "sam3_containment",
                "auto_verdict": "fail",
                "failure_mode": "containment_fail",
                "start_frame": 150,
                "end_frame": 169,
            },
            {
                "asset_id": "1002",
                "source_module": "sam3_containment",
                "auto_verdict": "fail",
                "failure_mode": "containment_fail",
                "start_frame": 120,
                "end_frame": 129,
            },
            {
                "asset_id": "1002",
                "source_module": "sam3_containment",
                "auto_verdict": "review",
                "failure_mode": "projection_review",
                "start_frame": 170,
                "end_frame": 199,
            }
        ],
    )
    summary_path = run_root / "xjgt" / "ledger" / "xjgt_100_summary.json"
    summary_path.write_text(
        json.dumps(
            {
                "asset_count": 100,
                "final_verdict_counts": {
                    "fail": 21,
                    "review": 64,
                    "pass_with_notes": 15,
                },
            }
        ),
        encoding="utf-8",
    )

    precheck_dir = run_root / "xjgt" / "precheck"
    (precheck_dir / "precheck_config.yaml").parent.mkdir(
        parents=True, exist_ok=True
    )
    (precheck_dir / "precheck_config.yaml").write_text(
        """
text_integrity:
  required_fields: [scene, task, text_en]
keypoint_temporal:
  min_angle_degrees: 5.0
  max_angle_degrees: 175.0
keypoint_morphology:
  duplicate_joint_distance_m: 0.00001
  min_palm_scale_m: 0.0001
  max_bone_length_ratio_spread_review: 3.0
  max_bone_length_ratio_spread_fail: 8.0
  max_normalized_bone_length_review: 3.0
  max_normalized_bone_length_fail: 6.0
  max_zero_length_bone_count_review: 1
  max_zero_length_bone_count_fail: 2
  max_duplicate_joint_pair_count_review: 1
  max_duplicate_joint_pair_count_fail: 3
  min_joint_angle_deg_review: 5.0
  min_joint_angle_deg_fail: 1.0
  max_joint_angle_violation_fraction_review: 0.15
  max_joint_angle_violation_fraction_fail: 0.40
skeleton_quality_score:
  joint_angle_change_deg_max_threshold: 10.0
  rotation_delta_max_threshold: 0.45
  joint_acceleration_m_s2_max_threshold: 15.0
  joint_displacement_m_max_threshold: 0.05
  candidate_gap_close_frames: 2
  candidate_min_seed_run_frames: 3
  candidate_pre_context_frames: 10
  candidate_post_context_frames: 10
  reject_missing_keypoints: true
  allowed_missing_keypoints_per_hand: 0
""".strip()
        + "\n",
        encoding="utf-8",
    )
    check_rows = [
        _check_result(
            0,
            "text_integrity",
            -1,
            _text_metrics(text_en_valid=False),
            True,
            "missing optional supplier text field",
        ),
        _check_result(
            0,
            "quality_score",
            -1,
            {"pass_ratio": 1.0, "pass_threshold": 0.9},
            True,
        ),
        _check_result(
            0,
            "skeleton_quality_score",
            0,
            {
                "keypoint_presence_invalid": 0.0,
                "valid_keypoint_count_left": 21.0,
                "valid_keypoint_count_right": 21.0,
                "skeleton_verdict": "review",
            },
            None,
        ),
        *[
            _check_result(
                0,
                "skeleton_quality_score",
                frame_idx,
                {
                    "keypoint_presence_invalid": 1.0,
                    "valid_keypoint_count_left": 20.0,
                    "valid_keypoint_count_right": 21.0,
                    "missing_keypoint_count_left": 1.0,
                    "missing_keypoint_count_right": 0.0,
                    "skeleton_verdict": "invalid",
                },
                True,
            )
            for frame_idx in range(80, 90)
        ],
        _check_result(
            0,
            "keypoint_morphology",
            0,
            {"morphology_verdict": "pass"},
            False,
        ),
        _check_result(
            0,
            "keypoint_morphology",
            -1,
            {"morphology_verdict": "pass"},
            False,
        ),
        _check_result(1, "text_integrity", -1, _text_metrics(), None),
        _check_result(
            1,
            "quality_score",
            0,
            {"frame_score": 0.0, "quality_left": 0.0, "quality_right": 1.0},
            None,
        ),
        _check_result(
            1,
            "quality_score",
            -1,
            {"pass_ratio": 0.5, "pass_threshold": 0.9},
            False,
        ),
        _check_result(
            1,
            "skeleton_quality_score",
            0,
            {
                # Supplier quality_hand may make the producer's combined
                # invalid flag true, but all 21 geometry points still exist.
                "keypoint_presence_invalid": 1.0,
                "low_quality_hand_invalid": 1.0,
                "valid_keypoint_count_left": 21.0,
                "valid_keypoint_count_right": 21.0,
                "skeleton_verdict": "review",
            },
            None,
        ),
        _check_result(
            1,
            "keypoint_morphology",
            0,
            {"morphology_verdict": "pass"},
            False,
        ),
        _check_result(
            1,
            "keypoint_morphology",
            -1,
            {"morphology_verdict": "pass"},
            False,
        ),
        _check_result(2, "text_integrity", -1, _text_metrics(), None),
        _check_result(
            2,
            "skeleton_quality_score",
            0,
            {
                "keypoint_presence_invalid": 0.0,
                "valid_keypoint_count_left": 21.0,
                "valid_keypoint_count_right": 21.0,
                "skeleton_verdict": "review",
            },
            None,
        ),
        _check_result(
            2,
            "keypoint_morphology",
            0,
            {"morphology_verdict": "pass"},
            False,
        ),
        _check_result(
            2,
            "keypoint_morphology",
            -1,
            {"morphology_verdict": "pass"},
            False,
        ),
    ]
    _write_parquet(precheck_dir / "check_results.parquet", check_rows)
    _write_parquet(
        precheck_dir / "clip_aggregates.parquet",
        [
            {
                "episode_idx": episode_idx,
                "check": check,
                "checked_frames": 1,
                "flagged_frames": 0,
                "uncalibrated_frames": 0,
                "clip_flag": False,
            }
            for episode_idx in range(3)
            for check in (
                "text_integrity",
                "quality_score",
                "skeleton_quality_score",
            )
        ]
        + [
            {
                "episode_idx": episode_idx,
                "check": "keypoint_morphology",
                "checked_frames": 1,
                "flagged_frames": 0,
                "uncalibrated_frames": 0,
                "clip_flag": False,
            }
            for episode_idx in (0, 1, 2)
        ],
    )
    (precheck_dir / "candidate_windows.json").write_text(
        json.dumps(
            [
                {
                    "asset_id": "1001",
                    "episode_idx": 0,
                    "start_frame": 0,
                    "end_frame": 49,
                    "peak_frame": 5,
                    "review_type": ["temporal_geometry_review"],
                    "trigger_reason": ["multi_signal_seed"],
                },
                {
                    "asset_id": "1002",
                    "episode_idx": 1,
                    "start_frame": 100,
                    "end_frame": 199,
                    "peak_frame": 150,
                    "review_type": ["temporal_geometry_review"],
                    "trigger_reason": ["multi_signal_seed"],
                },
                {
                    "asset_id": "1003",
                    "episode_idx": 2,
                    "start_frame": 0,
                    "end_frame": 20,
                    "peak_frame": 10,
                    "review_type": ["temporal_geometry_review"],
                    "trigger_reason": ["multi_signal_seed"],
                },
            ]
        ),
        encoding="utf-8",
    )

    video_dir = run_root / "xjgt" / "video_quality"
    _write_csv(
        video_dir / "video_quality_acceptance_summary.csv",
        [
            {"asset_id": asset_id, "status": "pass", "passed": True}
            for asset_id in ("1001", "1002", "1003")
        ],
    )
    (video_dir / "video_quality_results.json").write_text(
        json.dumps(
            [
                _video_quality_result(asset_id, frame_count)
                for asset_id, frame_count in (
                    ("1001", 100),
                    ("1002", 200),
                    ("1003", 100),
                )
            ]
        ),
        encoding="utf-8",
    )

    sam3_dir = run_root / "xjgt" / "sam3_containment"
    sam3_dir.mkdir(parents=True, exist_ok=True)
    (sam3_dir / "window_keypoint_containment_summary.json").write_text(
        json.dumps(
            [
                {
                    "asset_id": "1001",
                    "window_start_frame": 0,
                    "window_end_frame": 49,
                    "sampled_frame_count": 5,
                    "strong_fail_frame_count": 0,
                    "window_containment_verdict": "mixed_review",
                    "reason": "mixed containment evidence",
                },
                {
                    "asset_id": "1002",
                    "window_start_frame": 120,
                    "window_end_frame": 129,
                    "sampled_frame_count": 5,
                    "strong_fail_frame_count": 4,
                    "window_containment_verdict": "containment_fail",
                    "reason": "sustained strong keypoint-mask mismatch",
                },
                {
                    "asset_id": "1002",
                    "window_start_frame": 150,
                    "window_end_frame": 169,
                    "sampled_frame_count": 5,
                    "strong_fail_frame_count": 4,
                    "window_containment_verdict": "containment_fail",
                    "reason": "sustained strong keypoint-mask mismatch",
                },
                {
                    "asset_id": "1002",
                    "window_start_frame": 170,
                    "window_end_frame": 199,
                    "sampled_frame_count": 5,
                    "strong_fail_frame_count": 0,
                    "window_containment_verdict": "projection_review",
                    "reason": "only insufficient projection evidence",
                },
                {
                    "asset_id": "1003",
                    "window_start_frame": 0,
                    "window_end_frame": 20,
                    "sampled_frame_count": 5,
                    "strong_fail_frame_count": 0,
                    "window_containment_verdict": "review",
                    "reason": "uncertain containment evidence",
                },
            ]
        ),
        encoding="utf-8",
    )
    (sam3_dir / "frame_keypoint_containment.json").write_text("[]", encoding="utf-8")
    (sam3_dir / "clip_keypoint_containment.json").write_text("[]", encoding="utf-8")
    _write_csv(
        run_root
        / "xjgt"
        / "video_review_full"
        / "review_queue_with_clips.csv",
        [
            {
                "review_id": "rq_1001_0_49",
                "asset_id": "1001",
                "module": "sam3_containment",
                "source_level": "window",
                "window_start_frame": 0,
                "window_end_frame": 49,
            },
            {
                "review_id": "rq_1002_120_129",
                "asset_id": "1002",
                "module": "sam3_containment",
                "source_level": "window",
                "window_start_frame": 120,
                "window_end_frame": 129,
            },
            {
                "review_id": "rq_1002_150_169",
                "asset_id": "1002",
                "module": "sam3_containment",
                "source_level": "window",
                "window_start_frame": 150,
                "window_end_frame": 169,
            },
            {
                "review_id": "rq_1002_170_199",
                "asset_id": "1002",
                "module": "sam3_containment",
                "source_level": "window",
                "window_start_frame": 170,
                "window_end_frame": 199,
            },
            {
                "review_id": "rq_1003_0_20",
                "asset_id": "1003",
                "module": "sam3_containment",
                "source_level": "window",
                "window_start_frame": 0,
                "window_end_frame": 20,
            },
        ],
    )
    _write_csv(
        run_root
        / "xjgt"
        / "manual_review"
        / "manual_labels_autosave.normalized.csv",
        [
            {
                "asset_id": "1001",
                "manual_outcome": "true_positive",
                "window_start_frame": 0,
                "window_end_frame": 49,
                "affected_start_frame": 0,
                "affected_end_frame": 4,
                "failure_mode": "severe_keypoint_offset",
            },
            {
                "asset_id": "1001",
                "manual_outcome": "true_positive",
                "window_start_frame": 0,
                "window_end_frame": 49,
                "affected_start_frame": 4,
                "affected_end_frame": 9,
                "failure_mode": "severe_keypoint_offset",
            },
            {
                "asset_id": "1001",
                "manual_outcome": "acceptable_flagged",
                "window_start_frame": 50,
                "window_end_frame": 99,
                "affected_start_frame": 10,
                "affected_end_frame": 99,
                "failure_mode": "acceptable_minor_misalignment",
            },
            {
                "asset_id": "1002",
                "manual_outcome": "true_positive",
                "window_start_frame": 0,
                "window_end_frame": 99,
                "affected_start_frame": 10,
                "affected_end_frame": 19,
                "failure_mode": "hand_out_of_frame",
            },
            {
                "asset_id": "1002",
                "manual_outcome": "false_positive",
                "window_start_frame": 100,
                "window_end_frame": 149,
                "affected_start_frame": 20,
                "affected_end_frame": 199,
                "failure_mode": "severe_keypoint_offset",
            },
        ],
    )
    manual_patch = {
        "schema_version": "skeleton_qc_manual_patch.v1",
        "source": "manual_review_queue",
        "segments": [],
    }
    (
        run_root / "xjgt" / "manual_review" / "manual_labels_patch.json"
    ).write_text(json.dumps(manual_patch), encoding="utf-8")
    _write_csv(
        run_root / "deepreach" / "manifests" / "supplier_manifest_deepreach.csv",
        [
            {
                "supplier_id": "deepreach",
                "asset_id": f"task_{index:03d}__head",
                "video_path": f"/deepreach/task_{index:03d}.mp4",
            }
            for index in range(8)
        ],
    )


def _patch_frame_probe(monkeypatch) -> list[Path]:
    observed_paths: list[Path] = []

    def fake_probe(path: Path) -> int:
        observed_paths.append(path)
        return {
            "1001.mp4": 100,
            "1002.mp4": 200,
            "1003_video.mp4": 100,
            **{f"task_{index:03d}.mp4": 30 + index for index in range(8)},
        }[path.name]

    monkeypatch.setattr(
        "tools.build_weekly_supplier_acceptance_report.probe_video_frame_count",
        fake_probe,
    )
    return observed_paths


def test_xjgt_recomputes_frames_manual_ratio_and_final_status(
    tmp_path: Path, monkeypatch
) -> None:
    run_root = tmp_path / "acceptance_5x100"
    _prepare_inputs(run_root)
    observed_paths = _patch_frame_probe(monkeypatch)

    outputs = build_weekly_report(run_root)

    workbook = load_workbook(outputs.workbook_xlsx, read_only=True)
    assert workbook.sheetnames == [
        "五供应商总览",
        "星际归途",
        "DeepReach",
        "供应商3",
        "供应商4",
        "供应商5",
        "人工与难测问题统计",
        "人工问题与阈值",
    ]
    assert [cell.value for cell in workbook["五供应商总览"][1]] == SUMMARY_COLUMNS
    detail_headers = [cell.value for cell in workbook["星际归途"][1]]
    assert detail_headers[: len(CORE_DETAIL_COLUMNS)] == CORE_DETAIL_COLUMNS
    assert [
        column
        for column in detail_headers
        if column.startswith("text_field_")
    ] == [
        "text_field_present_scene_status",
        "text_field_nonempty_scene_status",
        "text_field_present_task_status",
        "text_field_nonempty_task_status",
        "text_field_present_text_en_status",
        "text_field_nonempty_text_en_status",
    ]
    assert [
        column
        for column in detail_headers
        if column.startswith("video_")
        and column.endswith("_status")
        and column != "video_quality_status"
    ] == [
        "video_black_screen_status",
        "video_underexposure_status",
        "video_overexposure_status",
        "video_blur_status",
        "video_freeze_stutter_status",
        "video_frame_alignment_status",
    ]
    assert "temporal_sam3_manual_status" not in detail_headers
    assert detail_headers.index("skeleton_missing_status") < detail_headers.index(
        "skeleton_morphology_status"
    )
    assert detail_headers.index("abnormal_frame_status") < detail_headers.index(
        "temporal_status"
    )
    rows = list(workbook["星际归途"].iter_rows(min_row=2, values_only=True))
    columns = {name: index for index, name in enumerate(detail_headers)}
    assert {
        "skeleton_missing_fail_frame_count",
        "skeleton_missing_fail_frame_ratio",
        "skeleton_morphology_fail_frame_count",
        "skeleton_morphology_fail_frame_ratio",
        "skeleton_static_fail_frame_count",
        "skeleton_static_fail_frame_ratio",
        "skeleton_static_status",
        "abnormal_fail_frame_count",
        "abnormal_fail_frame_ratio",
        "abnormal_frame_status",
        "fail_indicator_count",
        "video_quality_fail_frame_count",
        "video_quality_fail_frame_ratio",
        "auto_fail_frame_count",
        "reviewed_auto_fail_frame_count",
        "reviewed_auto_fail_true_positive_frame_count",
        "reviewed_auto_fail_false_positive_frame_count",
        "unreviewed_auto_fail_frame_count",
        "auto_fail_precision_on_reviewed",
        "abnormal_status_reason",
        "text_status_reason",
        "skeleton_static_status_reason",
        "mapped_precheck_checks",
        "missing_expected_checks",
        "precheck_mapping_status",
        "final_status_reason",
    } <= set(columns)
    by_asset = {row[columns["asset_id"]]: row for row in rows}

    assert by_asset["1001"][columns["total_frames"]] == 100
    assert by_asset["1001"][columns["frame_count_status"]] == "ok"
    assert by_asset["1001"][columns["manual_problem_frame_count"]] == 10
    assert by_asset["1001"][columns["manual_problem_frame_ratio_of_clip"]] == 0.1
    assert by_asset["1001"][columns["manual_reviewed_frame_count"]] == 100
    assert by_asset["1001"][columns["manual_problem_ratio_of_reviewed"]] == 0.1
    assert by_asset["1001"][columns["manual_review_status"]] == "fail"
    assert by_asset["1001"][columns["text_check_status"]] == "fail"
    assert "mapped_text_integrity=fail" in by_asset["1001"][columns["text_status_reason"]]
    assert by_asset["1001"][columns["skeleton_missing_status"]] == "fail"
    assert by_asset["1001"][columns["skeleton_morphology_status"]] == "pass"
    assert by_asset["1001"][columns["supplier_quality_signal"]] == "provided_ok"
    assert by_asset["1001"][columns["temporal_status"]] == "review"
    assert by_asset["1001"][columns["skeleton_static_fail_frame_count"]] == 10
    assert by_asset["1001"][columns["skeleton_static_fail_frame_ratio"]] == 0.1
    assert by_asset["1001"][columns["skeleton_static_status"]] == "fail"
    assert "missing_status=fail" in by_asset["1001"][columns["skeleton_static_status_reason"]]
    assert by_asset["1001"][columns["abnormal_fail_frame_count"]] == 10
    assert by_asset["1001"][columns["abnormal_fail_frame_ratio"]] == 0.1
    assert by_asset["1001"][columns["abnormal_frame_status"]] == "fail"
    assert by_asset["1001"][columns["fail_indicator_count"]] == 3
    assert by_asset["1001"][columns["acceptance_status"]] == "fail"
    assert by_asset["1001"][columns["abnormal_frame_status_v2"]] == "fail"
    assert by_asset["1001"][columns["acceptance_status_v2"]] == "fail"
    assert by_asset["1001"][columns["review_status_v2"]] == "completed"
    assert "skeleton_static_status" in by_asset["1001"][columns["final_status_reason"]]
    assert "abnormal_frame_status" in by_asset["1001"][columns["final_status_reason"]]
    assert by_asset["1001"][columns["precheck_mapping_status"]] == "mapped"
    assert "skeleton_quality_score" in by_asset["1001"][columns["mapped_precheck_checks"]]
    assert by_asset["1001"][columns["missing_expected_checks"]] in (None, "")

    assert by_asset["1002"][columns["total_frames"]] == 200
    assert by_asset["1002"][columns["frame_count_status"]] == "ok"
    assert by_asset["1002"][columns["manual_problem_frame_count"]] == 10
    assert by_asset["1002"][columns["manual_problem_frame_ratio_of_clip"]] == 0.05
    assert by_asset["1002"][columns["manual_reviewed_frame_count"]] == 150
    assert by_asset["1002"][columns["manual_problem_ratio_of_reviewed"]] == 10 / 150
    assert by_asset["1002"][columns["manual_review_status"]] == "pass"
    assert by_asset["1002"][columns["text_check_status"]] == "pass"
    assert "mapped_text_integrity=pass" in by_asset["1002"][columns["text_status_reason"]]
    assert by_asset["1002"][columns["skeleton_static_status"]] == "pass"
    assert by_asset["1002"][columns["abnormal_fail_frame_count"]] == 60
    assert by_asset["1002"][columns["abnormal_fail_frame_ratio"]] == 0.30
    assert by_asset["1002"][columns["abnormal_frame_status"]] == "fail"
    assert by_asset["1002"][columns["fail_indicator_count"]] == 1
    assert by_asset["1002"][columns["acceptance_status"]] == "fail"
    assert by_asset["1002"][columns["abnormal_frame_status_v2"]] == "fail"
    assert by_asset["1002"][columns["acceptance_status_v2"]] == "fail"
    assert by_asset["1002"][columns["review_status_v2"]] == "review"
    assert by_asset["1002"][columns["supplier_quality_signal"]] == "low"
    assert by_asset["1002"][columns["auto_fail_frame_count"]] == 30
    assert by_asset["1002"][columns["reviewed_auto_fail_frame_count"]] == 10
    assert (
        by_asset["1002"][columns["reviewed_auto_fail_true_positive_frame_count"]]
        == 0
    )
    assert (
        by_asset["1002"][columns["reviewed_auto_fail_false_positive_frame_count"]]
        == 10
    )
    assert by_asset["1002"][columns["unreviewed_auto_fail_frame_count"]] == 20
    assert by_asset["1002"][columns["auto_fail_precision_on_reviewed"]] == 0.0
    assert "unreviewed_auto_fail_frames=20" in by_asset["1002"][columns["abnormal_status_reason"]]
    assert "acceptance_status=fail" in by_asset["1002"][columns["final_status_reason"]]

    assert by_asset["1003"][columns["total_frames"]] == 100
    assert by_asset["1003"][columns["frame_count_status"]] == "ok"
    assert by_asset["1003"][columns["manual_problem_frame_count"]] == 0
    assert by_asset["1003"][columns["manual_reviewed_frame_count"]] == 0
    assert by_asset["1003"][columns["manual_review_status"]] == "not_reviewed"
    assert by_asset["1003"][columns["skeleton_morphology_status"]] == "pass"
    assert by_asset["1003"][columns["skeleton_static_status"]] == "pass"
    assert by_asset["1003"][columns["abnormal_frame_status"]] == "fail"
    assert by_asset["1003"][columns["fail_indicator_count"]] == 1
    assert by_asset["1003"][columns["acceptance_status"]] == "fail"
    assert by_asset["1003"][columns["abnormal_frame_status_v2"]] == "review"
    assert by_asset["1003"][columns["acceptance_status_v2"]] == "pass"
    assert by_asset["1003"][columns["review_status_v2"]] == "review"
    assert (
        by_asset["1003"][columns["supplier_quality_signal"]]
        == "not_provided"
    )
    assert by_asset["1003"][columns["precheck_mapping_status"]] == "mapped"
    assert Path(
        "/mnt/oss/egodata/XJGT_20260629/video/1003_video.mp4"
    ) in observed_paths

    with outputs.summary_csv.open(newline="", encoding="utf-8") as handle:
        summary = list(csv.DictReader(handle))
    assert list(summary[0]) == SUMMARY_COLUMNS == [
        "supplier_name",
        "sample_clip_count",
        "total_frame_count",
        "problem_frame_count",
        "problem_frame_ratio",
        "pass_clip_count",
        "pass_clip_ratio",
        "fail_clip_count",
        "fail_clip_ratio",
        "problem_frame_count_v2",
        "problem_frame_ratio_v2",
        "pass_clip_count_v2",
        "pass_clip_ratio_v2",
        "fail_clip_count_v2",
        "fail_clip_ratio_v2",
    ]
    xjgt = summary[0]
    assert xjgt["sample_clip_count"] == "3"
    assert xjgt["total_frame_count"] == "400"
    assert xjgt["problem_frame_count"] == "101"
    assert xjgt["problem_frame_ratio"] == str(101 / 400)
    assert xjgt["pass_clip_count"] == "0"
    assert xjgt["pass_clip_ratio"] == "0.0"
    assert xjgt["fail_clip_count"] == "3"
    assert xjgt["fail_clip_ratio"] == "1.0"
    assert xjgt["problem_frame_count_v2"] == "50"
    assert xjgt["problem_frame_ratio_v2"] == str(50 / 400)
    assert xjgt["pass_clip_count_v2"] == "1"
    assert xjgt["pass_clip_ratio_v2"] == str(1 / 3)
    assert xjgt["fail_clip_count_v2"] == "2"
    assert xjgt["fail_clip_ratio_v2"] == str(2 / 3)

    overview = workbook["五供应商总览"]
    overview_columns = {
        cell.value: cell.column for cell in overview[1]
    }
    for column in (
        "problem_frame_ratio",
        "pass_clip_ratio",
        "fail_clip_ratio",
        "problem_frame_ratio_v2",
        "pass_clip_ratio_v2",
        "fail_clip_ratio_v2",
    ):
        assert overview.cell(2, overview_columns[column]).number_format == "0.0%"


def test_final_sheet_combines_manual_issues_and_long_form_thresholds(
    tmp_path: Path, monkeypatch
) -> None:
    run_root = tmp_path / "acceptance_5x100"
    _prepare_inputs(run_root)
    _patch_frame_probe(monkeypatch)

    outputs = build_weekly_report(run_root)

    workbook = load_workbook(outputs.workbook_xlsx)
    sheet = workbook["人工问题与阈值"]
    assert set(sheet.tables) == {"ManualIssueStats", "EffectiveThresholdRules"}
    values = list(sheet.iter_rows(values_only=True))
    title_rows = {
        row[0]: index + 1
        for index, row in enumerate(values)
        if row and row[0] in {
            "人工确认问题类型统计",
            "本次验收生效阈值与规则",
        }
    }
    assert set(title_rows) == {
        "人工确认问题类型统计",
        "本次验收生效阈值与规则",
    }

    manual_title_row = title_rows["人工确认问题类型统计"]
    threshold_title_row = title_rows["本次验收生效阈值与规则"]
    manual_headers = list(values[manual_title_row][: len(MANUAL_ISSUE_COLUMNS)])
    threshold_headers = list(
        values[threshold_title_row][: len(THRESHOLD_RULE_COLUMNS)]
    )
    assert manual_headers == MANUAL_ISSUE_COLUMNS
    assert threshold_headers == THRESHOLD_RULE_COLUMNS
    assert threshold_title_row - (manual_title_row + 1) >= 5

    manual_rows = [
        dict(zip(MANUAL_ISSUE_COLUMNS, row))
        for row in values[manual_title_row + 1 : threshold_title_row - 2]
        if row[0] and row[0] != (
            "人工复核采用高风险定向抽样，本表适合分析供应商问题构成，"
            "不代表供应商全量数据的无偏问题率。"
        )
    ]
    assert len(manual_rows) == 2
    by_issue = {row["issue_type"]: row for row in manual_rows}
    assert by_issue["severe_keypoint_offset"]["confirmed_problem_frame_count"] == 10
    assert by_issue["severe_keypoint_offset"]["true_positive_count"] == 2
    assert by_issue["severe_keypoint_offset"]["false_positive_count"] == 1
    assert by_issue["hand_out_of_frame"]["confirmed_problem_frame_count"] == 10
    assert "acceptable_minor_misalignment" not in by_issue

    threshold_rows = [
        dict(zip(THRESHOLD_RULE_COLUMNS, row))
        for row in values[threshold_title_row + 1 :]
        if row[0]
    ]
    assert threshold_rows
    temporal_metrics = {
        row["metric_or_field"]
        for row in threshold_rows
        if row["check_name"] == "keypoint_temporal"
    }
    assert {
        "joint_angle_change_deg_max",
        "rotation_delta_max",
        "joint_acceleration_m_s2_max",
        "joint_displacement_m_max",
    } <= temporal_metrics

    morphology_levels = {
        row["threshold_level"]
        for row in threshold_rows
        if row["metric_or_field"] == "bone_length_ratio_spread"
    }
    assert {"review", "fail"} <= morphology_levels
    assert {
        row["metric_or_field"]
        for row in threshold_rows
        if row["rule_type"] == "required_field"
    } >= {"scene", "task", "text_en"}
    assert {
        row["metric_or_field"]
        for row in threshold_rows
        if row["module"] == "video_quality"
    } >= {
        "black_frame_ratio",
        "laplacian_p10",
        "frozen_frame_ratio",
        "drop_frame_ratio",
    }
    assert {
        row["check_name"]
        for row in threshold_rows
        if row["parent_indicator"] == "abnormal_frame"
    } >= {"abnormal_v1", "abnormal_v2"}
    assert any(
        row["parent_indicator"] == "final_acceptance"
        and row["metric_or_field"] == "top_level_fail_count"
        and row["operator"] == ">="
        and row["effective_value"] == 1
        for row in threshold_rows
    )
    assert any(
        row["metric_or_field"] == "supplier_quality_signal"
        and row["rule_type"] == "informational_only"
        and "does not determine skeleton_static_status" in row["notes"]
        for row in threshold_rows
    )
    assert any(
        row["effective_value"] == "unresolved"
        and row["value_source"] == "unavailable"
        for row in threshold_rows
    )
    hard_metric_count = next(
        row
        for row in threshold_rows
        if row["config_key"]
        == "skeleton_quality_score.hard_exceeded_metric_count"
    )
    assert hard_metric_count["effective_value"] == 3
    assert hard_metric_count["value_source"] == "code_default"
    assert hard_metric_count["source_path"].endswith(
        "configs/precheck_example.yaml"
    )
    keys = [
        (
            row["module"],
            row["check_name"],
            row["metric_or_field"],
            row["threshold_level"],
            row["config_key"],
        )
        for row in threshold_rows
    ]
    assert len(keys) == len(set(keys))


def test_single_skeleton_static_indicator_fails_final_clip() -> None:
    assert (
        recompute_xjgt_final_status(
            text_check_status="pass",
            video_quality_status="pass",
            skeleton_static_status="fail",
            abnormal_frame_status="pass",
            expected_module_missing=False,
        )
        == "fail"
    )


def test_pending_optional_text_rule_does_not_force_review() -> None:
    assert (
        recompute_xjgt_final_status(
            text_check_status="pending_rule",
            video_quality_status="pass",
            skeleton_static_status="pass",
            abnormal_frame_status="pass",
            expected_module_missing=False,
        )
        == "pass"
    )


def test_xjgt_text_mapping_uses_real_evidence_by_default() -> None:
    assert resolve_xjgt_text_status(
        source_status="readable",
        mapped_checks={"text_integrity"},
        observed_status="pass",
        has_unmatched_source_rows=False,
        skip_xjgt_text=False,
    ) == ("pass", "mapped_text_integrity=pass")
    assert resolve_xjgt_text_status(
        source_status="readable",
        mapped_checks={"text_integrity"},
        observed_status="fail",
        has_unmatched_source_rows=False,
        skip_xjgt_text=False,
    ) == ("fail", "mapped_text_integrity=fail")


def test_xjgt_text_mapping_reports_unresolved_source_states() -> None:
    assert resolve_xjgt_text_status(
        source_status="missing",
        mapped_checks=set(),
        observed_status="not_run",
        has_unmatched_source_rows=False,
        skip_xjgt_text=False,
    ) == ("not_run", "text_integrity_source_missing")
    assert resolve_xjgt_text_status(
        source_status="unreadable:ArrowInvalid",
        mapped_checks=set(),
        observed_status="not_run",
        has_unmatched_source_rows=False,
        skip_xjgt_text=False,
    ) == ("review", "text_integrity_source_unreadable:ArrowInvalid")
    assert resolve_xjgt_text_status(
        source_status="readable",
        mapped_checks=set(),
        observed_status="not_run",
        has_unmatched_source_rows=True,
        skip_xjgt_text=False,
    ) == ("review", "text_integrity_source_readable_asset_unmatched")
    assert resolve_xjgt_text_status(
        source_status="readable",
        mapped_checks={"skeleton_quality_score"},
        observed_status="not_run",
        has_unmatched_source_rows=False,
        skip_xjgt_text=False,
    ) == ("review", "text_integrity_expected_check_missing")


def test_xjgt_text_can_be_explicitly_skipped() -> None:
    assert resolve_xjgt_text_status(
        source_status="readable",
        mapped_checks={"text_integrity"},
        observed_status="fail",
        has_unmatched_source_rows=False,
        skip_xjgt_text=True,
    ) == ("not_applicable", "text_integrity_skipped_by_cli")


def test_text_atomic_columns_are_discovered_from_producer_metrics() -> None:
    columns = detail_columns_for_rows(
        [{"text_field_present_supplier_instruction_status": "pass"}]
    )
    assert "text_field_present_supplier_instruction_status" in columns


def test_video_atomic_columns_only_use_available_metric_groups() -> None:
    statuses = video_quality_atomic_statuses(
        {
            "qc_summary": {"reasons": [], "warn_reasons": []},
            "video_quality": {
                "metrics": {
                    "exposure_metrics": {
                        "black_frame_ratio": 0.0,
                    }
                }
            },
        }
    )
    assert statuses == {"video_black_screen_status": "pass"}


def test_skeleton_static_status_is_strict_for_review_and_missing() -> None:
    assert skeleton_static_status_from_parts(
        missing_status="pass",
        morphology_status="review",
        fail_frame_ratio=0.0,
    ) == "fail"
    assert skeleton_static_status_from_parts(
        missing_status="review",
        morphology_status="pass",
        fail_frame_ratio=0.0,
    ) == "fail"
    assert skeleton_static_status_from_parts(
        missing_status="pass",
        morphology_status="not_run",
        fail_frame_ratio=0.0,
    ) == "not_ready"
    assert skeleton_static_status_from_parts(
        missing_status="fail",
        morphology_status="pass",
        fail_frame_ratio=0.01,
    ) == "fail"


def test_fail_indicator_count_has_exactly_four_top_level_dimensions() -> None:
    assert set(inspect.signature(count_fail_indicators).parameters) == {
        "text_check_status",
        "video_quality_status",
        "skeleton_static_status",
        "abnormal_frame_status",
    }
    assert count_fail_indicators(
        text_check_status="pass",
        video_quality_status="pass",
        skeleton_static_status="fail",
        abnormal_frame_status="pass",
    ) == 1


def test_acceptance_and_review_status_are_independent() -> None:
    acceptance, review, _reason = acceptance_and_review_status(
        text_check_status="pass",
        video_quality_status="pass",
        skeleton_static_status="pass",
        abnormal_frame_status="review",
        required_outputs_ready=True,
        unresolved=["abnormal_frame_status=review"],
    )
    assert acceptance == "pass"
    assert review == "review"

    acceptance, review, _reason = acceptance_and_review_status(
        text_check_status="fail",
        video_quality_status="fail",
        skeleton_static_status="pass",
        abnormal_frame_status="review",
        required_outputs_ready=True,
        unresolved=["abnormal_frame_status=review"],
    )
    assert acceptance == "fail"
    assert review == "review"


def test_review_workflow_does_not_reduce_acceptance_pass_count() -> None:
    details = [
        {
            "asset_id": f"asset_{index:03d}",
            "total_frames": 100,
            "acceptance_status": "pass",
            "review_status": "review" if index < 86 else "completed",
            "manual_review_status": "not_reviewed",
            "manual_problem_frame_count": 0,
            "abnormal_frame_status": "review" if index < 86 else "pass",
            "_problem_intervals": [],
        }
        for index in range(100)
    ]
    summary = supplier_summary_from_details(
        supplier_name="XJGT",
        expected_clip_count=100,
        details=details,
        main_issue_type="",
        modules_completed="",
        blocked_modules="",
        notes="",
    )
    assert summary["pass_clip_count"] == 100
    assert summary["review_clip_count"] == 86


def test_problem_frame_count_uses_interval_union() -> None:
    summary = supplier_summary_from_details(
        supplier_name="XJGT",
        expected_clip_count=1,
        details=[
            {
                "asset_id": "1001",
                "total_frames": 20,
                "acceptance_status": "pass",
                "review_status": "completed",
                "manual_review_status": "pass",
                "manual_problem_frame_count": 10,
                "abnormal_frame_status": "pass",
                "_problem_intervals": [(0, 9), (5, 14)],
            }
        ],
        main_issue_type="",
        modules_completed="",
        blocked_modules="",
        notes="",
    )
    assert summary["problem_frame_count"] == 15


def test_single_abnormal_indicator_fails_final_clip() -> None:
    assert (
        recompute_xjgt_final_status(
            text_check_status="pass",
            video_quality_status="pass",
            skeleton_static_status="pass",
            abnormal_frame_status="fail",
            expected_module_missing=False,
        )
        == "fail"
    )


def test_static_and_abnormal_indicators_fail_final_clip() -> None:
    assert (
        recompute_xjgt_final_status(
            text_check_status="pass",
            video_quality_status="pass",
            skeleton_static_status="fail",
            abnormal_frame_status="fail",
            expected_module_missing=False,
        )
        == "fail"
    )


def test_video_and_abnormal_indicators_fail_final_clip() -> None:
    assert (
        recompute_xjgt_final_status(
            text_check_status="pass",
            video_quality_status="fail",
            skeleton_static_status="pass",
            abnormal_frame_status="fail",
            expected_module_missing=False,
        )
        == "fail"
    )


def test_text_and_static_indicators_fail_final_clip() -> None:
    assert (
        recompute_xjgt_final_status(
            text_check_status="fail",
            video_quality_status="pass",
            skeleton_static_status="fail",
            abnormal_frame_status="pass",
            expected_module_missing=False,
        )
        == "fail"
    )


def test_deepreach_and_placeholders_remain_honest(
    tmp_path: Path, monkeypatch
) -> None:
    run_root = tmp_path / "acceptance_5x100"
    _prepare_inputs(run_root)
    _patch_frame_probe(monkeypatch)

    outputs = build_weekly_report(run_root)

    workbook = load_workbook(outputs.workbook_xlsx, read_only=True)
    deepreach_rows = list(workbook["DeepReach"].iter_rows(min_row=2, values_only=True))
    headers = [cell.value for cell in workbook["DeepReach"][1]]
    columns = {name: index for index, name in enumerate(headers)}
    assert len(deepreach_rows) == 8
    assert {row[columns["text_check_status"]] for row in deepreach_rows} == {
        "pending_rule"
    }
    assert {row[columns["skeleton_missing_status"]] for row in deepreach_rows} == {
        "blocked"
    }
    assert {row[columns["sam3_containment_status"]] for row in deepreach_rows} == {
        "blocked"
    }
    assert {row[columns["video_quality_status"]] for row in deepreach_rows} == {
        "no_valid_output"
    }
    assert {
        row[columns["supplier_quality_signal"]]
        for row in deepreach_rows
    } == {"not_provided"}
    assert {row[columns["acceptance_status"]] for row in deepreach_rows} == {
        "not_ready"
    }
    assert {row[columns["review_status_v2"]] for row in deepreach_rows} == {
        "blocked"
    }
    assert sum(row[columns["total_frames"]] for row in deepreach_rows) == sum(
        30 + index for index in range(8)
    )
    assert {row[columns["frame_count_status"]] for row in deepreach_rows} == {
        "ok"
    }

    with outputs.summary_csv.open(newline="", encoding="utf-8") as handle:
        summary = list(csv.DictReader(handle))
    assert summary[1]["sample_clip_count"] == "8"
    assert summary[1]["pass_clip_count"] == ""
    for row in summary[2:]:
        assert row["sample_clip_count"] == "0"
        assert row["pass_clip_count"] == ""


def test_generation_prints_sanity_checks(
    tmp_path: Path, monkeypatch, capsys
) -> None:
    run_root = tmp_path / "acceptance_5x100"
    _prepare_inputs(run_root)
    _patch_frame_probe(monkeypatch)

    build_weekly_report(run_root)

    output = capsys.readouterr().out
    assert "XJGT total_frame_count=400 ok=True" in output
    assert "XJGT problem_frame_count=101 ok=True" in output
    assert "XJGT manual_problem_ratio_nonzero=True" in output
    assert "XJGT acceptance_status counts={'fail': 3}" in output
    assert "XJGT review_status counts={'completed': 3}" in output
    assert "DeepReach sample_clip_count=8" in output
    assert "video_quality=no_valid_output" in output


def test_hard_issue_sheet_contains_required_rows(
    tmp_path: Path, monkeypatch
) -> None:
    run_root = tmp_path / "acceptance_5x100"
    _prepare_inputs(run_root)
    _patch_frame_probe(monkeypatch)

    outputs = build_weekly_report(run_root)

    workbook = load_workbook(outputs.workbook_xlsx, read_only=True)
    assert [cell.value for cell in workbook["人工与难测问题统计"][1]] == HARD_ISSUE_COLUMNS
    issue_types = {
        row[0]
        for row in workbook["人工与难测问题统计"].iter_rows(
            min_row=2, values_only=True
        )
    }
    assert issue_types == {
        "hand_out_of_frame",
        "severe_keypoint_offset",
        "skeleton_pose_hallucination",
        "visual_skeleton_presence_mismatch",
        "occlusion_or_mask_undersegmentation",
        "projection_review",
        "video_quality_pending_colleague_thresholds",
        "text_check_pending_rules",
    }


def test_issue_sheet_reconciles_video_quality_fail_clips(
    tmp_path: Path, monkeypatch
) -> None:
    run_root = tmp_path / "acceptance_5x100"
    _prepare_inputs(run_root)
    _patch_frame_probe(monkeypatch)
    _write_csv(
        run_root / "xjgt" / "video_quality" / "video_quality_acceptance_summary.csv",
        [
            {"asset_id": "1001", "status": "pass", "passed": True},
            {"asset_id": "1002", "status": "fail", "passed": False},
            {"asset_id": "1003", "status": "pass", "passed": True},
        ],
    )

    outputs = build_weekly_report(run_root)

    workbook = load_workbook(outputs.workbook_xlsx, read_only=True)
    columns = {name: index for index, name in enumerate(HARD_ISSUE_COLUMNS)}
    rows = {
        row[columns["issue_type"]]: row
        for row in workbook["人工与难测问题统计"].iter_rows(
            min_row=2, values_only=True
        )
    }
    video_row = rows["video_quality_pending_colleague_thresholds"]
    assert video_row[columns["observed_count"]] == 1
    assert video_row[columns["observed_unit"]] == "clip"
    assert video_row[columns["denominator"]] == 3
    assert video_row[columns["denominator_definition"]] == "XJGT sampled clips"
    assert video_row[columns["source_module"]] == "video_quality"
    assert video_row[columns["observed_ratio"]] == 1 / 3


def test_asset_id_normalization_matches_numeric_and_file_stems() -> None:
    assert normalize_asset_id("1001") == "1001"
    assert normalize_asset_id("1001.0") == "1001"
    assert normalize_asset_id("/videos/1001_video.mp4") == "1001"
    assert normalize_asset_id("/hdf5/1001_hdf5.hdf5") == "1001"


def test_input_audit_distinguishes_missing_and_unmatched_sources(
    tmp_path: Path,
) -> None:
    run_root = tmp_path / "acceptance_5x100"
    _write_csv(
        run_root / "manifests" / "supplier_manifest_xjgt_100.csv",
        [{"supplier_id": "xjgt", "asset_id": "1001", "episode_idx": 0}],
    )
    _write_parquet(
        run_root / "xjgt" / "precheck" / "check_results.parquet",
        [_check_result(99, "text_integrity", -1, {}, None)],
    )

    audit = collect_input_audit(run_root)
    by_name = {row["input_name"]: row for row in audit["inputs"]}

    assert by_name["precheck_check_results"]["read_status"] == "readable_unmatched"
    assert by_name["precheck_check_results"]["asset_overlap_count"] == 0
    assert by_name["precheck_check_results"]["unmapped_asset_examples"] == [
        "episode_idx:99"
    ]
    assert by_name["precheck_clip_aggregates"]["read_status"] == "missing"


def test_manual_aggregation_keeps_eleven_pass_and_six_fail_assets() -> None:
    labels: list[dict[str, object]] = []
    totals = {f"asset_{index:02d}": 100 for index in range(17)}
    for index in range(6):
        labels.append(
            {
                "asset_id": f"asset_{index:02d}",
                "manual_outcome": "true_positive",
                "window_start_frame": 0,
                "window_end_frame": 99,
                "affected_start_frame": 0,
                "affected_end_frame": 9,
            }
        )
    for index in range(6, 11):
        labels.append(
            {
                "asset_id": f"asset_{index:02d}",
                "manual_outcome": "false_positive",
                "window_start_frame": 0,
                "window_end_frame": 99,
            }
        )
    for index in range(11, 17):
        labels.append(
            {
                "asset_id": f"asset_{index:02d}",
                "manual_outcome": "acceptable_flagged",
                "window_start_frame": 0,
                "window_end_frame": 99,
            }
        )

    aggregated = aggregate_manual_review(labels, totals)

    statuses = [row["manual_review_status"] for row in aggregated.values()]
    assert statuses.count("fail") == 6
    assert statuses.count("pass") == 11
    assert sum(row["manual_problem_frame_count"] for row in aggregated.values()) == 60
    assert all(
        row["manual_problem_frame_count"] == 0
        for asset_id, row in aggregated.items()
        if asset_id >= "asset_06"
    )


def test_audit_mode_is_read_only(
    tmp_path: Path, capsys
) -> None:
    run_root = tmp_path / "acceptance_5x100"
    _prepare_inputs(run_root)

    assert main(["--run-root", str(run_root), "--audit-inputs"]) == 0

    output = capsys.readouterr().out
    assert "precheck_check_results" in output
    assert "distinct_checks" in output
    assert "mapped_weekly_fields" in output
    assert not (run_root / "weekly_supplier_acceptance_report.xlsx").exists()
    assert not (run_root / "weekly_supplier_summary.csv").exists()
